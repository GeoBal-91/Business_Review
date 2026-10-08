"""Final screenshot layout contracts and display-only boundaries."""
from pathlib import Path
import unittest
from pptx import Presentation
from pptx.util import Inches
from presentation_update_fixtures import ROOT
from web.release_bookends import build_cover,build_closing

class ReleaseBookends(unittest.TestCase):
    def slide(self):
        prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
        return prs,prs.slides.add_slide(prs.slide_layouts[6])

    def test_cover_dynamic_fields_and_long_names(self):
        for name in ['Invet Zoo','პეტფუდი — შინაური ცხოველების კვებისა და ვეტერინარული მომსახურების ცენტრი · München / São Paulo']:
            prs,slide=self.slide();build_cover(slide,name,'January 2027','PETS')
            shapes={s.name:s for s in slide.shapes}
            self.assertEqual(shapes['Presentation identity title'].text,name)
            self.assertEqual(shapes['Cover review heading'].text,'Business Review')
            self.assertEqual(shapes['Cover reporting period'].text,'January 2027')
            self.assertEqual(shapes['Cover SBU'].text,'Sub-Business Unit: PETS')
            self.assertEqual(str(slide.background.fill.fore_color.rgb),'FFFFFF')
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left,0);self.assertGreaterEqual(shape.top,0)
                self.assertLessEqual(shape.left+shape.width,prs.slide_width)
                self.assertLessEqual(shape.top+shape.height,prs.slide_height)
            self.assertLess(shapes['Presentation identity title'].top+shapes['Presentation identity title'].height,shapes['Cover review heading'].top)

    def test_closing_identity_exclusions_and_period(self):
        prs,slide=self.slide();build_closing(slide,prs,'December 2025')
        texts=[s.text.replace('\x0b','\n') for s in slide.shapes if s.has_text_frame and s.text]
        self.assertEqual(texts,['Thank You','For Your Partnership','Let’s Turn The Agreed Actions\nInto Measurable Progress','December 2025'])
        self.assertEqual(str(slide.background.fill.fore_color.rgb),'FFD636')
        self.assertEqual({s.name for s in slide.shapes if s.shape_type==13},{'Thank You decorative visual','Bookend Glovo logo'})
        image=next(s for s in slide.shapes if s.name=='Thank You decorative visual')
        self.assertEqual(image.image.blob,(ROOT/'backend/V9_Beta2/assets/Thank_You.png').read_bytes())
        self.assertFalse(any((image.crop_left,image.crop_top,image.crop_right,image.crop_bottom)))
