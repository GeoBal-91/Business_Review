"""Output guard boundaries, source immutability and additional workspace contracts."""
from copy import deepcopy
import unittest
from pptx import Presentation
from pptx.util import Inches
from refinement_fixtures import results, D, U
from web.meeting_workspace import output_view, build_plan, build_additional_plan

class FinalRefinement(unittest.TestCase):
    def test_current_month_guard_and_unchanged_analysis(self):
        _, original, _ = next(results())
        for orders, uptime, suppressed in [(0,float('nan'),True),(float('nan'),float('nan'),True),(0,0,False),(10,float('nan'),False)]:
            with self.subTest(orders=orders,uptime=uptime):
                triage=deepcopy(original)
                raw=triage.analysis.data.partner_raw
                mask=raw.Period.eq(triage.analysis.data.reporting_period)
                raw.loc[mask,D]=orders;raw.loc[mask,U]=uptime
                actions=deepcopy(triage.actions);reviews=deepcopy(triage.analysis.reviews)
                view=output_view(triage)
                self.assertEqual(view.actions,[] if suppressed else actions)
                self.assertEqual(view.analysis.reviews,[] if suppressed else reviews)
                self.assertEqual(triage.actions,actions);self.assertEqual(triage.analysis.reviews,reviews)
                for field in ['data','stores','metrics','opportunities','summary','actions']:
                    self.assertIs(getattr(view.analysis,field),getattr(triage.analysis,field))
                self.assertIs(view.stores,triage.stores)
                prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
                build_plan(prs,view);build_additional_plan(prs)
                self.assertEqual(len(prs.slides),1 if suppressed else 2)
                texts=[x.text for s in prs.slides for x in s.shapes if x.has_text_frame]
                self.assertEqual(texts.count('Additional agreed actions'),1)
                if suppressed:self.assertFalse(any(t.startswith('PRIORITY ACTION ') for t in texts))

    def test_absent_current_row_does_not_use_historical_uptime(self):
        _,triage,_=next(results());raw=triage.analysis.data.partner_raw
        triage.analysis.data.partner_raw=raw.loc[raw.Period.ne(triage.analysis.data.reporting_period)].copy()
        self.assertEqual(output_view(triage).actions,[])

    def test_three_outputs_only_with_exact_filtered_review_counts(self):
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from triage_presentation import render
        from web.presentation_update import apply_update,title
        from web.meeting_workspace import apply_meeting_workspace,review_summary
        _, original, _=next(results())
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'base.pptx';render(original,path)
            for orders,uptime,suppressed in [(0,float('nan'),True),(float('nan'),float('nan'),True),(0,0,False),(10,float('nan'),False)]:
                triage=deepcopy(original);raw=triage.analysis.data.partner_raw
                mask=raw.Period.eq(triage.analysis.data.reporting_period)
                raw.loc[mask,D]=orders;raw.loc[mask,U]=uptime
                prs=apply_update(Presentation(path),triage.analysis)
                untouched={id(s):s._element.xml for s in prs.slides if not title(s).startswith(('Priority Opportunities','Action Plan','Operational & Data Review'))}
                apply_meeting_workspace(prs,triage)
                selected=output_view(triage)
                for s in prs.slides:
                    if id(s) in untouched:self.assertEqual(s._element.xml,untouched[id(s)])
                    values={x.name:x.text for x in s.shapes if x.has_text_frame}
                    if title(s).startswith('Priority Opportunities'):
                        shown=[x for x in s.shapes if x.name.startswith('Opportunity store:')]
                        self.assertEqual(len(shown),0 if suppressed else 1)
                    if 'Review summary items to review' in values:
                        for label,n in zip(['items to review','stores affected','critical observations'],review_summary(selected.analysis.reviews)):
                            self.assertEqual(values['Review summary '+label],str(n))
                        self.assertFalse(any(x.name.endswith(' location') for x in s.shapes))
                self.assertTrue(any(title(s).startswith('Store Diagnostics') for s in prs.slides))
