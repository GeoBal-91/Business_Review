"""Content, order, formatting and non-mutation contracts for meeting slides."""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from pptx import Presentation
from pptx.util import Inches
from refinement_fixtures import results, LONG_STORE
from web.meeting_workspace import (build_plan,build_reviews,review_summary,review_identity,plan_pages,apply_meeting_workspace,output_view,build_additional_plan)
from web.presentation_update import build_combined,apply_update,title,COMBINED_TITLE
from presentation import context
from triage_presentation import render as frozen_render
from triage_exports import export_triage
from theme import GREEN


def deck():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5);return prs

class MeetingWorkspace(unittest.TestCase):
    @classmethod
    def setUpClass(cls): cls.cases=list(results())

    def test_series_marker_fill_and_outline_exactly_match_line(self):
        for _,triage,_ in self.cases:
            slide=build_combined(deck(),context(triage.analysis))
            for shape in slide.shapes:
                if not shape.has_chart:continue
                chart=shape.chart
                series=chart._chartSpace.xpath('.//c:lineChart/c:ser')[0]
                line=series.xpath('./c:spPr/a:ln/a:solidFill/a:srgbClr')[0].get('val')
                self.assertEqual(line,str(GREEN))
                self.assertEqual(series.xpath('./c:marker/c:spPr/a:solidFill/a:srgbClr')[0].get('val'),line)
                self.assertEqual(series.xpath('./c:marker/c:spPr/a:ln/a:solidFill/a:srgbClr')[0].get('val'),line)
                self.assertFalse(series.xpath('./c:dPt'))
                self.assertNotIn('D24646',series.xml)
                self.assertEqual(chart._chartSpace.xpath('./c:chart/c:dispBlanksAs')[0].get('val'),'gap')

    def test_all_selected_actions_and_followup_kpis_survive(self):
        counts=[]
        for name,triage,_ in self.cases:
            prs=deck();build_plan(prs,triage)
            counts.append(len(triage.actions))
            self.assertEqual(sum(len(p) for p in plan_pages(triage.actions)),len(triage.actions))
            shown=[]
            for s in prs.slides:
                values={x.name:x.text for x in s.shapes if x.has_text_frame}
                self.assertEqual(values['Next check-in'],'[To agree]')
                for a in triage.actions:
                    prefix=f"Plan action {a['Triage Rank']}"
                    if prefix+' store' not in values:continue
                    shown.append(a['Store ID'])
                    self.assertEqual(values[prefix+' store'],a['Store'])
                    self.assertEqual(values[prefix+' recommendation'],a['Action'])
                    self.assertEqual(values[prefix+' primary'],a['Metric'])
                    self.assertEqual(values[prefix+' Follow-up KPI'],a['Metric'])
                    for f in ['agreed','Owner','Contact','ETA']:
                        self.assertEqual(values[prefix+' '+f],'[To agree]')
            self.assertEqual(shown,[a['Store ID'] for a in triage.actions])
        self.assertTrue({0,1,2,6}.issubset(counts),counts)

    def test_review_population_evidence_actions_types_order_and_summary(self):
        types=set()
        for name,triage,_ in self.cases:
            result=triage.analysis;prs=deck();build_reviews(prs,result)
            seen=[]
            expected=(len(result.reviews),len({x['Store ID'] for x in result.reviews}),sum(x['Type']=='Critical observation' for x in result.reviews))
            self.assertEqual(review_summary(result.reviews),expected)
            for slide in prs.slides:
                values={x.name:x.text for x in slide.shapes if x.has_text_frame}
                for label,n in zip(['items to review','stores affected','critical observations'],expected):
                    self.assertEqual(values['Review summary '+label],str(n))
                for i,r in enumerate(result.reviews):
                    p=f'Review {i}'
                    if p+' store' not in values:continue
                    seen.append(i);types.add(r['Type'])
                    self.assertEqual(values[p+' evidence'],r['Issue'])
                    self.assertEqual(values[p+' action'],r['Action'])
                    self.assertEqual(values[p+' type'],r['Type'])
                    self.assertNotIn(str(r['Store ID']),'\n'.join(values.values()))
            self.assertEqual(seen,list(range(len(result.reviews))))
            if not result.reviews:self.assertIn('Review empty state',values)
            if name=='Actions_6':
                self.assertGreater(len(prs.slides),1)
                self.assertGreater(expected[0],expected[1])
                self.assertEqual([title(s) for s in prs.slides],[f'Operational & Data Review · {p}/{len(prs.slides)}' for p in range(1,len(prs.slides)+1)])
        self.assertTrue({'Critical observation','Availability review','Data quality','Opening-hours hygiene'}.issubset(types),types)

    def test_explicit_locations_unicode_and_ambiguous_metadata(self):
        _,triage,_=self.cases[0];a=triage.analysis
        r=next(r for r in a.reviews if r['Store ID']=='PRIVATE-ID-0')
        self.assertEqual(review_identity(a,r),(LONG_STORE,''))
        copied=deepcopy(a);raw=copied.data.partner_raw
        raw.drop(columns=['City','Country'],inplace=True)
        self.assertEqual(review_identity(copied,r)[1],'')
        self.assertEqual(review_identity(copied,dict(r,Store=r['Store ID']))[0],'Store name unavailable')
        raw['City']='Tbilisi';raw['Store City']='Batumi'
        self.assertEqual(review_identity(copied,r)[1],'')

    def test_flow_non_target_slides_and_all_audit_files_unchanged(self):
        # Use the full missing-data, multi-page fixture for end-to-end transformation.
        _,triage,_=next(c for c in self.cases if c[0]=='Actions_6')
        with TemporaryDirectory() as tmp:
            path=Path(tmp);frozen_render(triage,path/'before.pptx')
            prs=apply_update(Presentation(path/'before.pptx'),triage.analysis)
            before={id(s):s._element.xml for s in prs.slides if not title(s).startswith(('Action Plan','Operational & Data Review','Priority Opportunities & Actions')) or 'Store Coverage' in title(s)}
            actions=deepcopy(triage.actions);reviews=deepcopy(triage.analysis.reviews)
            export_triage(triage,path/'audit_before')
            apply_meeting_workspace(prs,triage)
            export_triage(triage,path/'audit_after')
            files=lambda directory:{str(p.relative_to(directory)):p.read_bytes() for p in directory.rglob('*') if p.is_file()}
            self.assertEqual(files(path/'audit_before'),files(path/'audit_after'))
            self.assertEqual(triage.actions,actions);self.assertEqual(triage.analysis.reviews,reviews)
            headings=[title(s) for s in prs.slides]
            d=headings.index('Discussion');t=headings.index('Thank You')
            self.assertTrue(all(h.startswith('Action Plan') for h in headings[d+1:t]))
            self.assertEqual(t-d-1,len(plan_pages(output_view(triage).actions))+1)
            self.assertTrue(all(i<d for i,h in enumerate(headings) if h.startswith('Operational & Data Review')))
            self.assertTrue(headings[t+1].startswith('APPENDIX'))
            for slide in prs.slides:
                if id(slide) in before:self.assertEqual(slide._element.xml,before[id(slide)])
            prs.save(path/'after.pptx')
            from zipfile import ZipFile
            with ZipFile(path/'after.pptx') as z:self.assertEqual(len(z.namelist()),len(set(z.namelist())))
            self.assertEqual([title(s) for s in Presentation(path/'after.pptx').slides],headings)

    def test_meeting_shapes_within_slide_bounds(self):
        for _,triage,_ in self.cases:
            prs=deck();build_plan(prs,triage);build_reviews(prs,triage.analysis)
            for slide in prs.slides:
                for shape in slide.shapes:
                    self.assertGreaterEqual(shape.left,0,shape.name);self.assertGreaterEqual(shape.top,0,shape.name)
                    self.assertLessEqual(shape.left+shape.width,prs.slide_width,shape.name)
                    self.assertLessEqual(shape.top+shape.height,prs.slide_height,shape.name)
