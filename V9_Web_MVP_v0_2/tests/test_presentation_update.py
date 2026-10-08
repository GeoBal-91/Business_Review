"""Contracts for the web presentation adapter, separate from frozen snapshots."""
from collections import Counter
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile
import numpy as np
import pandas as pd
from pandas.testing import assert_frame_equal
from pptx import Presentation
from pptx.util import Inches
from presentation_update_fixtures import cases, ROOT, BACKEND
from business_classification import classify
from triage_presentation import render as render_frozen
from triage_exports import export_triage
from presentation import context
from web.presentation_update import (apply_update, title, COMBINED_TITLE, trend_frame,
    growth_values, comparison_values, build_combined, build_partner_visual, build_growth_visual, pair_chart, combo_chart)


class PresentationUpdate(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cases = list(cases())
        cls.tmp = TemporaryDirectory()
        cls.addClassCleanup(cls.tmp.cleanup)
        cls.decks = {}
        for name, result, _ in cls.cases:
            path = Path(cls.tmp.name) / (name + '.pptx')
            render_frozen(classify(result), path)
            cls.decks[name] = path.read_bytes()

    def deck(self, name):
        return Presentation(BytesIO(self.decks[name]))

    def test_flow_and_original_slides_unchanged(self):
        for name, result, _ in self.cases:
            with self.subTest(name=name):
                prs = self.deck(name)
                before = {title(s): s._element.xml for s in prs.slides}
                by_part = {str(s.part.partname): s._element.xml for s in prs.slides}
                old_order = [title(s) for s in prs.slides]
                apply_update(prs, result)
                headings = [title(s) for s in prs.slides]
                for prefix in ('Partner vs ', 'Growth vs '):
                    original = prefix + result.settings.partner_sbu
                    self.assertEqual(headings[headings.index(original)+1], original + ' · Visual')
                    self.assertEqual(next(s._element.xml for s in prs.slides if title(s) == original), before[original])
                self.assertEqual(headings.count(COMBINED_TITLE), 1)
                self.assertFalse(any(h in {'Executive Summary', 'Growth Drivers', 'Conversion Performance', 'Partner Performance Evolution'}
                                     or h.startswith('Store Performance Matrix') for h in headings))
                retained = [h for h in old_order if h not in {'Executive Summary', 'Growth Drivers', 'Conversion Performance', 'Partner Performance Evolution'}
                            and not h.startswith('Store Performance Matrix')]
                self.assertEqual([h for h in headings if h in retained], retained)
                for slide in prs.slides:
                    heading = title(slide)
                    if heading in before and heading != 'Operational Health':
                        self.assertEqual(slide._element.xml, by_part[str(slide.part.partname)], heading)

    def test_render_does_not_mutate_analysis_or_any_audit(self):
        for name, result, _ in self.cases:
            with self.subTest(name=name), TemporaryDirectory() as tmp:
                triage = classify(result)
                export_triage(triage, Path(tmp) / 'before')
                apply_update(self.deck(name), result)
                export_triage(triage, Path(tmp) / 'after')
                before = {str(p.relative_to(Path(tmp)/'before')): p.read_bytes()
                          for p in (Path(tmp)/'before').rglob('*') if p.is_file()}
                after = {str(p.relative_to(Path(tmp)/'after')): p.read_bytes()
                         for p in (Path(tmp)/'after').rglob('*') if p.is_file()}
                self.assertEqual(before, after)

    def test_operational_preserves_all_text_and_adds_three_charts(self):
        for name, result, _ in self.cases:
            prs = self.deck(name)
            slide = next(s for s in prs.slides if title(s) == 'Operational Health')
            before = Counter(s.text for s in slide.shapes if s.has_text_frame)
            apply_update(prs, result)
            after = Counter(s.text for s in slide.shapes if s.has_text_frame)
            self.assertFalse(before - after)
            self.assertEqual(sum(s.has_chart for s in slide.shapes), 3)
            if result.settings.partner_sbu == 'FLOWERS':
                self.assertIn('excluded from business Priority', '\n'.join(after))

    def test_exact_comparison_values_and_growth(self):
        for _, result, _ in self.cases:
            ctx = context(result)
            self.assertEqual([r[0] for r in comparison_values(ctx)], ['AOV', 'CVR', 'Uptime', 'PNA', 'Cancels'])
            for _, key, _, _, partner, sbu in comparison_values(ctx):
                for got, expected in [(partner,ctx['partner_comparison'][key]),(sbu,ctx['sbu_comparison'][key])]:
                    self.assertTrue((pd.isna(got) and pd.isna(expected)) or got == expected)
            for (_, p, s), key in zip(growth_values(ctx), ['€ GMV', '# Delivered Orders', '€ AOV Derived']):
                for got, current, previous in [(p,ctx['partner_comparison'],ctx['partner_comp_previous']),
                                                (s,ctx['sbu_comparison'],ctx['sbu_comp_previous'])]:
                    if previous is None or pd.isna(current[key]) or pd.isna(previous[key]) or previous[key] == 0:
                        self.assertIsNone(got)
                    else:
                        self.assertEqual(got, current[key]/previous[key]-1)

    def test_combo_charts_have_partner_only_native_data_and_gap_month(self):
        for name, result, _ in self.cases:
            prs = apply_update(self.deck(name), result)
            slide = next(s for s in prs.slides if title(s) == COMBINED_TITLE)
            charts = [s.chart for s in slide.shapes if s.has_chart]
            self.assertEqual(len(charts), 2)
            self.assertEqual([[s.name for s in c.series] for c in charts],
                             [['GMV (€)', 'Orders'], ['Partner Sessions', 'Partner CVR']])
            frame = trend_frame(context(result))
            for chart, keys in zip(charts, [('€ GMV','# Delivered Orders'),('Sessions Store','% CVR')]):
                self.assertEqual(len(chart.plots), 2)
                self.assertEqual(chart._chartSpace.xpath('./c:chart/c:dispBlanksAs')[0].get('val'), 'gap')
                for series, key in zip(chart.series, keys):
                    expected = [None if pd.isna(v) else float(v) for v in frame[key]]
                    self.assertEqual(list(series.values), expected)
            if name == 'FLOWERS_partial_missing':
                self.assertEqual(frame.index.strftime('%b').tolist(), ['Jun','Jul','Aug'])
                self.assertTrue(frame.loc['2026-07-01'].isna().all())
                for chart in charts:
                    for series in chart.series:
                        self.assertIsNone(series.values[1])

    def test_missing_series_and_partial_benchmarks_are_not_zero(self):
        ctx = deepcopy(context(self.cases[0][1]))
        ctx['monthly']['% CVR'] = np.nan
        ctx['monthly']['Sessions Store'] = np.nan
        ctx['partner_comparison']['% CVR'] = np.nan
        ctx['sbu_comparison']['% CVR'] = np.nan
        ctx['current_partner_coverage_incomplete'] = True
        prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
        slide = build_combined(prs, ctx)
        lower = [s.chart for s in slide.shapes if s.has_chart][1]
        self.assertTrue(all(v is None for s in lower.series for v in s.values))
        visual = build_partner_visual(prs, ctx)
        text = '\n'.join(s.text for s in visual.shapes if s.has_text_frame)
        self.assertIn('Not assessed — incomplete store coverage.', text)
        self.assertIn('Not assessed', text)
        self.assertNotIn('better than', text)

    def test_native_workbooks_keep_missing_cells_and_hyperlinks_resolve(self):
        for name, result, _ in self.cases:
            prs = apply_update(self.deck(name), result)
            stream = BytesIO(); prs.save(stream)
            with ZipFile(stream) as z:
                self.assertIsNone(z.testzip())
                self.assertEqual(len(z.namelist()), len(set(z.namelist())))
            reopened = Presentation(BytesIO(stream.getvalue()))
            slide_parts = {s.part.partname for s in reopened.slides}
            for slide in reopened.slides:
                for rel in slide.part.rels.values():
                    if rel.reltype.endswith('/slide'):
                        self.assertIn(rel.target_part.partname, slide_parts)
                for shape in slide.shapes:
                    if shape.has_chart:
                        import openpyxl
                        book = openpyxl.load_workbook(BytesIO(shape.chart.part.chart_workbook.xlsx_part.blob))
                        self.assertTrue(book.active.max_row >= 3)
                        if name == 'FLOWERS_partial_missing' and title(slide) == COMBINED_TITLE:
                            self.assertIsNone(book.active.cell(3,2).value)
                            self.assertIsNone(book.active.cell(3,3).value)


class MeetingLayout(unittest.TestCase):
    def test_cards_values_units_wording_and_geometry(self):
        from web.presentation_update import meeting_values
        from formatting import relative_change, pp_change, get_row_value, format_currency, format_number, format_percent
        for _, result, _ in cases():
            ctx = context(result)
            latest, changes, gap, insights, bottom = meeting_values(ctx)
            for key in ['€ GMV', '# Delivered Orders', 'Sessions Store', '% CVR']:
                expected = (pp_change if key == '% CVR' else relative_change)(ctx['latest_partner'].get(key), get_row_value(ctx['partner_previous'], key))
                self.assertEqual(changes[key], expected)
            prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
            slide = build_combined(prs, ctx)
            charts = [s for s in slide.shapes if s.has_chart]
            self.assertEqual(len(charts), 2)
            self.assertTrue(all(s.left+s.width < Inches(7.45) for s in charts))
            for label, key, formatter in [('GMV', '€ GMV', format_currency), ('Orders', '# Delivered Orders', format_number), ('Partner Sessions', 'Sessions Store', format_number), ('Partner CVR', '% CVR', format_percent)]:
                value = next(s for s in slide.shapes if s.name==f'Meeting KPI {label} 2')
                self.assertEqual(value.text, formatter(latest.get(key)))
                self.assertGreaterEqual(value.left, Inches(7.45))
            texts='\n'.join(s.text for s in slide.shapes if s.has_text_frame)
            self.assertNotIn('Store Sessions',texts)
            self.assertNotIn('worse than',texts)
            self.assertIn(bottom,texts)
            for s in slide.shapes:
                self.assertGreaterEqual(s.left,0); self.assertGreaterEqual(s.top,0)
                self.assertLessEqual(s.left+s.width,prs.slide_width)
                self.assertLessEqual(s.top+s.height,prs.slide_height)

    def test_benchmark_directions_missing_and_incomplete(self):
        from web.presentation_update import meeting_values
        ctx = deepcopy(context(next(cases())[1]))
        ctx['current_partner_coverage_incomplete']=False
        ctx['PARTNER_SBU']='PETS'
        ctx['SBU_CVR']=.264
        ctx['latest_partner']['% CVR']=.148
        self.assertEqual(meeting_values(ctx)[2]['text'], '11.6 pp below PETS benchmark')
        ctx['latest_partner']['% CVR']=.3
        self.assertEqual(meeting_values(ctx)[2]['text'], '3.6 pp above PETS benchmark')
        ctx['latest_partner']['% CVR']=.264
        self.assertEqual(meeting_values(ctx)[2]['text'], 'In line with PETS benchmark')
        ctx['SBU_CVR']=np.nan
        self.assertEqual(meeting_values(ctx)[2]['text'], 'Benchmark unavailable')
        ctx['current_partner_coverage_incomplete']=True
        self.assertEqual(meeting_values(ctx)[2]['text'], 'Not assessed — incomplete store coverage.')
        ctx['partner_previous']=None
        self.assertTrue(all(v is None for v in meeting_values(ctx)[1].values()))


class ChartRendering(unittest.TestCase):
    def test_zero_counts_and_missing_percentage_do_not_get_misleading_axes(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = pd.DataFrame({'sessions': [0., 0.], 'cvr': [np.nan, np.nan]},
                             index=pd.date_range('2026-07-01', periods=2, freq='MS'))
        chart = combo_chart(slide, frame, 'sessions', 'cvr', 'Partner Sessions', 'Partner CVR', 1, percent=True)
        primary = chart._chartSpace.xpath('.//c:valAx')[0]
        self.assertEqual(float(primary.find('{*}scaling/{*}max').get('val')), 1)
        self.assertEqual(float(primary.find('{*}majorUnit').get('val')), 1)
        secondary = chart._chartSpace.xpath('.//c:valAx')[1]
        self.assertEqual(secondary.find('{*}tickLblPos').get('val'), 'none')
        self.assertEqual([x.tag.rsplit('}',1)[-1] for x in secondary.find('{*}scaling')], ['max','min'])
        series_tags = [x.tag.rsplit('}',1)[-1] for x in chart._chartSpace.xpath('.//c:lineChart/c:ser')[0]]
        self.assertLess(series_tags.index('spPr'), series_tags.index('marker'))
        self.assertEqual(list(chart.series[1].values), [None, None])
        self.assertEqual(list(chart.series[0].values), [0., 0.])

    def test_signed_pair_uses_series_fills_and_explicit_negative_handling(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        chart = pair_chart(slide, [.5, -.25], ['Partner', 'SBU'], .5, 1, 7, 1.2,
                           '+0.0%;-0.0%;0.0%', (-.5, 1))
        self.assertEqual([list(s.values) for s in chart.series], [[.5, None], [None, -.25]])
        self.assertEqual(chart.plots[0].overlap, 100)
        for series in chart.series:
            self.assertFalse(series.invert_if_negative)
            self.assertFalse(series._element.xpath('./c:dPt'))
        self.assertEqual(chart.category_axis._element.find('{*}tickLblPos').get('val'), 'low')
        self.assertEqual((chart.value_axis.minimum_scale, chart.value_axis.maximum_scale), (-.5, 1))
        self.assertIn('-25.0%', [s.text for s in slide.shapes if s.has_text_frame])

    def test_pair_preserves_unknown_and_true_zero_separately(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        chart = pair_chart(slide, [np.nan, 0], ['Partner', 'SBU'], .5, 1, 7, 1.2, '0.0%')
        self.assertEqual([list(s.values) for s in chart.series], [[None, None], [None, 0]])
        strings = [s.text for s in slide.shapes if s.has_text_frame]
        self.assertIn('Not assessed', strings)
        self.assertIn('0.0%', strings)


if __name__ == '__main__':
    unittest.main()
