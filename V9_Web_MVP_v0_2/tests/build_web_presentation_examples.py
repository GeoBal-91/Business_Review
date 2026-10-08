"""Generate reproducible web-adapter decks and byte-level audit parity evidence."""
from pathlib import Path
from tempfile import TemporaryDirectory
import hashlib
import json
from pptx import Presentation
from presentation_update_fixtures import ROOT, cases
from business_classification import classify
from triage_exports import export_triage
from web.presentation_update import render, title


def main():
    out = ROOT / 'validation' / 'presentation_update'
    out.mkdir(parents=True, exist_ok=True)
    records = []
    for name, result, source in cases():
        triage = classify(result)
        with TemporaryDirectory() as tmp:
            before = Path(tmp) / 'before'
            export_triage(triage, before)
            render(triage, out / (name + '.pptx'))
            after = Path(tmp) / 'after'
            export_triage(triage, after)
            hashes = lambda folder: {str(p.relative_to(folder)): hashlib.sha256(p.read_bytes()).hexdigest()
                                     for p in folder.rglob('*') if p.is_file()}
            first, second = hashes(before), hashes(after)
            assert first == second, name
        prs = Presentation(out / (name + '.pptx'))
        record = dict(file=name+'.pptx', source=source, audit_files_unchanged=len(first),
                      audit_sha256=first, slides=[dict(number=i,title=title(s)) for i,s in enumerate(prs.slides,1)])
        records.append(record)
        print(name, len(prs.slides), 'slides;', len(first), 'audit files unchanged', flush=True)
    (out/'reports'/'generation.json').write_text(json.dumps(records, indent=2))


if __name__ == '__main__':
    main()
