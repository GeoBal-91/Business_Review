import os
import math
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# ============================================================
# 1. SETTINGS
# ============================================================

INPUT_FILE = "partner_data.xlsx"
SUBBUSINESS_FILE = "subbusiness_unit_data.xlsx"

OUTPUT_FILE = "Partner_Business_Review_V8_1.pptx"

LOGO_FILE = "Logo.png"
PARTNER_LOGO_FILE = "Partner_Logo.png"

PARTNER_SBU = "BEAUTY_PERSONAL_CARE"

MIN_ORDERS_FOR_CVR_CLASSIFICATION = 10
CVR_TOLERANCE_PP = 0.01
TREND_MONTHS = 12

# V8.1 Opportunity Score settings
CVR_GAP_FLOOR = 0.01       # 1 pp
PNA_GAP_FLOOR = 0.02       # 2 pp
CANCELS_GAP_FLOOR = 0.01   # 1 pp
UPTIME_GAP_FLOOR = 0.02    # 2 pp

COMMERCIAL_MAX_SCORE = 40
CVR_MAX_SCORE = 25
PNA_MAX_SCORE = 14
CANCELS_MAX_SCORE = 14
UPTIME_MAX_SCORE = 7

TOP_OPPORTUNITIES_COUNT = 6

FONT = "Montserrat"


# ============================================================
# 2. MATPLOTLIB FONT
# ============================================================

available_fonts = {
    f.name
    for f in fm.fontManager.ttflist
}

if FONT in available_fonts:
    plt.rcParams["font.family"] = FONT
    print(f"Matplotlib font loaded: {FONT}")
else:
    print(
        f"WARNING: {FONT} was not found by Matplotlib. "
        "Charts may use a fallback font."
    )

plt.rcParams["axes.titleweight"] = "bold"


# ============================================================
# 3. COLORS
# ============================================================

YELLOW = RGBColor(255, 194, 68)
GREEN = RGBColor(0, 160, 130)
RED = RGBColor(210, 70, 70)
ORANGE = RGBColor(230, 145, 50)

DARK = RGBColor(35, 35, 35)
GREY = RGBColor(110, 110, 110)
LIGHT_GREY = RGBColor(242, 242, 242)
HEADER_GREY = RGBColor(235, 235, 235)
MID_GREY = RGBColor(220, 220, 220)
WHITE = RGBColor(255, 255, 255)

MPL_YELLOW = "#FFC244"
MPL_GREEN = "#00A082"
MPL_RED = "#D24646"
MPL_ORANGE = "#E69132"
MPL_GREY = "#8A8A8A"
MPL_DARK = "#232323"
MPL_LIGHT_GREY = "#F2F2F2"


# ============================================================
# 4. LOAD FILES
# ============================================================

print("Loading partner data...")
df = pd.read_excel(INPUT_FILE)

print("Loading Sub-Business Unit data...")
sbu_df = pd.read_excel(SUBBUSINESS_FILE)

df.columns = df.columns.astype(str).str.strip()
sbu_df.columns = sbu_df.columns.astype(str).str.strip()

df = df.loc[:, ~df.columns.str.startswith("Unnamed")]
sbu_df = sbu_df.loc[:, ~sbu_df.columns.str.startswith("Unnamed")]


# ============================================================
# 5. VALIDATE FILES
# ============================================================

partner_required = [
    "Month",
    "Year",
    "Store Address Address",
    "# Delivered Orders",
    "€ GMV",
]

sbu_required = [
    "Month",
    "Year",
    "Store Sub Business Unit",
    "# Delivered Orders",
    "€ AOV",
    "% CVR",
    "€ GMV",
    "% PNA Orders",
    "% Cancels DTP",
    "% Uptime",
]

missing_partner = [
    col for col in partner_required
    if col not in df.columns
]

missing_sbu = [
    col for col in sbu_required
    if col not in sbu_df.columns
]

if missing_partner:
    raise ValueError(
        "Missing required columns in partner file: "
        + ", ".join(missing_partner)
    )

if missing_sbu:
    raise ValueError(
        "Missing required columns in SBU file: "
        + ", ".join(missing_sbu)
    )


# ============================================================
# 6. DATE PREPARATION
# ============================================================

def prepare_period(dataframe):
    dataframe = dataframe.copy()

    if pd.api.types.is_datetime64_any_dtype(
        dataframe["Year"]
    ):
        dataframe["Year_Number"] = (
            dataframe["Year"].dt.year
        )

    else:
        dataframe["Year_Number"] = pd.to_numeric(
            dataframe["Year"],
            errors="coerce"
        )

        if dataframe["Year_Number"].isna().all():
            dataframe["Year_Number"] = pd.to_datetime(
                dataframe["Year"],
                errors="coerce"
            ).dt.year

    dataframe["Year_Number"] = (
        dataframe["Year_Number"].astype("Int64")
    )

    dataframe["Period"] = pd.to_datetime(
        dataframe["Month"].astype(str).str.strip()
        + " "
        + dataframe["Year_Number"].astype(str),
        format="%B %Y",
        errors="coerce"
    )

    return dataframe.dropna(
        subset=["Period"]
    )


df = prepare_period(df)
sbu_df = prepare_period(sbu_df)


# ============================================================
# 7. NUMERIC CLEANING
# ============================================================

partner_numeric_columns = [
    "# Delivered Orders",
    "€ GMV",
    "€ AOV",
    "% CVR",
    "% Prime Orders",
    "% PNA Orders",
    "% Cancels DTP",
    "# Avoidable WTP (min)",
    "# Available Products",
    "% Uptime",
]

for col in partner_numeric_columns:
    if col in df.columns:
        df[col] = pd.to_numeric(
            df[col],
            errors="coerce"
        )

sbu_numeric_columns = [
    "# Delivered Orders",
    "€ GMV",
    "€ AOV",
    "% CVR",
    "% PNA Orders",
    "% Cancels DTP",
    "% Uptime",
]

for col in sbu_numeric_columns:
    if col in sbu_df.columns:
        sbu_df[col] = pd.to_numeric(
            sbu_df[col],
            errors="coerce"
        )


# ============================================================
# 8. NORMALIZE PERCENTAGES
# ============================================================

partner_percentage_columns = [
    "% CVR",
    "% Prime Orders",
    "% PNA Orders",
    "% Cancels DTP",
    "% Uptime",
]

for col in partner_percentage_columns:
    if col in df.columns:
        values = df[col].dropna()

        if len(values) > 0 and values.median() > 1:
            df[col] = df[col] / 100


sbu_percentage_columns = [
    "% CVR",
    "% PNA Orders",
    "% Cancels DTP",
    "% Uptime",
]

for col in sbu_percentage_columns:
    if col in sbu_df.columns:
        values = sbu_df[col].dropna()

        if len(values) > 0 and values.median() > 1:
            sbu_df[col] = sbu_df[col] / 100


# ============================================================
# 9. FILTER SBU
# ============================================================

sbu_df["Store Sub Business Unit"] = (
    sbu_df["Store Sub Business Unit"]
    .astype(str)
    .str.strip()
    .str.upper()
)

PARTNER_SBU = PARTNER_SBU.strip().upper()

selected_sbu_raw = sbu_df[
    sbu_df["Store Sub Business Unit"]
    == PARTNER_SBU
].copy()

if selected_sbu_raw.empty:
    raise ValueError(
        f"SBU '{PARTNER_SBU}' not found."
    )


# ============================================================
# 10. SBU MONTHLY
# ============================================================

sbu_rules = {
    "# Delivered Orders": "sum",
    "€ GMV": "sum",
    "% CVR": "mean",
    "% PNA Orders": "mean",
    "% Cancels DTP": "mean",
    "% Uptime": "mean",
}

selected_sbu = (
    selected_sbu_raw
    .groupby("Period", as_index=False)
    .agg(sbu_rules)
    .sort_values("Period")
    .reset_index(drop=True)
)

selected_sbu["€ AOV Derived"] = (
    selected_sbu["€ GMV"]
    /
    selected_sbu[
        "# Delivered Orders"
    ].replace(0, np.nan)
)


# ============================================================
# 11. PARTNER MONTHLY
# ============================================================

partner_rules = {
    "# Delivered Orders": "sum",
    "€ GMV": "sum",
    "% CVR": "mean",
    "% Prime Orders": "mean",
    "% PNA Orders": "mean",
    "% Cancels DTP": "mean",
    "# Avoidable WTP (min)": "mean",
    "# Available Products": "sum",
    "% Uptime": "mean",
}

available_partner_rules = {
    col: rule
    for col, rule in partner_rules.items()
    if col in df.columns
}

monthly = (
    df
    .groupby("Period", as_index=False)
    .agg(available_partner_rules)
    .sort_values("Period")
    .reset_index(drop=True)
)

monthly["€ AOV Derived"] = (
    monthly["€ GMV"]
    /
    monthly[
        "# Delivered Orders"
    ].replace(0, np.nan)
)


# ============================================================
# 12. STORE MONTHLY
# ============================================================

store_rules = {
    "# Delivered Orders": "sum",
    "€ GMV": "sum",
    "% CVR": "mean",
    "% Prime Orders": "mean",
    "% PNA Orders": "mean",
    "% Cancels DTP": "mean",
    "# Avoidable WTP (min)": "mean",
    "# Available Products": "sum",
    "% Uptime": "mean",
}

available_store_rules = {
    col: rule
    for col, rule in store_rules.items()
    if col in df.columns
}

store_monthly = (
    df
    .groupby(
        [
            "Store Address Address",
            "Period"
        ],
        as_index=False
    )
    .agg(available_store_rules)
)

store_monthly["€ AOV Derived"] = (
    store_monthly["€ GMV"]
    /
    store_monthly[
        "# Delivered Orders"
    ].replace(0, np.nan)
)


# ============================================================
# 13. PERIOD HELPERS
# ============================================================

def get_previous_month_row(
    dataframe,
    current_period
):
    target = (
        current_period
        - pd.DateOffset(months=1)
    )

    match = dataframe[
        (
            dataframe["Period"].dt.year
            == target.year
        )
        &
        (
            dataframe["Period"].dt.month
            == target.month
        )
    ]

    if match.empty:
        return None

    return match.iloc[-1]


def get_yoy_row(
    dataframe,
    current_period
):
    target_year = current_period.year - 1
    target_month = current_period.month

    match = dataframe[
        (
            dataframe["Period"].dt.year
            == target_year
        )
        &
        (
            dataframe["Period"].dt.month
            == target_month
        )
    ]

    if match.empty:
        return None

    return match.sort_values(
        "Period"
    ).iloc[-1]


# ============================================================
# 14. PERIOD LOGIC
# ============================================================

partner_latest_period = monthly["Period"].max()
sbu_latest_period = selected_sbu["Period"].max()

latest_partner = monthly[
    monthly["Period"] == partner_latest_period
].iloc[0]

latest_sbu = selected_sbu[
    selected_sbu["Period"] == sbu_latest_period
].iloc[0]

partner_previous = get_previous_month_row(
    monthly,
    partner_latest_period
)

sbu_previous = get_previous_month_row(
    selected_sbu,
    sbu_latest_period
)

sbu_yoy_row = get_yoy_row(
    selected_sbu,
    sbu_latest_period
)

partner_yoy_row = get_yoy_row(
    monthly,
    partner_latest_period
)

common_periods = sorted(
    set(monthly["Period"])
    .intersection(
        set(selected_sbu["Period"])
    )
)

if not common_periods:
    raise ValueError(
        "Partner and SBU have no common periods."
    )

comparison_period = max(common_periods)

partner_comparison = monthly[
    monthly["Period"] == comparison_period
].iloc[0]

sbu_comparison = selected_sbu[
    selected_sbu["Period"] == comparison_period
].iloc[0]

partner_comp_previous = get_previous_month_row(
    monthly,
    comparison_period
)

sbu_comp_previous = get_previous_month_row(
    selected_sbu,
    comparison_period
)


# ============================================================
# 15. FORMATTERS
# ============================================================

def format_percent(value):
    if value is None or pd.isna(value):
        return ""
    return f"{value:.1%}"


def format_currency(value):
    if value is None or pd.isna(value):
        return ""
    return f"€{value:,.0f}"


def format_currency_2(value):
    if value is None or pd.isna(value):
        return ""
    return f"€{value:,.2f}"


def format_number(value):
    if value is None or pd.isna(value):
        return ""
    return f"{value:,.0f}"


def relative_change(current, previous):
    if (
        current is None
        or previous is None
        or pd.isna(current)
        or pd.isna(previous)
        or previous == 0
    ):
        return None

    return current / previous - 1


def pp_change(current, previous):
    if (
        current is None
        or previous is None
        or pd.isna(current)
        or pd.isna(previous)
    ):
        return None

    return current - previous


def get_row_value(row, column):
    if row is None:
        return None

    return row.get(
        column,
        None
    )


# ============================================================
# 16. DELTA LOGIC
# ============================================================

def delta_info(
    change,
    metric_type="relative",
    higher_is_better=True,
    comparison_label="vs prev. month"
):
    if change is None or pd.isna(change):
        return {
            "text": "",
            "color": GREY
        }

    if abs(change) < 0.000001:
        return {
            "text": f"No change {comparison_label}",
            "color": GREY
        }

    improved = (
        change > 0
        if higher_is_better
        else change < 0
    )

    arrow = "▲" if change > 0 else "▼"

    if metric_type == "pp":
        number = f"{abs(change) * 100:.1f} pp"
    else:
        number = f"{abs(change):.1%}"

    return {
        "text":
            f"{arrow} {number} {comparison_label}",

        "color":
            GREEN if improved else RED
    }


def benchmark_gap_info(
    partner_value,
    benchmark_value,
    metric_type="pp",
    higher_is_better=True,
    label="SBU"
):
    if (
        partner_value is None
        or benchmark_value is None
        or pd.isna(partner_value)
        or pd.isna(benchmark_value)
    ):
        return {
            "text": "",
            "color": GREY
        }

    gap = partner_value - benchmark_value

    if abs(gap) < 0.000001:
        return {
            "text": f"In line with {label}",
            "color": GREY
        }

    if metric_type == "pp":
        gap_display = (
            f"{abs(gap) * 100:.1f} pp"
        )

    else:
        if benchmark_value == 0:
            return {
                "text": "",
                "color": GREY
            }

        relative_gap = (
            partner_value
            / benchmark_value
            - 1
        )

        gap_display = (
            f"{abs(relative_gap):.1%}"
        )

    if higher_is_better:
        better = (
            partner_value
            > benchmark_value
        )
    else:
        better = (
            partner_value
            < benchmark_value
        )

    return {
        "text":
            (
                f"{gap_display} better than {label}"
                if better
                else
                f"{gap_display} worse than {label}"
            ),

        "color":
            GREEN if better else RED
    }


# ============================================================
# 17. BENCHMARKS
# ============================================================

def valid_positive_rate(value):
    return (
        value is not None
        and not pd.isna(value)
        and value > 0
    )


def valid_zero_allowed_rate(value):
    return (
        value is not None
        and not pd.isna(value)
        and value >= 0
    )


SBU_CVR = (
    latest_sbu["% CVR"]
    if valid_positive_rate(
        latest_sbu["% CVR"]
    )
    else None
)

SBU_UPTIME = (
    latest_sbu["% Uptime"]
    if valid_positive_rate(
        latest_sbu["% Uptime"]
    )
    else None
)

SBU_PNA = (
    latest_sbu["% PNA Orders"]
    if valid_zero_allowed_rate(
        latest_sbu["% PNA Orders"]
    )
    else None
)

SBU_CANCELS = (
    latest_sbu["% Cancels DTP"]
    if valid_zero_allowed_rate(
        latest_sbu["% Cancels DTP"]
    )
    else None
)

SBU_AOV = latest_sbu[
    "€ AOV Derived"
]


# ============================================================
# 18. SHARE OF SBU
# ============================================================

share_data = pd.merge(
    monthly[
        ["Period", "€ GMV"]
    ],
    selected_sbu[
        ["Period", "€ GMV"]
    ],
    on="Period",
    how="inner",
    suffixes=(
        "_Partner",
        "_SBU"
    )
)

share_data["Partner Share"] = (
    share_data["€ GMV_Partner"]
    /
    share_data[
        "€ GMV_SBU"
    ].replace(0, np.nan)
)

share_data = (
    share_data
    .sort_values("Period")
    .reset_index(drop=True)
)

latest_share_row = share_data.iloc[-1]

latest_share = latest_share_row[
    "Partner Share"
]

share_period = latest_share_row[
    "Period"
]

previous_share_row = get_previous_month_row(
    share_data,
    share_period
)

previous_share = (
    previous_share_row[
        "Partner Share"
    ]
    if previous_share_row is not None
    else None
)

share_yoy_row = get_yoy_row(
    share_data,
    share_period
)

share_yoy = (
    share_yoy_row[
        "Partner Share"
    ]
    if share_yoy_row is not None
    else None
)


# ============================================================
# 19. LATEST STORE DATA
# ============================================================

latest_stores = (
    store_monthly[
        store_monthly["Period"]
        == partner_latest_period
    ]
    .copy()
    .reset_index(drop=True)
)

active_latest_stores = (
    latest_stores[
        latest_stores[
            "# Delivered Orders"
        ] > 0
    ]
    .copy()
)

partner_median_gmv = (
    active_latest_stores[
        "€ GMV"
    ].median()
    if len(active_latest_stores) > 0
    else None
)


# ============================================================
# 20. CONFIDENCE
# ============================================================

def get_confidence(orders):
    if (
        orders is None
        or pd.isna(orders)
        or orders <= 0
    ):
        return ""

    if orders < 10:
        return "Low"

    if orders < 30:
        return "Medium"

    return "High"


# ============================================================
# 21. ISSUE CANDIDATES
# ============================================================

def build_issue_candidates(row):
    orders = row.get(
        "# Delivered Orders",
        0
    )

    cvr = row.get(
        "% CVR",
        np.nan
    )

    uptime = row.get(
        "% Uptime",
        np.nan
    )

    pna = row.get(
        "% PNA Orders",
        np.nan
    )

    cancels = row.get(
        "% Cancels DTP",
        np.nan
    )

    candidates = []

    # PNA: only material gaps of at least 2 pp
    if (
        pd.notna(pna)
        and SBU_PNA is not None
        and (pna - SBU_PNA) >= PNA_GAP_FLOOR
    ):
        severity = (
            (pna - SBU_PNA)
            /
            max(SBU_PNA, 0.01)
        )

        candidates.append({
            "severity": severity,
            "text":
                f"PNA {format_percent(pna)} "
                f"vs {format_percent(SBU_PNA)}",
            "focus":
                "Improve product availability"
        })

    # Cancels: only material gaps of at least 1 pp
    if (
        pd.notna(cancels)
        and SBU_CANCELS is not None
        and (cancels - SBU_CANCELS) >= CANCELS_GAP_FLOOR
    ):
        severity = (
            (cancels - SBU_CANCELS)
            /
            max(SBU_CANCELS, 0.01)
        )

        candidates.append({
            "severity": severity,
            "text":
                f"Cancels {format_percent(cancels)} "
                f"vs {format_percent(SBU_CANCELS)}",
            "focus":
                "Reduce cancellations"
        })

    # Uptime: only material gaps of at least 2 pp
    if (
        pd.notna(uptime)
        and SBU_UPTIME is not None
        and (SBU_UPTIME - uptime) >= UPTIME_GAP_FLOOR
    ):
        severity = (
            (SBU_UPTIME - uptime)
            /
            max(1 - SBU_UPTIME, 0.01)
        )

        candidates.append({
            "severity": severity,
            "text":
                f"Uptime {format_percent(uptime)} "
                f"vs {format_percent(SBU_UPTIME)}",
            "focus":
                "Improve store uptime"
        })

    # CVR: only with enough orders and at least 1 pp material gap
    if (
        orders >= MIN_ORDERS_FOR_CVR_CLASSIFICATION
        and pd.notna(cvr)
        and SBU_CVR is not None
        and (SBU_CVR - cvr) >= CVR_GAP_FLOOR
    ):
        severity = (
            (SBU_CVR - cvr)
            /
            max(SBU_CVR, 0.01)
        )

        candidates.append({
            "severity": severity,
            "text":
                f"CVR {format_percent(cvr)} "
                f"vs {format_percent(SBU_CVR)}",
            "focus":
                "Improve conversion"
        })

    return candidates

# ============================================================
# 22. STORE CLASSIFICATION
# ============================================================

def classify_store(row):
    orders = row.get(
        "# Delivered Orders",
        np.nan
    )

    gmv = row.get(
        "€ GMV",
        np.nan
    )

    cvr = row.get(
        "% CVR",
        np.nan
    )

    uptime = row.get(
        "% Uptime",
        np.nan
    )

    pna = row.get(
        "% PNA Orders",
        np.nan
    )

    cancels = row.get(
        "% Cancels DTP",
        np.nan
    )

    if (
        pd.isna(orders)
        or orders <= 0
    ):
        return pd.Series({
            "Commercial Status":
                "Inactive / No Orders",
            "Operational Status":
                "",
            "Main Issue":
                "No delivered orders",
            "Recommended Focus":
                "Review store activity",
            "Priority":
                "Medium",
            "Confidence":
                "",
            "Operational Issues":
                0
        })

    confidence = get_confidence(
        orders
    )

    operational_issues = []

    if (
        pd.notna(uptime)
        and SBU_UPTIME is not None
        and (SBU_UPTIME - uptime) >= UPTIME_GAP_FLOOR
    ):
        operational_issues.append(
            "Uptime"
        )

    if (
        pd.notna(pna)
        and SBU_PNA is not None
        and (pna - SBU_PNA) >= PNA_GAP_FLOOR
    ):
        operational_issues.append(
            "PNA"
        )

    if (
        pd.notna(cancels)
        and SBU_CANCELS is not None
        and (cancels - SBU_CANCELS) >= CANCELS_GAP_FLOOR
    ):
        operational_issues.append(
            "Cancels"
        )

    if len(operational_issues) >= 2:
        operational_status = (
            "Operational Risk"
        )

    elif len(operational_issues) == 1:
        operational_status = "Watch"

    else:
        operational_status = "Healthy"

    if (
        orders
        < MIN_ORDERS_FOR_CVR_CLASSIFICATION
    ):
        commercial_status = (
            "Low Volume / Insufficient Data"
        )

    else:
        high_gmv = (
            partner_median_gmv is not None
            and gmv >= partner_median_gmv
        )

        if (
            pd.isna(cvr)
            or SBU_CVR is None
        ):
            commercial_status = "Stable"

        else:
            strong_cvr = (
                cvr >= SBU_CVR
            )

            clearly_low_cvr = (
                cvr
                < SBU_CVR
                - CVR_TOLERANCE_PP
            )

            if high_gmv and strong_cvr:
                commercial_status = (
                    "Top Performer"
                )

            elif (
                high_gmv
                and clearly_low_cvr
            ):
                commercial_status = (
                    "Conversion Opportunity"
                )

            elif (
                not high_gmv
                and strong_cvr
            ):
                commercial_status = (
                    "Growth Potential"
                )

            elif (
                not high_gmv
                and clearly_low_cvr
            ):
                commercial_status = (
                    "Needs Attention"
                )

            else:
                commercial_status = (
                    "Stable"
                )

    candidates = build_issue_candidates(
        row
    )

    if candidates:
        main_issue_object = max(
            candidates,
            key=lambda x: x["severity"]
        )

        main_issue = (
            main_issue_object["text"]
        )

        recommended_focus = (
            main_issue_object["focus"]
        )

    elif (
        commercial_status
        == "Growth Potential"
    ):
        main_issue = (
            f"Strong CVR "
            f"{format_percent(cvr)}"
        )

        recommended_focus = (
            "Scale commercial activity"
        )

    elif (
        commercial_status
        == "Low Volume / Insufficient Data"
    ):
        main_issue = (
            f"Only {format_number(orders)} orders"
        )

        recommended_focus = (
            "Build volume before deeper diagnosis"
        )

    else:
        main_issue = (
            "No major issue identified"
        )

        recommended_focus = (
            "Maintain performance"
        )

    if (
        operational_status
        == "Operational Risk"
        or commercial_status
        == "Conversion Opportunity"
    ):
        priority = "High"

    elif commercial_status in [
        "Growth Potential",
        "Needs Attention",
        "Low Volume / Insufficient Data"
    ]:
        priority = "Medium"

    else:
        priority = "Low"

    return pd.Series({
        "Commercial Status":
            commercial_status,
        "Operational Status":
            operational_status,
        "Main Issue":
            main_issue,
        "Recommended Focus":
            recommended_focus,
        "Priority":
            priority,
        "Confidence":
            confidence,
        "Operational Issues":
            len(operational_issues)
    })


classification_results = latest_stores.apply(
    classify_store,
    axis=1
)

latest_stores = pd.concat(
    [
        latest_stores,
        classification_results
    ],
    axis=1
)


# ============================================================
# 23. OPPORTUNITY SCORE
# ============================================================

active_for_score = latest_stores[
    latest_stores[
        "# Delivered Orders"
    ] > 0
].copy()

if len(active_for_score) > 0:
    active_for_score[
        "GMV Percentile"
    ] = active_for_score[
        "€ GMV"
    ].rank(pct=True)

    active_for_score[
        "Orders Percentile"
    ] = active_for_score[
        "# Delivered Orders"
    ].rank(pct=True)

else:
    active_for_score[
        "GMV Percentile"
    ] = 0

    active_for_score[
        "Orders Percentile"
    ] = 0


def calculate_opportunity_score(row):
    orders = row.get(
        "# Delivered Orders",
        0
    )

    if (
        pd.isna(orders)
        or orders <= 0
    ):
        return 0

    store_name = row.get(
        "Store Address Address"
    )

    matching = active_for_score[
        active_for_score[
            "Store Address Address"
        ] == store_name
    ]

    if not matching.empty:
        gmv_percentile = (
            matching[
                "GMV Percentile"
            ].iloc[0]
        )

        orders_percentile = (
            matching[
                "Orders Percentile"
            ].iloc[0]
        )

        # Commercial impact = 80% GMV + 20% Orders, max 40 points
        commercial_index = (
            0.80 * gmv_percentile
            +
            0.20 * orders_percentile
        )

        commercial_component = (
            commercial_index
            * COMMERCIAL_MAX_SCORE
        )

    else:
        commercial_component = 0

    cvr = row.get(
        "% CVR",
        np.nan
    )

    pna = row.get(
        "% PNA Orders",
        np.nan
    )

    cancels = row.get(
        "% Cancels DTP",
        np.nan
    )

    uptime = row.get(
        "% Uptime",
        np.nan
    )

    # Material-gap eligibility gate
    cvr_material_gap = False
    pna_material_gap = False
    cancels_material_gap = False
    uptime_material_gap = False

    if (
        orders >= MIN_ORDERS_FOR_CVR_CLASSIFICATION
        and pd.notna(cvr)
        and SBU_CVR is not None
        and SBU_CVR > 0
    ):
        cvr_material_gap = (
            SBU_CVR - cvr
        ) >= CVR_GAP_FLOOR

    if (
        pd.notna(pna)
        and SBU_PNA is not None
    ):
        pna_material_gap = (
            pna - SBU_PNA
        ) >= PNA_GAP_FLOOR

    if (
        pd.notna(cancels)
        and SBU_CANCELS is not None
    ):
        cancels_material_gap = (
            cancels - SBU_CANCELS
        ) >= CANCELS_GAP_FLOOR

    if (
        pd.notna(uptime)
        and SBU_UPTIME is not None
    ):
        uptime_material_gap = (
            SBU_UPTIME - uptime
        ) >= UPTIME_GAP_FLOOR

    eligible = any([
        cvr_material_gap,
        pna_material_gap,
        cancels_material_gap,
        uptime_material_gap
    ])

    # No material KPI gap = no opportunity
    if not eligible:
        return 0

    # CVR component: max 25 points
    cvr_component = 0

    if cvr_material_gap:
        cvr_gap_ratio = (
            (SBU_CVR - cvr)
            /
            SBU_CVR
        )

        cvr_component = min(
            CVR_MAX_SCORE,
            cvr_gap_ratio * 50
        )

    # PNA component: max 14 points
    pna_component = 0

    if (
        pna_material_gap
        and SBU_PNA is not None
        and SBU_PNA > 0
    ):
        pna_gap_ratio = (
            pna / SBU_PNA - 1
        )

        pna_component = min(
            PNA_MAX_SCORE,
            pna_gap_ratio * 7
        )

    # If SBU PNA is exactly zero, a material positive gap gets max severity.
    elif (
        pna_material_gap
        and SBU_PNA == 0
    ):
        pna_component = PNA_MAX_SCORE

    # Cancels component: max 14 points
    cancels_component = 0

    if (
        cancels_material_gap
        and SBU_CANCELS is not None
        and SBU_CANCELS > 0
    ):
        cancels_gap_ratio = (
            cancels / SBU_CANCELS - 1
        )

        cancels_component = min(
            CANCELS_MAX_SCORE,
            cancels_gap_ratio * 7
        )

    # If SBU Cancels is exactly zero, a material positive gap gets max severity.
    elif (
        cancels_material_gap
        and SBU_CANCELS == 0
    ):
        cancels_component = CANCELS_MAX_SCORE

    # Uptime component: max 7 points.
    # A 10 pp gap reaches the cap.
    uptime_component = 0

    if uptime_material_gap:
        uptime_gap = (
            SBU_UPTIME - uptime
        )

        uptime_component = min(
            UPTIME_MAX_SCORE,
            uptime_gap * 70
        )

    operational_component = (
        pna_component
        + cancels_component
        + uptime_component
    )

    total_score = (
        commercial_component
        + cvr_component
        + operational_component
    )

    return round(
        min(total_score, 100),
        0
    )

latest_stores[
    "Opportunity Score"
] = latest_stores.apply(
    calculate_opportunity_score,
    axis=1
)


# ============================================================
# 24. DIAGNOSTICS
# ============================================================

def store_diagnostic(row):
    orders = row.get(
        "# Delivered Orders",
        0
    )

    if orders <= 0:
        return (
            "No active commercial performance",
            "No delivered orders",
            "Review store activity"
        )

    cvr = row.get(
        "% CVR",
        np.nan
    )

    uptime = row.get(
        "% Uptime",
        np.nan
    )

    strengths = []

    if (
        pd.notna(cvr)
        and SBU_CVR is not None
        and cvr >= SBU_CVR
        and orders
        >= MIN_ORDERS_FOR_CVR_CLASSIFICATION
    ):
        strengths.append(
            "CVR above SBU"
        )

    if (
        pd.notna(uptime)
        and SBU_UPTIME is not None
        and uptime >= SBU_UPTIME
    ):
        strengths.append(
            "Uptime above SBU"
        )

    if not strengths:
        strengths.append(
            "Active order volume"
        )

    risk = row.get(
        "Main Issue",
        "No major KPI gap"
    )

    if risk == "No major issue identified":
        risk = "No major KPI gap"

    focus = row.get(
        "Recommended Focus",
        "Maintain performance"
    )

    return (
        strengths[0],
        risk,
        focus
    )


diagnostics = latest_stores.apply(
    store_diagnostic,
    axis=1
)

latest_stores["Strength"] = [
    x[0] for x in diagnostics
]

latest_stores["Risk"] = [
    x[1] for x in diagnostics
]

latest_stores["Focus"] = [
    x[2] for x in diagnostics
]


# ============================================================
# 25. POWERPOINT SETUP
# ============================================================

prs = Presentation()

prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


# ============================================================
# 26. POWERPOINT HELPERS
# ============================================================

def set_background(
    slide,
    color=WHITE
):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=12,
    bold=False,
    color=DARK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP
):
    if text is None:
        return None

    text = str(text)

    hidden_values = {
        "N/A",
        "YOY N/A",
        "BENCHMARK N/A",
        "NONE",
        "NAN"
    }

    if text.strip().upper() in hidden_values:
        return None

    if text.strip() == "":
        return None

    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign

    p = tf.paragraphs[0]
    p.text = text

    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align

    return box


def add_title(
    slide,
    title,
    subtitle=None
):
    add_text(
        slide,
        title,
        0.65,
        0.35,
        12,
        0.5,
        size=20,
        bold=True
    )

    if subtitle:
        add_text(
            slide,
            subtitle,
            0.67,
            0.9,
            11.8,
            0.35,
            size=10,
            color=GREY
        )


def add_footer(slide):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(7.37),
        Inches(13.333),
        Inches(0.13)
    )

    line.fill.solid()
    line.fill.fore_color.rgb = YELLOW
    line.line.fill.background()


def add_kpi_card(
    slide,
    x,
    y,
    w,
    h,
    title,
    value,
    delta_text=None,
    delta_color=GREY,
    secondary_text=None,
    secondary_color=GREY
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GREY
    shape.line.color.rgb = LIGHT_GREY

    add_text(
        slide,
        title,
        x + 0.2,
        y + 0.14,
        w - 0.4,
        0.28,
        size=9,
        color=GREY
    )

    add_text(
        slide,
        value,
        x + 0.2,
        y + 0.46,
        w - 0.4,
        0.43,
        size=19,
        bold=True
    )

    if delta_text:
        add_text(
            slide,
            delta_text,
            x + 0.2,
            y + 0.95,
            w - 0.4,
            0.27,
            size=8.5,
            bold=True,
            color=delta_color
        )

    if secondary_text:
        add_text(
            slide,
            secondary_text,
            x + 0.2,
            y + 1.26,
            w - 0.4,
            0.27,
            size=8,
            bold=True,
            color=secondary_color
        )


def add_section_card(
    slide,
    x,
    y,
    w,
    h,
    title,
    body
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GREY
    shape.line.color.rgb = LIGHT_GREY

    add_text(
        slide,
        title,
        x + 0.2,
        y + 0.15,
        w - 0.4,
        0.3,
        size=10,
        bold=True
    )

    add_text(
        slide,
        body,
        x + 0.2,
        y + 0.55,
        w - 0.4,
        h - 0.7,
        size=9,
        color=GREY
    )


def add_takeaway(
    slide,
    text
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.75),
        Inches(6.52),
        Inches(11.85),
        Inches(0.55)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GREY
    shape.line.color.rgb = LIGHT_GREY

    add_text(
        slide,
        text,
        1.0,
        6.67,
        11.3,
        0.25,
        size=9.5,
        bold=True
    )


def save_chart(
    fig,
    filename
):
    fig.savefig(
        filename,
        bbox_inches="tight",
        dpi=180,
        facecolor="white"
    )

    plt.close(fig)


def add_chart_image(
    slide,
    filename,
    x,
    y,
    w,
    h
):
    if os.path.exists(filename):
        slide.shapes.add_picture(
            filename,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h)
        )


# ============================================================
# 27. TABLE HELPERS
# ============================================================

def add_table_header(
    slide,
    headers,
    x_positions,
    widths,
    y,
    height=0.5,
    font_size=14
):
    total_x = min(x_positions)

    total_right = max(
        x + w
        for x, w
        in zip(
            x_positions,
            widths
        )
    )

    bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(total_x - 0.12),
        Inches(y - 0.08),
        Inches(
            total_right
            - total_x
            + 0.24
        ),
        Inches(height + 0.15)
    )

    bg.fill.solid()
    bg.fill.fore_color.rgb = HEADER_GREY
    bg.line.color.rgb = HEADER_GREY

    for i, header in enumerate(headers):
        add_text(
            slide,
            header,
            x_positions[i],
            y,
            widths[i],
            height,
            size=font_size,
            bold=True,
            color=DARK,
            valign=MSO_ANCHOR.MIDDLE
        )

    separator = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(total_x),
        Inches(y + height + 0.04),
        Inches(
            total_right
            - total_x
        ),
        Inches(0.015)
    )

    separator.fill.solid()
    separator.fill.fore_color.rgb = MID_GREY
    separator.line.fill.background()


# ============================================================
# 28. COVER
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(
    slide,
    YELLOW
)

if os.path.exists(LOGO_FILE):
    slide.shapes.add_picture(
        LOGO_FILE,
        Inches(0.75),
        Inches(0.55),
        width=Inches(1.65)
    )
else:
    print(
        f"WARNING: {LOGO_FILE} not found."
    )


if os.path.exists(PARTNER_LOGO_FILE):
    slide.shapes.add_picture(
        PARTNER_LOGO_FILE,
        Inches(10.55),
        Inches(0.55),
        width=Inches(2.0),
        height=Inches(0.95)
    )

else:
    partner_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(10.45),
        Inches(0.48),
        Inches(2.15),
        Inches(1.05)
    )

    partner_box.fill.solid()
    partner_box.fill.fore_color.rgb = YELLOW

    partner_box.line.color.rgb = RGBColor(
        225,
        166,
        35
    )

    add_text(
        slide,
        "PARTNER LOGO",
        10.55,
        0.82,
        1.95,
        0.3,
        size=8,
        bold=True,
        color=RGBColor(
            160,
            115,
            20
        ),
        align=PP_ALIGN.CENTER
    )


add_text(
    slide,
    "Partner Business Review",
    0.8,
    2.35,
    11.5,
    0.8,
    size=30,
    bold=True
)

partner_period_text = (
    f"{monthly['Period'].min().strftime('%B %Y')}"
    f" – "
    f"{partner_latest_period.strftime('%B %Y')}"
)

add_text(
    slide,
    partner_period_text,
    0.82,
    3.25,
    10,
    0.5,
    size=16
)

add_text(
    slide,
    f"Sub-Business Unit: {PARTNER_SBU}",
    0.82,
    3.75,
    10,
    0.4,
    size=12
)


# ============================================================
# 29. SBU BUSINESS DEVELOPMENT
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    f"{PARTNER_SBU} Business Development",
    (
        f"Latest SBU performance — "
        f"{sbu_latest_period.strftime('%B %Y')}"
    )
)

add_footer(slide)


sbu_gmv = latest_sbu["€ GMV"]

sbu_orders = latest_sbu[
    "# Delivered Orders"
]

sbu_aov = latest_sbu[
    "€ AOV Derived"
]

sbu_cvr = latest_sbu[
    "% CVR"
]


sbu_gmv_mom_change = relative_change(
    sbu_gmv,
    get_row_value(
        sbu_previous,
        "€ GMV"
    )
)

sbu_orders_mom_change = relative_change(
    sbu_orders,
    get_row_value(
        sbu_previous,
        "# Delivered Orders"
    )
)

sbu_aov_mom_change = relative_change(
    sbu_aov,
    get_row_value(
        sbu_previous,
        "€ AOV Derived"
    )
)

sbu_cvr_mom_change = pp_change(
    sbu_cvr,
    get_row_value(
        sbu_previous,
        "% CVR"
    )
)


sbu_gmv_mom = delta_info(
    sbu_gmv_mom_change
)

sbu_orders_mom = delta_info(
    sbu_orders_mom_change
)

sbu_aov_mom = delta_info(
    sbu_aov_mom_change
)

sbu_cvr_mom = delta_info(
    sbu_cvr_mom_change,
    metric_type="pp"
)


sbu_gmv_yoy = delta_info(
    relative_change(
        sbu_gmv,
        get_row_value(
            sbu_yoy_row,
            "€ GMV"
        )
    ),
    comparison_label="YoY"
)

sbu_orders_yoy = delta_info(
    relative_change(
        sbu_orders,
        get_row_value(
            sbu_yoy_row,
            "# Delivered Orders"
        )
    ),
    comparison_label="YoY"
)

sbu_aov_yoy = delta_info(
    relative_change(
        sbu_aov,
        get_row_value(
            sbu_yoy_row,
            "€ AOV Derived"
        )
    ),
    comparison_label="YoY"
)

sbu_yoy_cvr_value = get_row_value(
    sbu_yoy_row,
    "% CVR"
)

if valid_positive_rate(
    sbu_yoy_cvr_value
):
    sbu_cvr_yoy = delta_info(
        pp_change(
            sbu_cvr,
            sbu_yoy_cvr_value
        ),
        metric_type="pp",
        comparison_label="YoY"
    )
else:
    sbu_cvr_yoy = {
        "text": "",
        "color": GREY
    }


card_x = [
    0.45,
    3.5,
    6.55,
    9.6
]

add_kpi_card(
    slide,
    card_x[0],
    1.45,
    2.85,
    1.75,
    "SBU GMV",
    format_currency(sbu_gmv),
    sbu_gmv_mom["text"],
    sbu_gmv_mom["color"],
    sbu_gmv_yoy["text"],
    sbu_gmv_yoy["color"]
)

add_kpi_card(
    slide,
    card_x[1],
    1.45,
    2.85,
    1.75,
    "SBU Orders",
    format_number(sbu_orders),
    sbu_orders_mom["text"],
    sbu_orders_mom["color"],
    sbu_orders_yoy["text"],
    sbu_orders_yoy["color"]
)

add_kpi_card(
    slide,
    card_x[2],
    1.45,
    2.85,
    1.75,
    "SBU AOV",
    format_currency_2(sbu_aov),
    sbu_aov_mom["text"],
    sbu_aov_mom["color"],
    sbu_aov_yoy["text"],
    sbu_aov_yoy["color"]
)

add_kpi_card(
    slide,
    card_x[3],
    1.45,
    2.85,
    1.75,
    "SBU CVR",
    format_percent(sbu_cvr),
    sbu_cvr_mom["text"],
    sbu_cvr_mom["color"],
    sbu_cvr_yoy["text"],
    sbu_cvr_yoy["color"]
)


# Growth Composition

add_text(
    slide,
    "Growth Composition",
    0.65,
    3.65,
    5.4,
    0.35,
    size=12,
    bold=True
)

growth_box = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.65),
    Inches(4.08),
    Inches(5.6),
    Inches(2.15)
)

growth_box.fill.solid()
growth_box.fill.fore_color.rgb = LIGHT_GREY
growth_box.line.color.rgb = LIGHT_GREY


growth_items = [
    (
        "GMV",
        sbu_gmv_mom_change
    ),
    (
        "Orders",
        sbu_orders_mom_change
    ),
    (
        "AOV",
        sbu_aov_mom_change
    )
]

growth_y = 4.35

for name, change in growth_items:

    if change is None:
        continue

    value_text = f"{change:+.1%}"

    value_color = (
        GREEN
        if change >= 0
        else RED
    )

    add_text(
        slide,
        name,
        0.95,
        growth_y,
        1.35,
        0.35,
        size=10,
        bold=True
    )

    add_text(
        slide,
        value_text,
        2.25,
        growth_y,
        1.1,
        0.35,
        size=14,
        bold=True,
        color=value_color
    )

    bar_x = 3.35
    bar_y = growth_y + 0.08

    max_width = 2.25

    magnitude = min(
        abs(change) / 0.15,
        1
    )

    bar_width = max(
        0.12,
        magnitude * max_width
    )

    bar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(bar_x),
        Inches(bar_y),
        Inches(bar_width),
        Inches(0.16)
    )

    bar.fill.solid()
    bar.fill.fore_color.rgb = (
        GREEN if change >= 0 else RED
    )

    bar.line.fill.background()

    growth_y += 0.53


# Key Takeaways

add_text(
    slide,
    "Key Takeaways",
    6.75,
    3.65,
    5.6,
    0.35,
    size=12,
    bold=True
)

takeaway_box = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(6.75),
    Inches(4.08),
    Inches(5.9),
    Inches(2.15)
)

takeaway_box.fill.solid()
takeaway_box.fill.fore_color.rgb = LIGHT_GREY
takeaway_box.line.color.rgb = LIGHT_GREY


sbu_takeaways = []

if sbu_gmv_mom_change is not None:
    if sbu_gmv_mom_change >= 0:
        sbu_takeaways.append(
            f"GMV increased "
            f"{abs(sbu_gmv_mom_change):.1%} "
            f"vs previous month."
        )
    else:
        sbu_takeaways.append(
            f"GMV decreased "
            f"{abs(sbu_gmv_mom_change):.1%} "
            f"vs previous month."
        )


if (
    sbu_orders_mom_change is not None
    and sbu_aov_mom_change is not None
):
    if (
        sbu_orders_mom_change > 0
        and sbu_aov_mom_change < 0
    ):
        sbu_takeaways.append(
            "Growth was volume-led: orders increased "
            "while average basket softened."
        )

    elif (
        sbu_orders_mom_change > 0
        and sbu_aov_mom_change > 0
    ):
        sbu_takeaways.append(
            "Both order volume and average basket "
            "supported growth."
        )

    elif (
        sbu_orders_mom_change < 0
        and sbu_aov_mom_change > 0
    ):
        sbu_takeaways.append(
            "Higher basket size partly offset "
            "lower order volume."
        )

    else:
        sbu_takeaways.append(
            "Both order volume and basket size "
            "declined vs previous month."
        )


if sbu_cvr_mom_change is not None:
    if sbu_cvr_mom_change > 0:
        sbu_takeaways.append(
            f"CVR improved by "
            f"{abs(sbu_cvr_mom_change) * 100:.1f} pp "
            f"to {format_percent(sbu_cvr)}."
        )

    elif sbu_cvr_mom_change < 0:
        sbu_takeaways.append(
            f"CVR declined by "
            f"{abs(sbu_cvr_mom_change) * 100:.1f} pp "
            f"to {format_percent(sbu_cvr)}."
        )


takeaway_y = 4.37

for text in sbu_takeaways[:3]:

    add_text(
        slide,
        "• " + text,
        7.05,
        takeaway_y,
        5.25,
        0.45,
        size=9.5
    )

    takeaway_y += 0.55


# ============================================================
# 30. SBU PERFORMANCE TREND
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    f"{PARTNER_SBU} Performance Trend",
    f"Last {TREND_MONTHS} available months"
)

add_footer(slide)

sbu_trend = (
    selected_sbu
    .tail(TREND_MONTHS)
    .copy()
)

sbu_labels = (
    sbu_trend[
        "Period"
    ]
    .dt.strftime("%b %y")
)

fig, ax1 = plt.subplots(
    figsize=(10.5, 4.4)
)

ax1.bar(
    sbu_labels,
    sbu_trend["€ GMV"],
    color=MPL_YELLOW,
    alpha=0.85,
    label="GMV (€)"
)

ax1.set_ylabel("GMV (€)")
ax1.spines["top"].set_visible(False)

ax2 = ax1.twinx()

ax2.plot(
    sbu_labels,
    sbu_trend[
        "# Delivered Orders"
    ],
    marker="o",
    linewidth=2.5,
    color=MPL_GREEN,
    label="Orders"
)

ax2.set_ylabel("Orders")
ax2.spines["top"].set_visible(False)

plt.xticks(rotation=35)

plt.title(
    f"{PARTNER_SBU} — GMV & Orders",
    loc="left",
    fontsize=14,
    fontweight="bold"
)

handles1, labels1 = ax1.get_legend_handles_labels()
handles2, labels2 = ax2.get_legend_handles_labels()

ax1.legend(
    handles1 + handles2,
    labels1 + labels2,
    loc="upper center",
    bbox_to_anchor=(0.5, 1.02),
    ncol=2,
    frameon=False,
    fontsize=9
)

plt.tight_layout()

chart_file = "sbu_performance_trend.png"

save_chart(
    fig,
    chart_file
)

add_chart_image(
    slide,
    chart_file,
    0.8,
    1.35,
    11.8,
    4.75
)

if len(sbu_trend) >= 2:

    first_row = sbu_trend.iloc[0]
    last_row = sbu_trend.iloc[-1]

    trend_gmv_change = relative_change(
        last_row["€ GMV"],
        first_row["€ GMV"]
    )

    trend_orders_change = relative_change(
        last_row[
            "# Delivered Orders"
        ],
        first_row[
            "# Delivered Orders"
        ]
    )

    if (
        trend_gmv_change is not None
        and trend_orders_change is not None
    ):
        add_takeaway(
            slide,
            (
                f"Across the displayed period, "
                f"{PARTNER_SBU} GMV changed "
                f"{trend_gmv_change:+.1%}, while "
                f"orders changed "
                f"{trend_orders_change:+.1%}."
            )
        )


# ============================================================
# 31. PARTNER VS SBU
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    f"Partner vs {PARTNER_SBU}",
    (
        f"Comparable period: "
        f"{comparison_period.strftime('%B %Y')}"
    )
)

add_footer(slide)


comparison_metrics = [
    (
        "AOV",
        "€ AOV Derived",
        "currency",
        True
    ),
    (
        "CVR",
        "% CVR",
        "percent",
        True
    ),
    (
        "Uptime",
        "% Uptime",
        "percent",
        True
    ),
    (
        "PNA",
        "% PNA Orders",
        "percent",
        False
    ),
    (
        "Cancels",
        "% Cancels DTP",
        "percent",
        False
    )
]


headers = [
    "KPI",
    "Partner",
    PARTNER_SBU,
    "Gap"
]

x_positions = [
    0.9,
    4.0,
    6.4,
    8.9
]

widths = [
    2.5,
    1.8,
    1.8,
    3.2
]

add_table_header(
    slide,
    headers,
    x_positions,
    widths,
    1.42,
    height=0.42,
    font_size=14
)

y = 2.18

for (
    metric_name,
    metric_column,
    metric_format,
    higher_is_better
) in comparison_metrics:

    partner_value = (
        partner_comparison.get(
            metric_column,
            np.nan
        )
    )

    sbu_value = (
        sbu_comparison.get(
            metric_column,
            np.nan
        )
    )

    if metric_format == "currency":
        partner_display = (
            format_currency_2(
                partner_value
            )
        )

        sbu_display = (
            format_currency_2(
                sbu_value
            )
        )

        gap_info = benchmark_gap_info(
            partner_value,
            sbu_value,
            metric_type="relative",
            higher_is_better=
                higher_is_better,
            label=PARTNER_SBU
        )

    else:
        partner_display = (
            format_percent(
                partner_value
            )
        )

        sbu_display = (
            format_percent(
                sbu_value
            )
        )

        gap_info = benchmark_gap_info(
            partner_value,
            sbu_value,
            metric_type="pp",
            higher_is_better=
                higher_is_better,
            label=PARTNER_SBU
        )

    add_text(
        slide,
        metric_name,
        x_positions[0],
        y,
        widths[0],
        0.42,
        size=13,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE
    )

    add_text(
        slide,
        partner_display,
        x_positions[1],
        y,
        widths[1],
        0.42,
        size=13,
        valign=MSO_ANCHOR.MIDDLE
    )

    add_text(
        slide,
        sbu_display,
        x_positions[2],
        y,
        widths[2],
        0.42,
        size=13,
        valign=MSO_ANCHOR.MIDDLE
    )

    add_text(
        slide,
        gap_info["text"],
        x_positions[3],
        y,
        widths[3],
        0.42,
        size=12,
        bold=True,
        color=gap_info["color"],
        valign=MSO_ANCHOR.MIDDLE
    )

    y += 0.75


add_takeaway(
    slide,
    (
        f"Benchmark comparison uses the latest "
        f"common period available in both datasets: "
        f"{comparison_period.strftime('%B %Y')}."
    )
)


# ============================================================
# 32. GROWTH VS SBU
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    f"Growth vs {PARTNER_SBU}",
    (
        f"MoM growth comparison — "
        f"{comparison_period.strftime('%B %Y')}"
    )
)

add_footer(slide)


growth_metrics = [
    (
        "GMV Growth",
        "€ GMV"
    ),
    (
        "Orders Growth",
        "# Delivered Orders"
    ),
    (
        "AOV Growth",
        "€ AOV Derived"
    )
]

headers = [
    "KPI",
    "Partner",
    PARTNER_SBU,
    "Growth Gap"
]

x_positions = [
    0.9,
    4.0,
    6.4,
    8.9
]

widths = [
    2.5,
    1.8,
    1.8,
    3.2
]

add_table_header(
    slide,
    headers,
    x_positions,
    widths,
    1.42,
    height=0.42,
    font_size=14
)

y = 2.2

for (
    metric_name,
    column
) in growth_metrics:

    partner_growth = None
    sbu_growth = None

    if partner_comp_previous is not None:
        partner_growth = relative_change(
            partner_comparison[column],
            partner_comp_previous[column]
        )

    if sbu_comp_previous is not None:
        sbu_growth = relative_change(
            sbu_comparison[column],
            sbu_comp_previous[column]
        )

    if (
        partner_growth is not None
        and sbu_growth is not None
    ):
        growth_gap = (
            partner_growth
            - sbu_growth
        )

        gap_color = (
            GREEN
            if growth_gap >= 0
            else RED
        )

        gap_text = (
            f"{growth_gap * 100:+.1f} pp"
        )

    else:
        gap_text = ""
        gap_color = GREY

    partner_text = (
        f"{partner_growth:+.1%}"
        if partner_growth is not None
        else ""
    )

    sbu_text = (
        f"{sbu_growth:+.1%}"
        if sbu_growth is not None
        else ""
    )

    add_text(
        slide,
        metric_name,
        x_positions[0],
        y,
        widths[0],
        0.42,
        size=13,
        bold=True
    )

    add_text(
        slide,
        partner_text,
        x_positions[1],
        y,
        widths[1],
        0.42,
        size=13
    )

    add_text(
        slide,
        sbu_text,
        x_positions[2],
        y,
        widths[2],
        0.42,
        size=13
    )

    add_text(
        slide,
        gap_text,
        x_positions[3],
        y,
        widths[3],
        0.42,
        size=12,
        bold=True,
        color=gap_color
    )

    y += 0.85


# ============================================================
# 33. PARTNER SHARE OF SBU
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    f"Partner Share of {PARTNER_SBU}",
    (
        f"Partner GMV as a share of "
        f"{PARTNER_SBU} GMV — "
        f"{share_period.strftime('%B %Y')}"
    )
)

add_footer(slide)


share_mom = delta_info(
    pp_change(
        latest_share,
        previous_share
    ),
    metric_type="pp",
    higher_is_better=True
)

share_yoy_delta = delta_info(
    pp_change(
        latest_share,
        share_yoy
    ),
    metric_type="pp",
    higher_is_better=True,
    comparison_label="YoY"
)

add_kpi_card(
    slide,
    0.8,
    1.6,
    3.6,
    2.0,
    f"Share of {PARTNER_SBU} GMV",
    format_percent(latest_share),
    share_mom["text"],
    share_mom["color"],
    share_yoy_delta["text"],
    share_yoy_delta["color"]
)


share_chart_data = (
    share_data
    .tail(TREND_MONTHS)
    .copy()
)

share_labels = (
    share_chart_data[
        "Period"
    ]
    .dt.strftime("%b %y")
)

fig, ax = plt.subplots(
    figsize=(7.2, 3.7)
)

x_values = np.arange(
    len(share_chart_data)
)

share_values = (
    share_chart_data[
        "Partner Share"
    ] * 100
)

ax.plot(
    x_values,
    share_values,
    marker="o",
    linewidth=2.5,
    color=MPL_GREEN,
    label="Partner Share"
)

ax.fill_between(
    x_values,
    share_values,
    alpha=0.08,
    color=MPL_GREEN
)

ax.set_xticks(x_values)

ax.set_xticklabels(
    share_labels,
    rotation=35
)

ax.set_ylabel(
    f"Share of {PARTNER_SBU} GMV (%)"
)

ax.spines["top"].set_visible(False)
ax.spines["right"].set_visible(False)

ax.legend(
    frameon=False,
    loc="upper center"
)

plt.title(
    "Partner Share Evolution",
    loc="left",
    fontsize=13,
    fontweight="bold"
)

plt.tight_layout()

chart_file = "partner_sbu_share.png"

save_chart(
    fig,
    chart_file
)

add_chart_image(
    slide,
    chart_file,
    4.8,
    1.45,
    7.6,
    4.6
)

if previous_share is not None:

    share_change_pp = (
        latest_share
        - previous_share
    ) * 100

    if share_change_pp > 0:
        share_takeaway = (
            f"Partner gained "
            f"{abs(share_change_pp):.2f} pp "
            f"of {PARTNER_SBU} GMV share "
            f"vs previous month."
        )

    elif share_change_pp < 0:
        share_takeaway = (
            f"Partner lost "
            f"{abs(share_change_pp):.2f} pp "
            f"of {PARTNER_SBU} GMV share "
            f"vs previous month."
        )

    else:
        share_takeaway = (
            f"Partner share of "
            f"{PARTNER_SBU} GMV "
            f"was unchanged vs previous month."
        )

    add_takeaway(
        slide,
        share_takeaway
    )


# ============================================================
# 34. EXECUTIVE SUMMARY
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    "Executive Summary",
    (
        f"Latest partner performance — "
        f"{partner_latest_period.strftime('%B %Y')}"
    )
)

add_footer(slide)


latest_gmv = latest_partner["€ GMV"]

latest_orders = latest_partner[
    "# Delivered Orders"
]

latest_aov = latest_partner[
    "€ AOV Derived"
]

latest_cvr = latest_partner.get(
    "% CVR",
    np.nan
)

previous_gmv = get_row_value(
    partner_previous,
    "€ GMV"
)

previous_orders = get_row_value(
    partner_previous,
    "# Delivered Orders"
)

previous_aov = get_row_value(
    partner_previous,
    "€ AOV Derived"
)

previous_cvr = get_row_value(
    partner_previous,
    "% CVR"
)


gmv_delta = delta_info(
    relative_change(
        latest_gmv,
        previous_gmv
    )
)

orders_delta = delta_info(
    relative_change(
        latest_orders,
        previous_orders
    )
)

aov_delta = delta_info(
    relative_change(
        latest_aov,
        previous_aov
    )
)

cvr_delta = delta_info(
    pp_change(
        latest_cvr,
        previous_cvr
    ),
    metric_type="pp"
)

aov_gap = benchmark_gap_info(
    latest_aov,
    SBU_AOV,
    metric_type="relative",
    higher_is_better=True,
    label=PARTNER_SBU
)

cvr_gap = benchmark_gap_info(
    latest_cvr,
    SBU_CVR,
    metric_type="pp",
    higher_is_better=True,
    label=PARTNER_SBU
)


add_kpi_card(
    slide,
    0.45,
    1.45,
    2.85,
    1.75,
    "GMV",
    format_currency(latest_gmv),
    gmv_delta["text"],
    gmv_delta["color"]
)

add_kpi_card(
    slide,
    3.5,
    1.45,
    2.85,
    1.75,
    "Orders",
    format_number(latest_orders),
    orders_delta["text"],
    orders_delta["color"]
)

add_kpi_card(
    slide,
    6.55,
    1.45,
    2.85,
    1.75,
    "AOV",
    format_currency_2(latest_aov),
    aov_delta["text"],
    aov_delta["color"],
    aov_gap["text"],
    aov_gap["color"]
)

add_kpi_card(
    slide,
    9.6,
    1.45,
    2.85,
    1.75,
    "Avg Store CVR",
    format_percent(latest_cvr),
    cvr_delta["text"],
    cvr_delta["color"],
    cvr_gap["text"],
    cvr_gap["color"]
)


insights = []

gmv_change = relative_change(
    latest_gmv,
    previous_gmv
)

orders_change = relative_change(
    latest_orders,
    previous_orders
)

aov_change = relative_change(
    latest_aov,
    previous_aov
)

if gmv_change is not None:
    insights.append(
        f"GMV changed "
        f"{gmv_change:+.1%} "
        f"vs previous month."
    )

if (
    orders_change is not None
    and aov_change is not None
):
    if (
        abs(orders_change)
        > abs(aov_change)
    ):
        insights.append(
            "Order volume moved more sharply "
            "than basket size during the month."
        )

    else:
        insights.append(
            "Basket-size movement was larger "
            "than the change in order volume."
        )

if (
    pd.notna(latest_cvr)
    and SBU_CVR is not None
):
    cvr_gap_pp = (
        latest_cvr
        - SBU_CVR
    ) * 100

    if cvr_gap_pp >= 0:
        insights.append(
            f"Average store CVR is "
            f"{abs(cvr_gap_pp):.1f} pp "
            f"better than the latest available "
            f"{PARTNER_SBU} benchmark."
        )

    else:
        insights.append(
            f"Average store CVR is "
            f"{abs(cvr_gap_pp):.1f} pp "
            f"worse than the latest available "
            f"{PARTNER_SBU} benchmark."
        )

y = 3.75

for insight in insights[:3]:
    add_text(
        slide,
        "• " + insight,
        0.8,
        y,
        11.7,
        0.45,
        size=11.5
    )

    y += 0.62


# ============================================================
# 35. PARTNER PERFORMANCE EVOLUTION
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    "Partner Performance Evolution",
    "Monthly GMV and delivered orders"
)

add_footer(slide)

partner_trend = (
    monthly
    .tail(TREND_MONTHS)
    .copy()
)

partner_labels = (
    partner_trend[
        "Period"
    ]
    .dt.strftime("%b %y")
)

fig, ax1 = plt.subplots(
    figsize=(10.5, 4.5)
)

ax1.bar(
    partner_labels,
    partner_trend["€ GMV"],
    color=MPL_YELLOW,
    alpha=0.85,
    label="GMV (€)"
)

ax1.set_ylabel("GMV (€)")
ax1.spines["top"].set_visible(False)

ax2 = ax1.twinx()

ax2.plot(
    partner_labels,
    partner_trend[
        "# Delivered Orders"
    ],
    marker="o",
    linewidth=2.5,
    color=MPL_GREEN,
    label="Orders"
)

ax2.set_ylabel("Orders")
ax2.spines["top"].set_visible(False)

plt.xticks(rotation=35)

plt.title(
    "GMV & Orders Trend",
    loc="left",
    fontsize=14,
    fontweight="bold"
)

handles1, labels1 = ax1.get_legend_handles_labels()
handles2, labels2 = ax2.get_legend_handles_labels()

ax1.legend(
    handles1 + handles2,
    labels1 + labels2,
    loc="upper center",
    bbox_to_anchor=(0.5, 1.02),
    ncol=2,
    frameon=False,
    fontsize=9
)

plt.tight_layout()

chart_file = (
    "partner_performance_trend.png"
)

save_chart(
    fig,
    chart_file
)

add_chart_image(
    slide,
    chart_file,
    0.8,
    1.35,
    11.8,
    4.9
)


# ============================================================
# 36. GROWTH DRIVERS
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    "Growth Drivers",
    "Order volume and average basket evolution"
)

add_footer(slide)

fig, ax1 = plt.subplots(
    figsize=(10.5, 4.5)
)

ax1.plot(
    partner_labels,
    partner_trend[
        "# Delivered Orders"
    ],
    marker="o",
    linewidth=2.5,
    color=MPL_GREEN,
    label="Orders"
)

ax1.set_ylabel("Orders")
ax1.spines["top"].set_visible(False)

ax2 = ax1.twinx()

ax2.plot(
    partner_labels,
    partner_trend[
        "€ AOV Derived"
    ],
    marker="o",
    linewidth=2.5,
    color=MPL_YELLOW,
    label="AOV (€)"
)

ax2.set_ylabel("AOV (€)")
ax2.spines["top"].set_visible(False)

plt.xticks(rotation=35)

plt.title(
    "Orders vs AOV",
    loc="left",
    fontsize=14,
    fontweight="bold"
)

handles1, labels1 = ax1.get_legend_handles_labels()
handles2, labels2 = ax2.get_legend_handles_labels()

ax1.legend(
    handles1 + handles2,
    labels1 + labels2,
    loc="upper center",
    bbox_to_anchor=(0.5, 1.02),
    ncol=2,
    frameon=False,
    fontsize=9
)

plt.tight_layout()

chart_file = "growth_drivers.png"

save_chart(
    fig,
    chart_file
)

add_chart_image(
    slide,
    chart_file,
    0.8,
    1.35,
    11.8,
    4.9
)


# ============================================================
# 37. CONVERSION PERFORMANCE
# ============================================================

if "% CVR" in monthly.columns:

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        "Conversion Performance",
        (
            f"Partner average store CVR vs "
            f"{PARTNER_SBU}"
        )
    )

    add_footer(slide)

    merged_cvr = pd.merge(
        monthly[
            ["Period", "% CVR"]
        ],
        selected_sbu[
            ["Period", "% CVR"]
        ],
        on="Period",
        how="inner",
        suffixes=(
            "_Partner",
            "_SBU"
        )
    )

    merged_cvr = (
        merged_cvr
        .tail(TREND_MONTHS)
        .copy()
    )

    merged_cvr.loc[
        merged_cvr[
            "% CVR_SBU"
        ] <= 0,
        "% CVR_SBU"
    ] = np.nan

    fig, ax = plt.subplots(
        figsize=(10.5, 4.5)
    )

    cvr_labels = (
        merged_cvr[
            "Period"
        ]
        .dt.strftime("%b %y")
    )

    ax.plot(
        cvr_labels,
        merged_cvr[
            "% CVR_Partner"
        ] * 100,
        marker="o",
        linewidth=2.5,
        color=MPL_GREEN,
        label="Partner Avg Store CVR"
    )

    ax.plot(
        cvr_labels,
        merged_cvr[
            "% CVR_SBU"
        ] * 100,
        marker="o",
        linewidth=2.2,
        linestyle="--",
        color=MPL_YELLOW,
        label=f"{PARTNER_SBU} CVR"
    )

    ax.set_ylabel("CVR (%)")

    ax.spines[
        "top"
    ].set_visible(False)

    ax.spines[
        "right"
    ].set_visible(False)

    ax.legend(
        frameon=False,
        loc="upper center",
        ncol=2
    )

    plt.xticks(rotation=35)

    plt.title(
        "CVR Trend",
        loc="left",
        fontsize=14,
        fontweight="bold"
    )

    plt.tight_layout()

    chart_file = (
        "conversion_performance.png"
    )

    save_chart(
        fig,
        chart_file
    )

    add_chart_image(
        slide,
        chart_file,
        0.8,
        1.35,
        11.8,
        4.9
    )


# ============================================================
# 38. OPERATIONAL HEALTH
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    "Operational Health",
    (
        f"Latest partner performance vs "
        f"{PARTNER_SBU} benchmark "
        f"({sbu_latest_period.strftime('%b %Y')})"
    )
)

add_footer(slide)

operational_cards = []

if "% Uptime" in monthly.columns:

    current = latest_partner[
        "% Uptime"
    ]

    previous_value = get_row_value(
        partner_previous,
        "% Uptime"
    )

    mom = delta_info(
        pp_change(
            current,
            previous_value
        ),
        metric_type="pp",
        higher_is_better=True
    )

    gap = benchmark_gap_info(
        current,
        SBU_UPTIME,
        metric_type="pp",
        higher_is_better=True,
        label=PARTNER_SBU
    )

    operational_cards.append(
        (
            "Avg Store Uptime",
            current,
            mom,
            gap
        )
    )


if "% PNA Orders" in monthly.columns:

    current = latest_partner[
        "% PNA Orders"
    ]

    previous_value = get_row_value(
        partner_previous,
        "% PNA Orders"
    )

    mom = delta_info(
        pp_change(
            current,
            previous_value
        ),
        metric_type="pp",
        higher_is_better=False
    )

    gap = benchmark_gap_info(
        current,
        SBU_PNA,
        metric_type="pp",
        higher_is_better=False,
        label=PARTNER_SBU
    )

    operational_cards.append(
        (
            "Avg Store PNA",
            current,
            mom,
            gap
        )
    )


if "% Cancels DTP" in monthly.columns:

    current = latest_partner[
        "% Cancels DTP"
    ]

    previous_value = get_row_value(
        partner_previous,
        "% Cancels DTP"
    )

    mom = delta_info(
        pp_change(
            current,
            previous_value
        ),
        metric_type="pp",
        higher_is_better=False
    )

    gap = benchmark_gap_info(
        current,
        SBU_CANCELS,
        metric_type="pp",
        higher_is_better=False,
        label=PARTNER_SBU
    )

    operational_cards.append(
        (
            "Avg Store Cancels",
            current,
            mom,
            gap
        )
    )


op_x = [
    0.75,
    4.8,
    8.85
]

for i, (
    title,
    value,
    mom,
    gap
) in enumerate(
    operational_cards[:3]
):

    add_kpi_card(
        slide,
        op_x[i],
        1.5,
        3.7,
        1.75,
        title,
        format_percent(value),
        mom["text"],
        mom["color"],
        gap["text"],
        gap["color"]
    )


add_text(
    slide,
    "Interpretation",
    0.85,
    3.8,
    3.5,
    0.35,
    size=12,
    bold=True
)

add_text(
    slide,
    (
        "Higher is better for Uptime. "
        "Lower is better for PNA and Cancels. "
        "Green indicates improvement; "
        "red indicates deterioration."
    ),
    0.85,
    4.2,
    11.2,
    0.5,
    size=10,
    color=GREY
)


# ============================================================
# 39. STORE PERFORMANCE MATRIX
# ============================================================

matrix_data = latest_stores[
    latest_stores[
        "# Delivered Orders"
    ]
    >= MIN_ORDERS_FOR_CVR_CLASSIFICATION
].copy()

matrix_data = matrix_data.dropna(
    subset=[
        "% CVR",
        "€ GMV"
    ]
)


if len(matrix_data) > 0:

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        "Store Performance Matrix",
        (
            f"GMV vs CVR — "
            f"{partner_latest_period.strftime('%B %Y')} | "
            f"Minimum "
            f"{MIN_ORDERS_FOR_CVR_CLASSIFICATION} "
            f"orders"
        )
    )

    add_footer(slide)

    matrix_data = (
        matrix_data
        .sort_values(
            "€ GMV",
            ascending=False
        )
        .reset_index(drop=True)
    )

    STORE_COLOR_PALETTE = [
        "#1F77B4",
        "#FF7F0E",
        "#2CA02C",
        "#D62728",
        "#9467BD",
        "#8C564B",
        "#E377C2",
        "#7F7F7F",
        "#BCBD22",
        "#17BECF",
        "#393B79",
        "#637939",
        "#8C6D31",
        "#843C39",
        "#7B4173",
        "#3182BD",
        "#31A354",
        "#756BB1",
        "#636363",
        "#E6550D",
        "#6BAED6",
        "#74C476",
        "#9E9AC8",
        "#969696",
        "#FD8D3C",
        "#9ECAE1",
        "#A1D99B",
        "#BCBDDC",
        "#BDBDBD",
        "#FDAE6B",
    ]

    number_of_stores = len(
        matrix_data
    )

    store_colors = [
        STORE_COLOR_PALETTE[
            i % len(STORE_COLOR_PALETTE)
        ]
        for i in range(
            number_of_stores
        )
    ]

    orders = (
        matrix_data[
            "# Delivered Orders"
        ]
        .fillna(0)
    )

    min_orders = orders.min()
    max_orders = orders.max()

    if max_orders > min_orders:

        bubble_sizes = (
            120
            +
            360
            *
            (
                (
                    orders - min_orders
                )
                /
                (
                    max_orders - min_orders
                )
            )
        )

    else:
        bubble_sizes = np.repeat(
            240,
            number_of_stores
        )

    cvr_values = (
        matrix_data[
            "% CVR"
        ] * 100
    )

    if SBU_CVR is not None:

        benchmark_cvr = (
            SBU_CVR * 100
        )

        x_min = min(
            cvr_values.min(),
            benchmark_cvr
        ) - 4

        x_max = max(
            cvr_values.max(),
            benchmark_cvr
        ) + 4

    else:

        benchmark_cvr = None

        x_min = (
            cvr_values.min() - 4
        )

        x_max = (
            cvr_values.max() + 4
        )

    if x_max <= x_min:
        x_max = x_min + 10

    y_max = (
        matrix_data[
            "€ GMV"
        ].max() * 1.18
    )

    if (
        pd.isna(y_max)
        or y_max <= 0
    ):
        y_max = 100

    fig, ax = plt.subplots(
        figsize=(9.4, 4.7)
    )

    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    if (
        benchmark_cvr is not None
        and partner_median_gmv is not None
    ):

        ax.fill_between(
            [
                benchmark_cvr,
                x_max
            ],
            partner_median_gmv,
            y_max,
            color="#EAF7F4",
            alpha=0.55
        )

        ax.fill_between(
            [
                x_min,
                benchmark_cvr
            ],
            partner_median_gmv,
            y_max,
            color="#FFF4E5",
            alpha=0.55
        )

        ax.fill_between(
            [
                benchmark_cvr,
                x_max
            ],
            0,
            partner_median_gmv,
            color="#FFF9E8",
            alpha=0.55
        )

        ax.fill_between(
            [
                x_min,
                benchmark_cvr
            ],
            0,
            partner_median_gmv,
            color="#FCEEEE",
            alpha=0.45
        )

    for i, row in matrix_data.iterrows():

        ax.scatter(
            row["% CVR"] * 100,
            row["€ GMV"],
            s=bubble_sizes.iloc[i],
            color=store_colors[i],
            alpha=0.88,
            edgecolors="white",
            linewidth=1.2,
            zorder=5
        )

    if benchmark_cvr is not None:

        ax.axvline(
            benchmark_cvr,
            linestyle="--",
            linewidth=1.6,
            color=MPL_YELLOW,
            zorder=3
        )

        ax.text(
            benchmark_cvr + 0.35,
            y_max * 0.58,
            f"{PARTNER_SBU} CVR {benchmark_cvr:.1f}%",
            rotation=90,
            va="center",
            ha="left",
            fontsize=7.5,
            fontweight="bold",
            color=MPL_ORANGE
        )

    if partner_median_gmv is not None:

        ax.axhline(
            partner_median_gmv,
            linestyle="--",
            linewidth=1.2,
            color=MPL_GREY,
            zorder=3
        )

        ax.text(
            x_max - 0.5,
            partner_median_gmv + y_max * 0.018,
            f"Partner Median GMV €{partner_median_gmv:,.0f}",
            ha="right",
            va="bottom",
            fontsize=7.5,
            fontweight="bold",
            color=MPL_GREY
        )

    if (
        benchmark_cvr is not None
        and partner_median_gmv is not None
    ):

        ax.text(
            0.98,
            0.96,
            "TOP PERFORMERS",
            transform=ax.transAxes,
            ha="right",
            va="top",
            fontsize=8,
            fontweight="bold",
            color=MPL_GREEN
        )

        ax.text(
            0.02,
            0.96,
            "CONVERSION\nOPPORTUNITY",
            transform=ax.transAxes,
            ha="left",
            va="top",
            fontsize=7.5,
            fontweight="bold",
            color=MPL_ORANGE
        )

        ax.text(
            0.98,
            0.05,
            "GROWTH POTENTIAL",
            transform=ax.transAxes,
            ha="right",
            va="bottom",
            fontsize=7.5,
            fontweight="bold",
            color=MPL_DARK
        )

        ax.text(
            0.02,
            0.05,
            "NEEDS ATTENTION",
            transform=ax.transAxes,
            ha="left",
            va="bottom",
            fontsize=7.5,
            fontweight="bold",
            color=MPL_RED
        )

    ax.set_xlabel("CVR (%)")
    ax.set_ylabel("GMV (€)")

    ax.set_xlim(
        x_min,
        x_max
    )

    ax.set_ylim(
        0,
        y_max
    )

    ax.grid(
        axis="both",
        alpha=0.12
    )

    ax.spines[
        "top"
    ].set_visible(False)

    ax.spines[
        "right"
    ].set_visible(False)

    ax.tick_params(
        labelsize=8
    )

    ax.set_title(
        "Store Commercial Positioning",
        loc="left",
        fontsize=12,
        fontweight="bold",
        pad=10
    )

    plt.tight_layout()

    chart_file = (
        "store_matrix_v8.png"
    )

    fig.savefig(
        chart_file,
        dpi=200,
        facecolor="white",
        bbox_inches="tight"
    )

    plt.close(fig)

    slide.shapes.add_picture(
        chart_file,
        Inches(0.45),
        Inches(1.35),
        width=Inches(9.05),
        height=Inches(5.45)
    )

    add_text(
        slide,
        "Stores",
        9.68,
        1.38,
        3.0,
        0.35,
        size=10,
        bold=True
    )

    add_text(
        slide,
        "Bubble size = delivered orders",
        9.68,
        1.72,
        3.0,
        0.28,
        size=7.5,
        color=GREY
    )

    legend_start_y = 2.08
    legend_row_height = 0.29
    max_rows_per_column = 16

    for i, row in matrix_data.iterrows():

        column_number = (
            i // max_rows_per_column
        )

        row_number = (
            i % max_rows_per_column
        )

        legend_x = (
            9.68
            + column_number * 1.45
        )

        legend_y = (
            legend_start_y
            + row_number
            * legend_row_height
        )

        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(legend_x),
            Inches(
                legend_y + 0.03
            ),
            Inches(0.11),
            Inches(0.11)
        )

        hex_color = (
            store_colors[i]
            .replace("#", "")
        )

        r = int(
            hex_color[0:2],
            16
        )

        g = int(
            hex_color[2:4],
            16
        )

        b = int(
            hex_color[4:6],
            16
        )

        dot.fill.solid()

        dot.fill.fore_color.rgb = (
            RGBColor(r, g, b)
        )

        dot.line.fill.background()

        store_name = str(
            row[
                "Store Address Address"
            ]
        )

        if len(store_name) > 31:
            store_name = (
                store_name[:29]
                + "…"
            )

        add_text(
            slide,
            store_name,
            legend_x + 0.17,
            legend_y,
            1.25,
            0.22,
            size=6.5,
            color=DARK
        )


# ============================================================
# 40. TOP PERFORMING STORES
# ============================================================

top_stores = (
    latest_stores[
        latest_stores[
            "# Delivered Orders"
        ] > 0
    ]
    .sort_values(
        "€ GMV",
        ascending=False
    )
    .head(5)
)

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    "Top Performing Stores",
    (
        f"Highest GMV stores — "
        f"{partner_latest_period.strftime('%B %Y')}"
    )
)

add_footer(slide)

headers = [
    "Store",
    "GMV",
    "Orders",
    "CVR",
    "Commercial",
    "Operations"
]

x_positions = [
    0.6,
    4.45,
    5.65,
    6.75,
    7.8,
    10.55
]

widths = [
    3.65,
    1.05,
    1.0,
    0.9,
    2.45,
    2.0
]

add_table_header(
    slide,
    headers,
    x_positions,
    widths,
    1.42,
    height=0.42,
    font_size=12
)

y = 2.18

for _, row in top_stores.iterrows():

    store_name = str(
        row[
            "Store Address Address"
        ]
    )

    if len(store_name) > 38:
        store_name = (
            store_name[:35]
            + "..."
        )

    values = [
        store_name,
        format_currency(
            row["€ GMV"]
        ),
        format_number(
            row[
                "# Delivered Orders"
            ]
        ),
        format_percent(
            row.get("% CVR")
        ),
        row[
            "Commercial Status"
        ],
        row[
            "Operational Status"
        ]
    ]

    for i, value in enumerate(values):

        add_text(
            slide,
            value,
            x_positions[i],
            y,
            widths[i],
            0.52,
            size=10.5,
            bold=(i == 0)
        )

    y += 0.76


# ============================================================
# 41. BIGGEST OPPORTUNITIES
# ============================================================

opportunities = (
    latest_stores[
        (latest_stores["# Delivered Orders"] > 0)
        &
        (latest_stores["Opportunity Score"] > 0)
    ]
    .copy()
    .sort_values(
        [
            "Opportunity Score",
            "€ GMV"
        ],
        ascending=[
            False,
            False
        ]
    )
    .head(TOP_OPPORTUNITIES_COUNT)
)

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    "Biggest Opportunities",
    (
        "Prioritized using commercial impact, "
        "CVR gap and operational gaps"
    )
)

add_footer(slide)

headers = [
    "Store",
    "Score",
    "Commercial",
    "Operations",
    "Main KPI Gap",
    "Confidence"
]

x_positions = [
    0.55,
    3.8,
    4.7,
    7.05,
    9.05,
    11.55
]

widths = [
    3.05,
    0.75,
    2.15,
    1.8,
    2.25,
    1.15
]

add_table_header(
    slide,
    headers,
    x_positions,
    widths,
    1.42,
    height=0.42,
    font_size=11
)

y = 2.15

for _, row in opportunities.iterrows():

    name = str(
        row[
            "Store Address Address"
        ]
    )

    if len(name) > 32:
        name = (
            name[:29]
            + "..."
        )

    values = [
        name,
        format_number(
            row[
                "Opportunity Score"
            ]
        ),
        row[
            "Commercial Status"
        ],
        row[
            "Operational Status"
        ],
        row[
            "Main Issue"
        ],
        row[
            "Confidence"
        ]
    ]

    for i, value in enumerate(values):

        add_text(
            slide,
            value,
            x_positions[i],
            y,
            widths[i],
            0.52,
            size=9.5,
            bold=(
                i == 0
                or i == 1
            )
        )

    y += 0.72


# ============================================================
# 42. RECOMMENDATIONS
# ============================================================

recommendation_rows = (
    opportunities.head(4)
)

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    "Recommendations",
    (
        "Store-specific priorities based "
        "on the opportunity ranking"
    )
)

add_footer(slide)

y = 1.35

for _, row in recommendation_rows.iterrows():

    store = str(
        row[
            "Store Address Address"
        ]
    )

    if len(store) > 34:
        store = (
            store[:31]
            + "..."
        )

    focus = row[
        "Recommended Focus"
    ]

    issue = row[
        "Main Issue"
    ]

    score = row[
        "Opportunity Score"
    ]

    recommendation_text = (
        f"{issue}. "
        f"Recommended focus: "
        f"{focus}."
    )

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.75),
        Inches(y),
        Inches(11.85),
        Inches(1.05)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GREY
    shape.line.color.rgb = LIGHT_GREY

    add_text(
        slide,
        store,
        1.0,
        y + 0.13,
        3.2,
        0.3,
        size=10,
        bold=True
    )

    add_text(
        slide,
        f"Score {score:.0f}",
        4.05,
        y + 0.13,
        1.0,
        0.3,
        size=9,
        bold=True,
        color=ORANGE
    )

    add_text(
        slide,
        recommendation_text,
        5.05,
        y + 0.12,
        7.1,
        0.65,
        size=9
    )

    y += 1.22


# ============================================================
# 43. ACTION PLAN
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(slide)

add_title(
    slide,
    "Action Plan",
    (
        "Proposed actions — owner and deadline "
        "to be agreed during the meeting"
    )
)

add_footer(slide)

headers = [
    "Recommended Action",
    "Owner",
    "Deadline",
    "Expected Impact"
]

x_positions = [
    0.75,
    6.2,
    8.1,
    9.9
]

widths = [
    5.1,
    1.5,
    1.5,
    2.7
]

add_table_header(
    slide,
    headers,
    x_positions,
    widths,
    1.42,
    height=0.42,
    font_size=11
)

action_candidates = []

if (
    "% PNA Orders" in monthly.columns
    and SBU_PNA is not None
    and latest_partner[
        "% PNA Orders"
    ] > SBU_PNA
):
    action_candidates.append(
        (
            "Review availability and PNA in priority stores",
            "",
            "",
            "Lower PNA / improve fulfilled demand"
        )
    )

if (
    "% Cancels DTP" in monthly.columns
    and SBU_CANCELS is not None
    and latest_partner[
        "% Cancels DTP"
    ] > SBU_CANCELS
):
    action_candidates.append(
        (
            "Review cancellation drivers in priority stores",
            "",
            "",
            "Reduce avoidable cancellations"
        )
    )

if (
    pd.notna(latest_cvr)
    and SBU_CVR is not None
    and latest_cvr
    < SBU_CVR - CVR_TOLERANCE_PP
):
    action_candidates.append(
        (
            "Review conversion gaps in high-volume stores",
            "",
            "",
            "Improve conversion"
        )
    )

if (
    "% Uptime" in monthly.columns
    and SBU_UPTIME is not None
    and latest_partner[
        "% Uptime"
    ] < SBU_UPTIME
):
    action_candidates.append(
        (
            "Improve uptime in stores below benchmark",
            "",
            "",
            "Increase available selling hours"
        )
    )

if not action_candidates:
    action_candidates.append(
        (
            "Maintain current performance and track key KPIs",
            "",
            "",
            "Protect performance"
        )
    )

y = 2.2

for row in action_candidates[:4]:

    for i, value in enumerate(row):

        add_text(
            slide,
            value,
            x_positions[i],
            y,
            widths[i],
            0.65,
            size=9.5
        )

    y += 0.9


# ============================================================
# 44. Q&A
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(
    slide,
    YELLOW
)

add_text(
    slide,
    "Questions & Discussion",
    0.9,
    2.7,
    11.5,
    0.8,
    size=30,
    bold=True,
    align=PP_ALIGN.CENTER
)


# ============================================================
# 45. APPENDIX COVER
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(
    slide,
    DARK
)

add_text(
    slide,
    "Appendix",
    0.9,
    2.6,
    11.5,
    0.8,
    size=30,
    bold=True,
    color=WHITE,
    align=PP_ALIGN.CENTER
)

add_text(
    slide,
    "Store-level performance detail",
    0.9,
    3.45,
    11.5,
    0.5,
    size=14,
    color=WHITE,
    align=PP_ALIGN.CENTER
)


# ============================================================
# 46. STORE APPENDIX
# ============================================================

stores = (
    store_monthly[
        "Store Address Address"
    ]
    .dropna()
    .unique()
)

for store_name in stores:

    store_data = (
        store_monthly[
            store_monthly[
                "Store Address Address"
            ] == store_name
        ]
        .sort_values("Period")
    )

    if store_data.empty:
        continue

    latest_store = (
        store_data.iloc[-1]
    )

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    display_name = str(
        store_name
    )

    if len(display_name) > 70:
        display_name = (
            display_name[:67]
            + "..."
        )

    add_title(
        slide,
        display_name,
        "Store Performance Overview"
    )

    add_footer(slide)

    status_match = latest_stores[
        latest_stores[
            "Store Address Address"
        ] == store_name
    ]

    if not status_match.empty:

        status = status_match.iloc[0]

        commercial_status = (
            status[
                "Commercial Status"
            ]
        )

        operational_status = (
            status[
                "Operational Status"
            ]
        )

        confidence = (
            status["Confidence"]
        )

        strength = (
            status["Strength"]
        )

        risk = status["Risk"]
        focus = status["Focus"]

        opportunity_score = (
            status[
                "Opportunity Score"
            ]
        )

    else:
        commercial_status = ""
        operational_status = ""
        confidence = ""

        strength = ""
        risk = ""
        focus = ""

        opportunity_score = 0


    add_kpi_card(
        slide,
        0.5,
        1.3,
        2.2,
        1.55,
        "GMV",
        format_currency(
            latest_store[
                "€ GMV"
            ]
        )
    )

    add_kpi_card(
        slide,
        2.85,
        1.3,
        2.2,
        1.55,
        "Orders",
        format_number(
            latest_store[
                "# Delivered Orders"
            ]
        )
    )

    cvr_secondary = (
        f"{PARTNER_SBU}: "
        f"{format_percent(SBU_CVR)}"
        if SBU_CVR is not None
        else ""
    )

    add_kpi_card(
        slide,
        5.2,
        1.3,
        2.2,
        1.55,
        "CVR",
        format_percent(
            latest_store.get(
                "% CVR"
            )
        ),
        secondary_text=
            cvr_secondary
    )

    add_kpi_card(
        slide,
        7.55,
        1.3,
        2.45,
        1.55,
        "Commercial",
        commercial_status
    )

    add_kpi_card(
        slide,
        10.15,
        1.3,
        2.45,
        1.55,
        "Operations",
        operational_status
    )

    add_section_card(
        slide,
        0.65,
        3.05,
        3.7,
        1.05,
        "Strength",
        strength
    )

    add_section_card(
        slide,
        4.55,
        3.05,
        3.7,
        1.05,
        "Risk",
        risk
    )

    add_section_card(
        slide,
        8.45,
        3.05,
        3.7,
        1.05,
        "Recommended Focus",
        focus
    )

    score_text = (
        f"Opportunity Score: "
        f"{opportunity_score:.0f}/100"
    )

    if confidence:
        score_text += (
            f"   |   "
            f"Confidence: "
            f"{confidence}"
        )

    add_text(
        slide,
        score_text,
        0.8,
        4.28,
        5.8,
        0.35,
        size=9,
        bold=True,
        color=GREY
    )

    fig, ax1 = plt.subplots(
        figsize=(10.5, 2.6)
    )

    store_chart_data = (
        store_data
        .tail(TREND_MONTHS)
    )

    store_labels = (
        store_chart_data[
            "Period"
        ]
        .dt.strftime("%b %y")
    )

    ax1.bar(
        store_labels,
        store_chart_data[
            "€ GMV"
        ],
        color=MPL_YELLOW,
        alpha=0.85,
        label="GMV (€)"
    )

    ax1.set_ylabel("GMV (€)")

    ax1.spines[
        "top"
    ].set_visible(False)

    ax2 = ax1.twinx()

    ax2.plot(
        store_labels,
        store_chart_data[
            "# Delivered Orders"
        ],
        marker="o",
        linewidth=2,
        color=MPL_GREEN,
        label="Orders"
    )

    ax2.set_ylabel("Orders")

    ax2.spines[
        "top"
    ].set_visible(False)

    handles1, labels1 = ax1.get_legend_handles_labels()
    handles2, labels2 = ax2.get_legend_handles_labels()

    ax1.legend(
        handles1 + handles2,
        labels1 + labels2,
        loc="upper center",
        bbox_to_anchor=(0.5, 1.02),
        ncol=2,
        frameon=False,
        fontsize=8
    )

    plt.xticks(rotation=30)

    plt.tight_layout()

    safe_store = (
        str(store_name)
        .replace("/", "_")
        .replace("\\", "_")
        .replace(":", "_")
    )

    chart_file = (
        "store_"
        + safe_store[:35]
        + ".png"
    )

    save_chart(
        fig,
        chart_file
    )

    add_chart_image(
        slide,
        chart_file,
        0.8,
        4.65,
        11.8,
        2.15
    )


# ============================================================
# 47. THANK YOU
# ============================================================

slide = prs.slides.add_slide(
    prs.slide_layouts[6]
)

set_background(
    slide,
    YELLOW
)

if os.path.exists(LOGO_FILE):

    slide.shapes.add_picture(
        LOGO_FILE,
        Inches(5.45),
        Inches(1.35),
        width=Inches(2.4)
    )

add_text(
    slide,
    "Thank You",
    0.9,
    3.3,
    11.5,
    0.8,
    size=32,
    bold=True,
    align=PP_ALIGN.CENTER
)


# ============================================================
# 48. SAVE
# ============================================================

prs.save(
    OUTPUT_FILE
)

print()
print("============================================")
print("V8 PRESENTATION CREATED SUCCESSFULLY")
print("Output:", OUTPUT_FILE)
print("Font:", FONT)
print("SBU:", PARTNER_SBU)

print(
    "Partner latest:",
    partner_latest_period.strftime(
        "%B %Y"
    )
)

print(
    "SBU latest:",
    sbu_latest_period.strftime(
        "%B %Y"
    )
)

print(
    "Comparison period:",
    comparison_period.strftime(
        "%B %Y"
    )
)

print(
    f"Partner share of {PARTNER_SBU}:",
    format_percent(
        latest_share
    )
)

print("============================================")
