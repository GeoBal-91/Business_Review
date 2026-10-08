"""Dev6.1.1 display helpers. Analysis records are read, never modified."""
import colorsys
import textwrap
import matplotlib.pyplot as plt
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from data_pipeline import G, period_row, present
from diagnostics import box, BORDER, AMBER
from formatting import format_currency, format_percent
from ppt_core import add_text, add_title, add_footer, add_chart_image, save_chart, add_picture_contain
from theme import ASSETS_DIR, GREEN, GREY, WHITE

MATRIX_FOOTNOTE = 'Positioning uses GMV and conversion; opportunity ranking uses the assessed KPI gaps and exposure.'
COVERAGE_CVR_NOTE = 'Partner vs SBU CVR: Not assessed — incomplete store coverage.'
CREAM = RGBColor(255, 251, 237)
# Categorical identity only. The fixed palette is assigned to sorted stable IDs,
# independently of input order, GMV, CVR, status, severity and opportunity rank.
STORE_PALETTE = (
    '#00A082', '#4477AA', '#CC6677', '#AA4499', '#E69F00', '#332288',
    '#117733', '#88CCEE', '#882255', '#999933', '#661100', '#DDCC77',
    '#0077BB', '#EE7733', '#44AA99', '#994F88', '#669900', '#CC3311',
    '#777777', '#BB88CC',
)


def current_coverage_incomplete(result):
    row = period_row(result.data.coverage, result.data.reporting_period)
    return row is not None and row['Missing Stores'] > 0


def store_color_mapping(store_ids):
    mapping = {}
    for index, sid in enumerate(sorted({str(x) for x in store_ids})):
        if index < len(STORE_PALETTE):
            color = STORE_PALETTE[index]
        else:
            step = index - len(STORE_PALETTE)
            while True:
                rgb = colorsys.hls_to_rgb((step * .618033988749895) % 1, .43, .65)
                color = '#' + ''.join(f'{round(c * 255):02X}' for c in rgb)
                if color.lower() not in {v.lower() for v in mapping.values()}:
                    break
                step += 1
        mapping[sid] = color
    return mapping


def store_label(row):
    address = str(row['Store Address Address'])
    name = row.get('Store Name')
    return f'{name} — {address}' if isinstance(name, str) and name.strip() and name != address else address


def matrix_figure(active, benchmark, conversion_label='CVR'):
    colors = store_color_mapping(active['Store ID'])
    fig, ax = plt.subplots(figsize=(7.3, 4.7))
    ax.scatter(active[G], active['% CVR'] * 100, s=90,
               c=[colors[str(sid)] for sid in active['Store ID']])
    ax.axvline(active[G].median(), color='grey', linestyle='--', label='Median active-store GMV')
    if benchmark is not None and present(benchmark.get('% CVR')):
        ax.axhline(benchmark['% CVR'] * 100, color='#FFC244', linestyle='--', label='SBU conversion')
    ax.set_xlabel('GMV (€)'); ax.set_ylabel(conversion_label + ' (%)')
    ax.legend(loc='lower left', bbox_to_anchor=(0, 1.01), ncol=2, fontsize=8.5, frameon=False)
    ax.spines[['top', 'right']].set_visible(False)
    return fig


def matrix_legend_pages(active):
    """Keep full names at 10.5 pt; continue longer legends at the same size."""
    colors = store_color_mapping(active['Store ID'])
    pages, page, used = [], [], 0
    for _, row in active.sort_values('Store ID').iterrows():
        label = store_label(row)
        lines = sum(max(1, len(textwrap.wrap(line, width=43))) for line in label.splitlines())
        height = max(.34, .175 * lines + .025)
        if page and used + height > 4.54:
            pages.append(page); page, used = [], 0
        page.append((str(row['Store ID']), label, colors[str(row['Store ID'])], height))
        used += height
    if page:
        pages.append(page)
    return pages


def _page(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title, subtitle); add_footer(slide)
    return slide


def build_matrix(prs, result, active, charts):
    period = result.summary['reporting_period']
    subtitle = f'{period} · Commercial position is separate from operational status'
    if active.empty:
        slide = _page(prs, 'Store Performance Matrix', subtitle)
        add_text(slide, 'No complete active-store observations for the matrix.', .8, 2, 11, 1, size=18)
        return
    fig = matrix_figure(active, result.data.benchmark_row, result.settings.conversion_label)
    path = charts('store_matrix.png'); save_chart(fig, path)
    pages = matrix_legend_pages(active)
    shown = 0
    for index, entries in enumerate(pages):
        title = 'Store Performance Matrix' + (f' · {index + 1}/{len(pages)}' if len(pages) > 1 else '')
        slide = _page(prs, title, subtitle)
        add_chart_image(slide, path, .73, 1.55, 7.35, 4.98)
        box(slide, 8.35, 1.55, 4.32, 5.08, CREAM)
        legend_title = 'Store legend' if len(pages) == 1 else f'Store legend · {shown + 1}–{shown + len(entries)} of {len(active)}'
        add_text(slide, legend_title, 8.53, 1.72, 3.96, .3, size=12, bold=True)
        y = 2.12
        for sid, label, color, height in entries:
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.55), Inches(y + .035), Inches(.13), Inches(.13))
            dot.fill.solid(); dot.fill.fore_color.rgb = RGBColor.from_string(color[1:])
            dot.line.fill.background(); dot.name = 'Store color: ' + sid
            dot._element.spPr.append(OxmlElement('a:effectLst'))
            for effect in dot._element.xpath('./p:style/a:effectRef'):
                effect.set('idx', '0')
            name = add_text(slide, label, 8.82, y, 3.63, height - .005, size=10.5)
            name.name = 'Store legend: ' + sid
            y += height
        shown += len(entries)
        add_text(slide, MATRIX_FOOTNOTE, .8, 6.80, 11.7, .42, size=11, color=GREY)


def top_gmv_concentration(result, displayed):
    """A display-only share of current partner GMV; no ranking or eligibility."""
    count = len(displayed)
    current = period_row(result.data.monthly, result.data.reporting_period)
    total = current.get(G) if current is not None else None
    if current_coverage_incomplete(result):
        return count, None, 'GMV share not assessed — incomplete store coverage.'
    if not count:
        return count, None, 'No positive-GMV stores available for the current period.'
    numerator = displayed[G].sum(skipna=False)
    if not present(total) or not present(numerator):
        return count, None, 'GMV share not assessed — missing GMV evidence.'
    if total <= 0:
        return count, None, 'GMV share unavailable — partner GMV is zero.'
    return count, float(numerator / total), ''


def build_top_gmv(prs, result, displayed, title='Top GMV Stores'):
    slide = _page(prs, title, result.summary['reporting_period'])
    box(slide, .67, 1.55, 8.55, .49, CREAM)
    columns = [(1.10, 2.68), (3.93, 1.03), (5.12, 1.95), (7.24, 1.82)]
    for label, (x, width) in zip(['Store', 'GMV', 'Commercial status', 'Operational status'], columns):
        add_text(slide, label, x, 1.68, width, .28, size=10, bold=True, color=GREEN)
    for i, (_, row) in enumerate(displayed.iterrows()):
        y = 2.17 + i * .77
        box(slide, .67, y - .02, 8.55, .71, WHITE if i % 2 else RGBColor(247, 249, 249))
        add_text(slide, str(i + 1), .80, y + .12, .24, .25, size=11, bold=True, color=GREEN)
        values = [row['Store Address Address'], format_currency(row[G]), row['Commercial Status'], row['Operational Status']]
        for value, (x, width) in zip(values, columns):
            add_text(slide, value, x, y + .06, width, .62, size=10)
    if displayed.empty:
        add_text(slide, 'No positive-GMV stores available for the current period.', .85, 2.4, 8, .6, size=15)
    box(slide, 9.48, 1.55, 3.19, 5.33, AMBER)
    add_text(slide, 'Top Stores', 9.70, 1.77, 2.7, .32, size=15, bold=True)
    add_picture_contain(slide, ASSETS_DIR / 'goldCrownCelebration.png', 9.78, 2.15, 2.55, 1.87)
    count, share, reason = top_gmv_concentration(result, displayed)
    period = result.data.reporting_period.strftime('%B %Y')
    if share is not None:
        add_text(slide, format_percent(share), 9.70, 4.12, 2.7, .55, size=29, bold=True, color=GREEN)
        noun = 'store' if count == 1 else 'stores'
        add_text(slide, f'Top {count} {noun} generated {share:.1%} of total partner GMV in {period}.',
                 9.70, 4.98, 2.72, 1.08, size=13)
    else:
        noun = 'store' if count == 1 else 'stores'
        add_text(slide, f'{count} {noun} displayed', 9.70, 4.17, 2.72, .64, size=18, bold=True, color=GREEN)
        add_text(slide, reason, 9.70, 5.04, 2.72, 1.08, size=12)
    current = period_row(result.data.monthly, result.data.reporting_period)
    total = current.get(G) if current is not None else None
    add_text(slide, 'Reported partner GMV: ' + format_currency(total), 9.70, 6.36, 2.72, .32, size=9, color=GREY)


def guide_cards(result):
    s = result.settings
    priority = ('Each eligible KPI’s own severity and exposure, with dominant-KPI protection.'
                if s.priority_model == 'paired' else 'Challenger: Priority = Severity × Impact / 100; no positive score floor.')
    uptime = '95% target. Restoration exposure uses healthy history where available.'
    if not s.include_uptime:
        uptime += ' FLOWERS uptime: hygiene only, excluded from business scoring.'
    return [
        ('Severity', 'Describes the size of the KPI problem.', 'sbu_pharmacy_otc.png'),
        ('Impact', 'Order-related exposure, not euros.', 'storeLocationPin.png'),
        ('Priority', priority, 'goldStarCelebration.png'),
        ('Exposure components', 'Components may overlap. The headline uses the largest component; it is not a guaranteed recovery forecast.', 'sbu_toys_games.png'),
        ('Selection threshold', f'Priority ≥{s.min_priority:g}. The {s.min_priority:g}-point selection threshold is provisional and requires business calibration.', 'sbu_jewelry.png'),
        ('Missing evidence', 'Missing evidence produces an incomplete assessment. Critical observations/inactivity stay in a separate review list.', 'sbu_books.png'),
        ('Benchmark month', f"{result.summary['benchmark_period'] or 'Unavailable'}. SBU peer composition is not available.", 'sbu_electronics.png'),
        ('Uptime', uptime, 'sbu_services.png'),
    ]


def build_guide(prs, result):
    slide = _page(prs, 'How to Read This Review', f"{result.summary['reporting_period']} · {result.settings.conversion_description}")
    for i, (title, body, asset) in enumerate(guide_cards(result)):
        x, y = .65 + (i % 2) * 4.43, 1.52 + (i // 2) * 1.37
        box(slide, x, y, 4.25, 1.25, CREAM if i % 3 == 0 else WHITE, BORDER)
        add_picture_contain(slide, ASSETS_DIR / asset, x + .11, y + .15, .55, .55)
        add_text(slide, title, x + .79, y + .13, 3.30, .29, size=12, bold=True)
        add_text(slide, body, x + .79, y + .50, 3.30, .70, size=10.5)
    box(slide, 9.53, 1.52, 3.14, 5.36, AMBER)
    add_text(slide, 'Key takeaway', 9.77, 1.76, 2.66, .38, size=15, bold=True)
    add_text(slide, 'This review combines data, context and a consistent methodology to highlight where action can improve partner performance.',
             9.77, 2.32, 2.66, 1.43, size=13)
    add_text(slide, 'Remember', 9.77, 4.00, 2.66, .33, size=13, bold=True, color=GREEN)
    reminders = [
        'Performance and operational status are separate.',
        'Exposure is an opportunity signal, not a guaranteed forecast.',
        'Review Store Diagnostics for full store-level context.',
    ]
    for i, text in enumerate(reminders):
        add_text(slide, str(i + 1), 9.77, 4.54 + i * .69, .22, .29, size=11, bold=True, color=GREEN)
        add_text(slide, text, 10.10, 4.54 + i * .69, 2.32, .61, size=11)
