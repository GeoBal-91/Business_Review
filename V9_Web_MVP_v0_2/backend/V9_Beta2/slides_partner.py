"""Existing Dev5 slide narrative/layout, now invoked with explicit per-run context."""
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from theme import GREY, MPL_GREEN, MPL_YELLOW
from presentation_hardening import benchmark_display, COVERAGE_BENCHMARK_NOTE
from formatting import benchmark_gap_info, delta_info, format_percent, get_row_value, pp_change
from ppt_core import add_chart_image, add_footer, add_kpi_card, add_text, add_title, save_chart, set_background


def build(prs, ctx, charts):
    chart_path = charts
    conversion_label = ctx["conversion_label"]
    PARTNER_SBU = ctx['PARTNER_SBU']
    SBU_CANCELS = ctx['SBU_CANCELS']
    SBU_PNA = ctx['SBU_PNA']
    TREND_MONTHS = ctx['TREND_MONTHS']
    UPTIME_TARGET = ctx['UPTIME_TARGET']
    benchmark_period_label = ctx['benchmark_period_label']
    incomplete = ctx['current_partner_coverage_incomplete']
    include_uptime = ctx['include_uptime']
    latest_partner = ctx['latest_partner']
    monthly = ctx['monthly']
    partner_previous = ctx['partner_previous']
    selected_sbu = ctx['selected_sbu']
    uptime_coverage_note = ctx['uptime_coverage_note']


    # 35. PARTNER PERFORMANCE EVOLUTION
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        "Partner Performance Evolution",
        "Monthly GMV and delivered orders"
    )

    add_footer(slide)

    partner_trend = (
        monthly
        .tail(TREND_MONTHS)
        .copy()
    )

    partner_labels = (
        partner_trend[
            "Period"
        ]
        .dt.strftime("%b %y")
    )

    fig, ax1 = plt.subplots(
        figsize=(10.5, 4.5)
    )

    ax1.bar(
        partner_labels,
        partner_trend["€ GMV"],
        color=MPL_YELLOW,
        alpha=0.85,
        label="GMV (€)"
    )

    ax1.set_ylabel("GMV (€)")
    ax1.spines["top"].set_visible(False)

    ax2 = ax1.twinx()

    ax2.plot(
        partner_labels,
        partner_trend[
            "# Delivered Orders"
        ],
        marker="o",
        linewidth=2.5,
        color=MPL_GREEN,
        label="Orders"
    )

    ax2.set_ylabel("Orders")
    ax2.spines["top"].set_visible(False)

    plt.xticks(rotation=35)

    plt.title(
        "GMV & Orders Trend",
        loc="left",
        fontsize=14,
        fontweight="bold"
    )

    handles1, labels1 = ax1.get_legend_handles_labels()
    handles2, labels2 = ax2.get_legend_handles_labels()

    ax1.legend(
        handles1 + handles2,
        labels1 + labels2,
        loc="upper center",
        bbox_to_anchor=(0.5, 1.02),
        ncol=2,
        frameon=False,
        fontsize=9
    )

    plt.tight_layout()

    chart_file = chart_path("partner_performance_trend.png")

    save_chart(
        fig,
        chart_file
    )

    add_chart_image(
        slide,
        chart_file,
        0.8,
        1.35,
        11.8,
        4.9
    )

    # ============================================================
    # 36. GROWTH DRIVERS
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        "Growth Drivers",
        "Order volume and average basket evolution"
    )

    add_footer(slide)

    fig, ax1 = plt.subplots(
        figsize=(10.5, 4.5)
    )

    ax1.plot(
        partner_labels,
        partner_trend[
            "# Delivered Orders"
        ],
        marker="o",
        linewidth=2.5,
        color=MPL_GREEN,
        label="Orders"
    )

    ax1.set_ylabel("Orders")
    ax1.spines["top"].set_visible(False)

    ax2 = ax1.twinx()

    ax2.plot(
        partner_labels,
        partner_trend[
            "€ AOV Derived"
        ],
        marker="o",
        linewidth=2.5,
        color=MPL_YELLOW,
        label="AOV (€)"
    )

    ax2.set_ylabel("AOV (€)")
    ax2.spines["top"].set_visible(False)

    plt.xticks(rotation=35)

    plt.title(
        "Orders vs AOV",
        loc="left",
        fontsize=14,
        fontweight="bold"
    )

    handles1, labels1 = ax1.get_legend_handles_labels()
    handles2, labels2 = ax2.get_legend_handles_labels()

    ax1.legend(
        handles1 + handles2,
        labels1 + labels2,
        loc="upper center",
        bbox_to_anchor=(0.5, 1.02),
        ncol=2,
        frameon=False,
        fontsize=9
    )

    plt.tight_layout()

    chart_file = chart_path("growth_drivers.png")

    save_chart(
        fig,
        chart_file
    )

    add_chart_image(
        slide,
        chart_file,
        0.8,
        1.35,
        11.8,
        4.9
    )


    # ============================================================
    # 37. CONVERSION PERFORMANCE
    # ============================================================

    if "% CVR" in monthly.columns:

        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        set_background(slide)

        add_title(
            slide,
            "Conversion Performance",
            (
                f"Partner CVR vs "
                f"{PARTNER_SBU}"
            )
        )

        add_footer(slide)

        merged_cvr = pd.merge(
            monthly[
                ["Period", "% CVR"]
            ],
            selected_sbu[
                ["Period", "% CVR"]
            ],
            on="Period",
            how="inner",
            suffixes=(
                "_Partner",
                "_SBU"
            )
        )

        merged_cvr = (
            merged_cvr
            .tail(TREND_MONTHS)
            .copy()
        )

        merged_cvr.loc[
            merged_cvr[
                "% CVR_SBU"
            ] <= 0,
            "% CVR_SBU"
        ] = np.nan

        fig, ax = plt.subplots(
            figsize=(10.5, 4.5)
        )

        cvr_labels = (
            merged_cvr[
                "Period"
            ]
            .dt.strftime("%b %y")
        )

        ax.plot(
            cvr_labels,
            merged_cvr[
                "% CVR_Partner"
            ] * 100,
            marker="o",
            linewidth=2.5,
            color=MPL_GREEN,
            label=f"Partner {conversion_label}"
        )

        ax.plot(
            cvr_labels,
            merged_cvr[
                "% CVR_SBU"
            ] * 100,
            marker="o",
            linewidth=2.2,
            linestyle="--",
            color=MPL_YELLOW,
            label=f"{PARTNER_SBU} {conversion_label}"
        )

        ax.set_ylabel(f"{conversion_label} (%)")

        ax.spines[
            "top"
        ].set_visible(False)

        ax.spines[
            "right"
        ].set_visible(False)

        ax.legend(
            frameon=False,
            loc="upper center",
            ncol=2
        )

        plt.xticks(rotation=35)

        plt.title(
            f"{conversion_label} Trend",
            loc="left",
            fontsize=14,
            fontweight="bold"
        )

        plt.tight_layout()

        chart_file = chart_path("conversion_performance.png")

        save_chart(
            fig,
            chart_file
        )

        add_chart_image(
            slide,
            chart_file,
            0.8,
            1.35,
            11.8,
            4.9
        )


    # ============================================================
    # 38. OPERATIONAL HEALTH
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        "Operational Health",
        (
            f"Latest partner performance — Uptime vs 95% target; "
            f"PNA/Cancels vs {PARTNER_SBU} "
            f"({benchmark_period_label})"
        )
    )

    add_footer(slide)

    operational_cards = []

    if "% Uptime" in monthly.columns:

        current = latest_partner[
            "% Uptime"
        ]

        previous_value = get_row_value(
            partner_previous,
            "% Uptime"
        )

        mom = delta_info(
            pp_change(
                current,
                previous_value
            ),
            metric_type="pp",
            higher_is_better=True
        )

        gap = benchmark_display(
            current,
            UPTIME_TARGET,
            incomplete,
            raw_label='Target: ',
            metric_type="pp",
            higher_is_better=True,
            label="95% target"
        )

        operational_cards.append(
            (
                "Operating Store Uptime",
                current,
                mom,
                gap
            )
        )


    if "% PNA Orders" in monthly.columns:

        current = latest_partner[
            "% PNA Orders"
        ]

        previous_value = get_row_value(
            partner_previous,
            "% PNA Orders"
        )

        mom = delta_info(
            pp_change(
                current,
                previous_value
            ),
            metric_type="pp",
            higher_is_better=False
        )

        gap = benchmark_display(
            current,
            SBU_PNA,
            incomplete,
            raw_label='SBU PNA: ',
            metric_type="pp",
            higher_is_better=False,
            label=PARTNER_SBU
        )

        operational_cards.append(
            (
                "Partner PNA",
                current,
                mom,
                gap
            )
        )


    if "% Cancels DTP" in monthly.columns:

        current = latest_partner[
            "% Cancels DTP"
        ]

        previous_value = get_row_value(
            partner_previous,
            "% Cancels DTP"
        )

        mom = delta_info(
            pp_change(
                current,
                previous_value
            ),
            metric_type="pp",
            higher_is_better=False
        )

        gap = benchmark_display(
            current,
            SBU_CANCELS,
            incomplete,
            raw_label='SBU Cancels: ',
            metric_type="pp",
            higher_is_better=False,
            label=PARTNER_SBU
        )

        operational_cards.append(
            (
                "Partner Cancels",
                current,
                mom,
                gap
            )
        )


    op_x = [
        0.75,
        4.8,
        8.85
    ]

    for i, (
        title,
        value,
        mom,
        gap
    ) in enumerate(
        operational_cards[:3]
    ):

        add_kpi_card(
            slide,
            op_x[i],
            1.5,
            3.7,
            1.75,
            title,
            format_percent(value),
            mom["text"],
            mom["color"],
            gap["text"],
            gap["color"]
        )


    if incomplete:
        add_text(slide, "Partner benchmark / target comparisons: " + COVERAGE_BENCHMARK_NOTE,
                 .85, 3.45, 11.2, .35, size=12, bold=True, color=GREY)

    add_text(
        slide,
        "Interpretation",
        0.85,
        3.8,
        3.5,
        0.35,
        size=12,
        bold=True
    )

    add_text(
        slide,
        (
            ("Raw rates reflect stores with current data. "
             + ("FLOWERS uptime remains opening-hours hygiene, excluded from business Priority. " if not include_uptime else "")
             + uptime_coverage_note) if incomplete else
            (
                "Uptime is evaluated against a 95% operating target. "
                "Observed Uptime is treated as opening-hours configuration hygiene "
                "and does not increase business priority. "
                f"Lower is better for PNA and Cancels. {uptime_coverage_note}"
                if not include_uptime
                else
                "Uptime is evaluated against a 95% operating target. "
                "Lower is better for PNA and Cancels. "
                f"Green indicates improvement; red indicates deterioration. {uptime_coverage_note}"
            )
        ),
        0.85,
        4.2,
        11.2,
        0.5,
        size=10,
        color=GREY
    )


    # ============================================================
