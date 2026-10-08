import hashlib
import json
import tempfile
import unittest
import zipfile
from copy import deepcopy
from dataclasses import replace
from pathlib import Path
import pandas as pd
from pandas.testing import assert_frame_equal
from pptx import Presentation
from business_classification import classify, RULES, RULE_FIELDS, calibration_summary, signal
from triage_config import TriageSettings
from triage_exports import export_triage, validation_frame
from triage_presentation import render, texts, action_pages
from triage_fixtures import contract_case, engine_case, flowers_ten
from beta_fixtures import large_partner, visual_fixtures
from presentation_hardening_fixtures import current_partial
from test_regression import run, fixture, zero
from main import export_analysis
from data_pipeline import G, D, N, T, U


def control(a, settings=None):
    t = classify(a, settings)
    return t, t.stores.set_index('Store ID').loc['Control']


class RulesV21(unittest.TestCase):
    def test_base_priority_inclusive_no_flag_required(self):
        a=contract_case(5, cvr=.2, pna=.08, cancels=.025, uptime=.98)
        t,r=control(a); self.assertEqual(r['Action Reason'],'BASE_PRIORITY');self.assertEqual(r['Business Classification'],'ACTION')
        self.assertEqual(a.stores.iloc[0]['Critical Flags'],'')

    def test_priority_full_precision_entry_not_rounded(self):
        for p in [4.99999999999,5.0]:
            t,r=control(contract_case(p,cvr=.2,pna=.08,cancels=.025,uptime=.98))
            self.assertEqual(r['Base Priority Rule'],p>=5)

    def test_vapes_485_single_store_contract(self):
        t,r=control(contract_case(4.85,sbu='VAPES',stores=1))
        self.assertEqual(r['Action Reason'],'SINGLE_STORE_OVERRIDE');self.assertTrue(r['Human Review Required'])
        self.assertTrue(r['Low Scale Flag'])

    def test_operating_count_excludes_lifecycle_rows(self):
        a=run([fixture('Control',**{N:90,'% CVR':.02}),zero('NoActivity'),zero('Closed',**{'Lifecycle Status':'closed'}),fixture('Missing','July')])
        a.stores.loc[a.stores['Store ID']=='Control','Opportunity Score']=4.85
        t,r=control(a);self.assertEqual(r['Operating Commercial Store Count'],1);self.assertEqual(r['Action Reason'],'SINGLE_STORE_OVERRIDE')

    def test_active_incomplete_store_still_counts_as_operating(self):
        a=run([fixture('Control',**{N:90,'% CVR':.02}),fixture('Incomplete',**{'% CVR':float('nan')})])
        a.stores.loc[a.stores['Store ID']=='Control','Opportunity Score']=4.85
        t,r=control(a);self.assertEqual(r['Operating Commercial Store Count'],2);self.assertFalse(r['Single-Store Override'])

    def test_single_store_primary_floor(self):
        t,r=control(contract_case(4.85,stores=1,primary='Cancels',attempts=3,delivered=1,cvr=.2,uptime=.98,pna=.08,cancels=.5))
        self.assertFalse(r['Single-Store Override']);self.assertEqual(r['Business Classification'],'MONITOR')

    def test_pets_424_multi_kpi_contract(self):
        t,r=control(contract_case(4.24))
        self.assertEqual(r['Action Reason'],'MULTI_KPI_ESCALATION');self.assertTrue(r['Low Scale Flag'])
        self.assertIn('SEVERE_PRIMARY_CORROBORATION',r['Supporting Rule Matches'])

    def test_pharmacy_453_ultra_low_scale_contract(self):
        t,r=control(contract_case(4.53,sbu='PHARMACY_OTC',sessions=78,attempts=5,delivered=2,gmv=14,cvr=.064))
        self.assertEqual(r['Action Reason'],'MULTI_KPI_ESCALATION');self.assertTrue(r['Ultra Low Scale Flag']);self.assertTrue(r['Human Review Required'])

    def test_uptime_841_vs_852_severe_boundary_contract(self):
        for p,u,expected in [(4.13,.841,'ACTION'),(4.91,.852,'MONITOR')]:
            t,r=control(contract_case(p,sbu='BEAUTY_PERSONAL_CARE',sessions=100,cvr=.19,pna=.08,cancels=.025,uptime=u))
            self.assertEqual(r['Business Classification'],expected)
        self.assertEqual(r['Uptime Interpretation'],'MATERIALLY_ADVERSE')

    def test_single_uptime_band_does_not_create_action(self):
        for u in [.85,.8500001,.90,.9000001,.9499,.95]:
            t,r=control(contract_case(4.91,cvr=.2,pna=.08,cancels=.025,uptime=u))
            self.assertEqual(r['Business Classification'],'MONITOR')

    def test_evidence_floor_boundaries(self):
        for metric,low,high in [('CVR',49,50),('PNA',9,10),('Cancels',9,10)]:
            for n,expected in [(low,False),(high,True)]:
                s=signal({'metric':metric,'eligible':True,'state':'Assessed','raw_rate':.02 if metric=='CVR' else .5,'benchmark':.2 if metric=='CVR' else .05,'denominator':n},True)
                self.assertEqual(s['evidence_floor'],expected)

    def test_pna_cancels_tiny_denominator_not_independent_domains(self):
        t,r=control(contract_case(4.24,cvr=.2,uptime=.98,attempts=3,delivered=1,pna=1/3,cancels=1/3))
        self.assertEqual(r['Business Classification'],'MONITOR');self.assertEqual(r['Qualified Adverse Domains'],[])

    def test_electronics_382_corroboration_contract(self):
        t,r=control(contract_case(3.82,sbu='ELECTRONICS',sessions=189,attempts=3,delivered=1,cvr=.021,benchmark_cvr=.065,pna=1/3,cancels=1/3,uptime=.922))
        self.assertEqual(r['Action Reason'],'SEVERE_PRIMARY_CORROBORATION')
        self.assertTrue(r['Ultra Low Scale Flag']);self.assertTrue(r['Human Review Required'])
        self.assertIn('corroboration only',r['Likely Operational Driver']);self.assertIn('root cause',r['V2.1 Recommended Action'])

    def test_severe_cvr_requires_both_relative_and_absolute_gap(self):
        for raw,b,expected in [(.03,.05,False),(.06,.10,True),(.061,.1,False)]:
            s=signal({'metric':'CVR','state':'Assessed','eligible':True,'raw_rate':raw,'benchmark':b,'denominator':100},True)
            self.assertEqual(s['severe'],expected)

    def test_corroboration_is_not_trivial_miss(self):
        t,r=control(contract_case(3.82,uptime=.922,pna=.081,cancels=.026))
        self.assertFalse(r['Severe Primary + Corroboration']);self.assertEqual(r['Business Classification'],'MONITOR')

    def test_zero_benchmark_tiny_miss_not_meaningful_ratio(self):
        s=signal({'metric':'PNA','state':'Assessed','eligible':True,'raw_rate':.001,'benchmark':0,'denominator':100},True)
        self.assertFalse(s['meaningful']);self.assertFalse(s['material'])
        s=signal({'metric':'PNA','state':'Assessed','eligible':True,'raw_rate':.03,'benchmark':0,'denominator':3},True)
        self.assertTrue(s['meaningful']);self.assertFalse(s['evidence_floor'])

    def test_corroboration_needs_cvr_sessions_floor(self):
        t,r=control(contract_case(3.82,sessions=49));self.assertFalse(r['Severe Primary + Corroboration'])

    def test_uptime_corroboration_92_inclusive(self):
        for u,expected in [(.92,True),(.920001,False)]:
            t,r=control(contract_case(3.82,uptime=u,pna=.08,cancels=.025))
            self.assertEqual(r['Severe Primary + Corroboration'],expected)

    def test_home_356_materiality_contract(self):
        a=contract_case(3.56,sbu='HOME_DECORATION',sessions=400,attempts=100,delivered=80,gmv=6708,cvr=.2,pna=.234,cancels=.146,uptime=.897,benchmark_pna=.149,benchmark_cancels=.094)
        # Partner GMV specified in the brief, kept as a classification context fixture.
        a.data.monthly.loc[a.data.monthly['Period']==a.data.reporting_period,G]=22680
        t,r=control(a);self.assertEqual(r['Action Reason'],'BUSINESS_MATERIALITY_OVERRIDE')
        self.assertAlmostEqual(r['Store GMV Share'],6708/22680);self.assertEqual(r['Opportunity Score'],3.56)

    def test_top_two_alone_is_insufficient(self):
        t,r=control(contract_case(3.56,cvr=.2,sessions=100,attempts=20,delivered=5,gmv=10,pna=.23,cancels=.15,uptime=.89))
        self.assertFalse(r['Business Materiality Override'])

    def test_top_two_exposure_guard_delivery_or_sessions(self):
        for d,n,expected in [(19,199,False),(20,199,True),(19,200,True)]:
            t,r=control(contract_case(3.56,cvr=.2,delivered=d,sessions=n,attempts=100,gmv=10,pna=.23,cancels=.15,uptime=.89))
            self.assertEqual(r['Business Materiality Override'],expected)

    def test_materiality_requires_two_qualified_ops(self):
        t,r=control(contract_case(3.56,cvr=.2,sessions=300,attempts=3,delivered=1,gmv=6708,pna=.5,cancels=.5,uptime=.89))
        self.assertFalse(r['Business Materiality Override'])

    def test_materiality_partial_coverage_not_inferred(self):
        a=current_partial();t=classify(a)
        self.assertFalse(t.stores['Materiality Assessed'].any());self.assertFalse(t.stores['Business Materiality Override'].any())

    def test_precedence_and_supporting_rules_once(self):
        t,r=control(contract_case(4.85,stores=1))
        self.assertEqual(r['Action Reason'],'SINGLE_STORE_OVERRIDE');self.assertIn('MULTI_KPI_ESCALATION',r['Supporting Rule Matches'])
        self.assertEqual(t.summary['business_classification_counts']['ACTION'],1)

    def test_monitor_priority_only_controls_need_extra_evidence(self):
        # Brief supplies only these priorities: this tests non-promotion by score alone, not real-row replay.
        for p in [3.71,3.61,3.62,3.47,3.27,4.91,1.43]:
            t,r=control(contract_case(p,cvr=.2,pna=.08,cancels=.025,uptime=.98))
            self.assertEqual(r['Business Classification'],'MONITOR')

    def test_no_activity_closed_missing_and_incomplete_unclassified(self):
        a=run([zero('NoActivity'),zero('Closed',**{'Lifecycle Status':'closed'}),fixture('Missing','July'),fixture('Incomplete',**{'% CVR':float('nan')})])
        t=classify(a);self.assertTrue((t.stores['Business Classification']=='UNCLASSIFIED').all())
        states=t.stores.set_index('Store ID')['Operational/Data State']
        self.assertIn('NO_ACTIVITY',states['NoActivity']);self.assertIn('CLOSED',states['Closed']);self.assertIn('MISSING_CURRENT',states['Missing']);self.assertIn('DATA_REVIEW',states['Incomplete'])

    def test_zero_attempts_not_missing_when_engine_says_no_exposure(self):
        from data_pipeline import P,C
        a=run([fixture(**{D:0,G:0,N:10000,T:0,P:0,C:0,'% CVR':0})])
        t=classify(a);self.assertEqual(t.stores.iloc[0]['Business Classification'],'ACTION')

    def test_healthy_tolerance_seven_sbu_synthetic_examples(self):
        for sbu in ['PETS','PHARMACY_OTC','VAPES','BEAUTY_PERSONAL_CARE','HOME_DECORATION','ELECTRONICS','FLOWERS']:
            t,r=control(contract_case(.2,sbu=sbu,cvr=.195,pna=.082,cancels=.027,uptime=.94))
            self.assertEqual(r['Business Classification'],'HEALTHY');self.assertTrue(r['Healthy Tolerance Candidate'])

    def test_tolerance_cannot_hide_severe_or_material_issues(self):
        for kwargs in [{'uptime':.90},{'pna':.20},{'cvr':.02}]:
            args=dict(cvr=.2,pna=.08,cancels=.025,uptime=.95);args.update(kwargs)
            t,r=control(contract_case(.1,**args));self.assertEqual(r['Business Classification'],'MONITOR')

    def test_tolerance_config_isolated_and_conservative(self):
        a=contract_case(.2,cvr=.195,pna=.08,cancels=.025,uptime=.95)
        before=a.stores.copy(deep=True);t,r=control(a,TriageSettings(healthy_cvr_gap_pp=.1))
        self.assertEqual(r['Business Classification'],'MONITOR');assert_frame_equal(before,a.stores,check_exact=True)
        with self.assertRaises(ValueError):TriageSettings(healthy_priority_below=3)

    def test_low_scale_flags_never_change_classification_or_scores(self):
        a=contract_case(4.24);b=deepcopy(a)
        b.stores.loc[b.stores['Store ID']=='Control',[G,D]]= [0,0]
        ta,ra=control(a);tb,rb=control(b)
        self.assertEqual(ra['Opportunity Score'],rb['Opportunity Score']);self.assertEqual(ra['Business Classification'],rb['Business Classification'])

    def test_flowers_zhiuli_503_base_contract(self):
        t,r=control(contract_case(5.03,sbu='FLOWERS',uptime=.626))
        self.assertEqual(r['Action Reason'],'BASE_PRIORITY');self.assertIn('HYGIENE_REVIEW',r['Operational/Data State'])

    def test_flowers_tarieli_143_uptime_cannot_promote(self):
        t,r=control(contract_case(1.43,sbu='FLOWERS',cvr=.2,pna=.08,cancels=.025,uptime=.01))
        self.assertEqual(r['Business Classification'],'MONITOR');self.assertIn('HYGIENE_REVIEW',r['Operational/Data State'])

    def test_flowers_uptime_cannot_enter_any_commercial_override(self):
        for p in [3.56,3.82,4.24]:
            t,r=control(contract_case(p,sbu='FLOWERS',pna=.08,cancels=.025,uptime=.01))
            self.assertEqual(r['Business Classification'],'MONITOR')
            self.assertNotIn('Uptime',r['Qualified Adverse Domains']);self.assertNotIn('availability',r['Likely Operational Driver'].lower())

    def test_flowers_ten_mix_and_explicit_duplicate_profile_alias(self):
        t=classify(flowers_ten())
        self.assertEqual(t.summary['business_classification_counts'],{'ACTION':1,'MONITOR':2,'HEALTHY':3,'UNCLASSIFIED':4})
        s=calibration_summary({'FLOWERS#3':t,'FLOWERS#4':deepcopy(t)},{'FLOWERS#4':'FLOWERS#3'})
        self.assertEqual(s['known_stores'],10);self.assertEqual(s['canonical_profiles'],1)

    def test_conflicting_duplicate_profile_is_not_silently_dropped(self):
        a=classify(flowers_ten());b=deepcopy(a);b.analysis.stores.loc[0,G]+=1
        with self.assertRaises(AssertionError):calibration_summary({'3':a,'4':b},{'4':'3'})

    def test_equal_addresses_different_ids_never_merged(self):
        a=run([fixture('A',**{'Store Address Address':'Same'}),fixture('B',**{'Store Address Address':'Same'})])
        t=classify(a);self.assertEqual(len(t.stores),2)

    def test_critical_flags_do_not_drive_classification(self):
        a=contract_case(3.2,cvr=.2,pna=.08,cancels=.025,uptime=.98)
        b=deepcopy(a);b.stores['Critical Flags']='Critical PNA'
        for k in ['Business Classification','Action Reason','Opportunity Score']:
            self.assertEqual(control(a)[1][k],control(b)[1][k])

    def test_primary_kpi_not_replaced_by_likely_driver(self):
        a=contract_case(4.24,primary='CVR');t,r=control(a)
        self.assertEqual(r['Primary Action KPI'],'CVR');self.assertIn('availability',r['Likely Operational Driver'])
        self.assertEqual(a.stores.loc[a.stores['Store ID']=='Control','Main Driver'].iloc[0],'CVR')

    def test_causal_wording_guard(self):
        for args in [{},{'primary':'Cancels'},{'primary':'PNA'},{'attempts':3,'delivered':1,'uptime':.922}]:
            t,r=control(contract_case(4.24,**args))
            joined=r['Likely Operational Driver']+' '+r['V2.1 Recommended Action']
            self.assertNotIn(' because ',joined.lower());self.assertNotIn('guaranteed',joined.lower());self.assertNotIn('forecast',joined.lower())


class ExportsAndFreeze(unittest.TestCase):
    def test_all_beta_source_tests_and_assets_unchanged(self):
        root=Path(__file__).parents[1];baseline=json.loads((root/'tests/beta1_0_frozen_files.json').read_text())
        for name,digest in baseline.items():self.assertEqual(hashlib.sha256((root/name).read_bytes()).hexdigest(),digest,name)

    def test_classification_is_read_only_for_all_beta_fixtures(self):
        for name,a in visual_fixtures():
            old=deepcopy(a);t=classify(a)
            for key in ['stores','metrics','opportunities']:assert_frame_equal(getattr(old,key),getattr(a,key),check_exact=True)
            self.assertEqual(old.actions,a.actions);self.assertEqual(old.reviews,a.reviews);self.assertEqual(old.summary,a.summary)
            assert_frame_equal(t.stores[a.stores.columns],a.stores,check_exact=True)

    def test_deterministic_classification_and_full_rule_traces(self):
        a=large_partner(30);x=classify(a);y=classify(a)
        assert_frame_equal(x.stores,y.stores,check_exact=True);self.assertEqual(x.traces,y.traces)
        for trace in x.traces:
            self.assertEqual(list(trace['rules']),list(RULES))
            for fields in trace['rules'].values():self.assertEqual(set(fields),{'entry_range','evidence_floor','severe_or_material_condition','materiality','outcome'})

    def test_action_order_uses_same_priority_impact_severity_id_keys(self):
        a=large_partner(30);t=classify(a)
        expected=t.stores[t.stores['Business Classification']=='ACTION'].sort_values(['Opportunity Score','Business Impact Score','Performance Severity Score','Store ID'],ascending=[False,False,False,True],kind='mergesort').head(6)['Store ID'].tolist()
        self.assertEqual([x['Store ID'] for x in t.actions],expected)

    def test_legacy_exports_byte_identical_and_new_exports_precise(self):
        a=large_partner(30);t=classify(a)
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);export_analysis(a,p/'old');export_triage(t,p/'new')
            for f in (p/'old').iterdir():self.assertEqual(f.read_bytes(),(p/'new'/f.name).read_bytes(),f.name)
            re=pd.read_csv(p/'new/Business_Classification.csv',float_precision='round_trip')
            for name in ['Opportunity Score','Performance Severity Score','Business Impact Score']:
                pd.testing.assert_series_equal(re[name],a.stores[name],check_exact=True)
            for f in (p/'new').glob('*.json'):json.loads(f.read_text(),parse_constant=lambda x: self.fail('Nonfinite JSON '+x))

    def test_validation_manual_columns_blank_and_original_values_exact(self):
        a=current_partial();t=classify(a);f=validation_frame(t)
        self.assertTrue((f.filter(regex='^AM_')=='').all().all());self.assertEqual(len(f),len(a.stores))
        self.assertTrue(pd.isna(f.set_index('Store ID').loc['Large','Priority']))

    def test_empty_actions_keep_csv_schema(self):
        with tempfile.TemporaryDirectory() as tmp:
            export_triage(classify(run()),tmp)
            f=pd.read_csv(Path(tmp)/'Business_Actions.csv');self.assertTrue(f.empty);self.assertIn('Action Reason',f)

    def test_new_cli_analysis_only(self):
        from main_v2_1 import main
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'test.pptx';t=main(['--analysis-only','--output',str(out)])
            self.assertFalse(out.exists());self.assertTrue((Path(tmp)/'test_audit/Classification_Summary.json').is_file())

    def test_challenger_not_silently_classified_as_production(self):
        a=run(priority_model='simple_product');t=classify(a)
        self.assertTrue((t.stores['Business Classification']=='UNCLASSIFIED').all())


class TriageDeck(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory();cls.t=classify(large_partner(30));cls.path=render(cls.t,Path(cls.tmp.name)/'V2.1.pptx');cls.prs=Presentation(cls.path)

    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def test_new_actions_have_reason_driver_without_severity_impact(self):
        slides=[s for s in self.prs.slides if any(t.startswith('Priority Opportunities & Actions') for t in texts(s))]
        joined='\n'.join('\n'.join(texts(s)) for s in slides)
        self.assertIn('Likely operational driver:',joined);self.assertIn('ACTION ·',joined)
        self.assertNotIn('Severity',joined);self.assertNotIn('Impact',joined)
        ids=[s.name.split(': ',1)[1] for slide in slides for s in slide.shapes if s.name.startswith('Opportunity store: ')]
        self.assertEqual(ids,[a['Store ID'] for a in self.t.actions])

    def test_action_focus_and_pagination_consistent(self):
        self.assertEqual([len(x) for x in action_pages([1,2,3,4,5])],[2,2,1])
        self.assertEqual([len(x) for x in action_pages([1,2,3,4])],[2,2])
        alltext='\n'.join('\n'.join(texts(s)) for s in self.prs.slides)
        self.assertIn('Highest',alltext);self.assertIn('Lowest',alltext);self.assertIn('selected of',alltext)

    def test_main_order_and_all_diagnostics_retained(self):
        tt=[texts(s)[0] for s in self.prs.slides]
        self.assertLess(tt.index('Top GMV Stores'),tt.index('How We Prioritize Opportunities'))
        self.assertLess(tt.index('Thank You'),tt.index('APPENDIX — Store Diagnostics'))
        self.assertEqual(tt.count('Store Diagnostics'),30)

    def test_overview_business_distinct_from_status_and_all_links_retained(self):
        over=[s for s in self.prs.slides if texts(s)[0].startswith('Store Overview')]
        self.assertTrue(all('Business' in texts(s) and 'Status' in texts(s) for s in over))
        self.assertEqual(sum(s.name.startswith('Overview store: ') for slide in over for s in slide.shapes),30)
        for slide in over:
            for s in slide.shapes:
                if s.name.startswith('Overview store: '):self.assertIsNotNone(s.click_action.target_slide)

    def test_diagnostics_keep_exact_score_labels_and_add_business_metadata(self):
        diag=[s for s in self.prs.slides if texts(s)[0]=='Store Diagnostics']
        for s in diag:
            self.assertIn('Severity',texts(s));self.assertIn('Impact',texts(s));self.assertTrue(any(x.startswith('Business:') for x in texts(s)))

    def test_beta11_visible_version_and_native_bounds(self):
        joined='\n'.join('\n'.join(texts(s)) for s in self.prs.slides)
        self.assertIn('V9_Beta2',joined);self.assertNotIn('V9 Beta 1.0',joined)
        for i,slide in enumerate(self.prs.slides):
            for s in slide.shapes:
                self.assertGreaterEqual(s.left,-9144);self.assertGreaterEqual(s.top,-9144)
                self.assertLessEqual(s.left+s.width,self.prs.slide_width+9144,(i,s.name))
                self.assertLessEqual(s.top+s.height,self.prs.slide_height+9144,(i,s.name))

    def test_unique_pptx_parts_after_section_replacement(self):
        with zipfile.ZipFile(self.path) as z:
            self.assertEqual(len(z.namelist()),len(set(z.namelist())))
            self.assertIsNone(z.testzip())

    def test_partial_coverage_guards_remain(self):
        p=render(classify(current_partial()),Path(self.tmp.name)/'Partial.pptx');prs=Presentation(p)
        for title in ['Executive Summary','Operational Health']:
            s=next(s for s in prs.slides if title in texts(s));joined=' '.join(texts(s))
            self.assertIn('incomplete store coverage',joined.lower())

    def test_override_appears_in_actions_plan_and_diagnostics(self):
        a=contract_case(3.56,cvr=.2,sessions=400,attempts=100,delivered=80,gmv=6708,pna=.234,cancels=.146,uptime=.897,benchmark_pna=.149,benchmark_cancels=.094)
        t=classify(a);p=render(t,Path(self.tmp.name)/'Override.pptx');prs=Presentation(p)
        joined='\n'.join('\n'.join(texts(s)) for s in prs.slides)
        self.assertIn('Business Materiality Override',joined);self.assertIn('Priority 3.56',joined)
        self.assertIn('1 selected of 1 ACTION store',joined);self.assertNotIn('selected of 1 ACTION stores',joined)
        plan=next(s for s in prs.slides if 'Action Plan' in texts(s));self.assertIn(t.actions[0]['Action'],'\n'.join(texts(plan)))
