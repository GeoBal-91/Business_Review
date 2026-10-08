"""Dev6.1 acceptance criteria A–G; fixture-based, no threshold calibration."""
import unittest
import tempfile
import json
from pathlib import Path
from dataclasses import replace
from datetime import date
import pandas as pd
import numpy as np
from pptx import Presentation
from pptx.util import Inches
from test_regression import run,fixture,zero
from data_pipeline import D,N,T,P,C,G,U,load_frames,DataValidationError
from config import Settings
from analysis import analyze
from main import export_analysis
from presentation import context,render
from coverage_guard import comparison,comparison_exports,NOT_COMPARABLE
from diagnostics import build as diagnostics_build,trend_data,confidence_label,metric_records


def text_of(deck):
    return '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)


def diagnostic_deck(a):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    with tempfile.TemporaryDirectory() as tmp:
        diagnostics_build(prs,a,lambda name:str(Path(tmp)/name))
    return prs


def disappearance():
    rows=[]
    for month,year in [('August',2025),('July',2026)]:
        rows.extend([fixture('Large',month,**{D:1000,T:2000,N:5000,G:10000,'Year':year}),fixture('Small',month,**{D:100,T:200,N:500,G:1000,'Year':year})])
    rows.append(fixture('Small',**{D:100,T:200,N:500,G:1000}))
    return run(rows)


class CVRContract(unittest.TestCase):
    def test_store_scoring_uses_reported(self):
        a=run([fixture(**{'% CVR':.07})]);self.assertEqual(a.metrics.set_index('metric').loc['CVR','raw_rate'],.07)
        self.assertEqual(a.stores.iloc[0]['% CVR'],.07)
    def test_partner_reported_weighted(self):
        a=run([fixture(**{'% CVR':.07}),fixture('B',**{N:100,'% CVR':.5})])
        self.assertAlmostEqual(a.data.monthly.iloc[0]['% CVR'],120/1100)
    def test_sbu_reported_weighted(self):
        a=run(benchmarks=[fixture(**{'% CVR':.07}),fixture('Peer',**{N:100,'% CVR':.5})],allow_partitioned_rows=True)
        self.assertAlmostEqual(a.data.benchmark_row['% CVR'],120/1100)
    def test_proxy_never_substitutes(self):
        a=run([fixture(**{'% CVR':.07})]);r=a.stores.iloc[0]
        self.assertEqual(r['Delivered / Sessions Diagnostic'],.2)
        self.assertAlmostEqual(r['CVR Reconciliation Gap'],-.13)
        for mode in ['converted_sessions','delivered_per_session']:
            with self.assertRaises(ValueError): Settings(cvr_mode=mode)
    def test_missing_reported_incomplete(self):
        a=run([fixture(**{'% CVR':np.nan})]);r=a.stores.iloc[0]
        self.assertEqual(r['Scoring Status'],'Incomplete');self.assertTrue(pd.isna(r['% CVR']))
        self.assertEqual(r['Delivered / Sessions Diagnostic'],.2)
        self.assertTrue(pd.isna(r['Opportunity Score']))
    def test_missing_reported_column_rejected(self):
        row=fixture();del row['% CVR']
        with self.assertRaises(DataValidationError):run([row])
    def test_store_observation_retained_with_zero_sessions(self):
        a=run([fixture(**{N:0,'% CVR':.07})]);self.assertEqual(a.stores.iloc[0]['% CVR'],.07)
        self.assertEqual(a.stores.iloc[0]['Scoring Status'],'Incomplete')
    def test_missing_rate_confidence_not_presented_as_percent(self):
        a=run([fixture(**{P:np.nan})]);m=metric_records(a,'A')['PNA'];label,weight=confidence_label(m)
        self.assertIn('n=400',label);self.assertIn('Rate unavailable',label);self.assertNotIn('%',label);self.assertIsNone(weight)


class CoverageSafety(unittest.TestCase):
    def test_missing_current_unknown(self):
        a=disappearance();r=a.stores.set_index('Store ID').loc['Large']
        self.assertEqual(r['Store Status'],'Unknown');self.assertEqual(r['Data Quality'],'Missing current data')
    def test_raw_totals_retained_both_comparisons_suppressed(self):
        a=disappearance();self.assertEqual(a.data.monthly.iloc[-1][D],100)
        for months in [1,12]:
            c=comparison(a.data,a.data.reporting_period,a.data.reporting_period-pd.DateOffset(months=months))
            self.assertFalse(c.comparable);self.assertEqual(c.reason,NOT_COMPARABLE)
        c=context(a);self.assertIsNone(c['partner_previous']);self.assertTrue(pd.isna(c['previous_share']))
    def test_coverage_counts_and_names_exported(self):
        a=disappearance();r=a.data.coverage.iloc[-1]
        self.assertEqual((r['Expected Stores'],r['Present Stores'],r['Missing Stores']),(2,1,1))
        self.assertEqual(r['Missing Store Names'],'Large')
        rows=comparison_exports(a.data);rows=rows[rows.Period==a.data.reporting_period]
        self.assertTrue(rows[D+' Change'].isna().all())
        with tempfile.TemporaryDirectory() as tmp:
            export_analysis(a,tmp)
            self.assertEqual(pd.read_csv(Path(tmp)/'coverage.csv').iloc[-1]['Missing Stores'],1)
    def test_previous_period_partial_also_suppressed(self):
        a=run([fixture('A','June'),fixture('B','June'),fixture('A','July'),fixture('A'),fixture('B')])
        self.assertFalse(comparison(a.data,pd.Timestamp('2026-08-01'),pd.Timestamp('2026-07-01')).comparable)
    def test_future_stores_not_expected_before_first_observation(self):
        a=run([fixture('A','July'),fixture('A'),fixture('New')])
        self.assertEqual(a.data.coverage.iloc[0]['Expected Stores'],1)
        self.assertEqual(a.data.coverage.iloc[0]['Missing Stores'],0)
    def test_coverage_does_not_change_store_scores(self):
        one=run([fixture('B',**{'% CVR':.01})])
        two=run([fixture('A','July'),fixture('B',**{'% CVR':.01})])
        self.assertEqual(one.stores.iloc[0]['Opportunity Score'],two.stores.set_index('Store ID').loc['B','Opportunity Score'])
    def test_no_false_growth_in_full_deck(self):
        a=disappearance()
        with tempfile.TemporaryDirectory() as tmp:
            deck=Presentation(render(a,Path(tmp)/'coverage.pptx'));text=text_of(deck)
        self.assertIn('Partial Store Coverage',text);self.assertIn('Expected 2 · Present 1 · Missing 1',text)
        self.assertIn('Missing: Large',text);self.assertIn('MoM: '+NOT_COMPARABLE,text)
        self.assertNotIn('-90.9%',text);self.assertNotIn('▼ 90.9%',text)
        self.assertNotIn('GMV changed',text);self.assertNotIn('Order volume moved more sharply',text)


class DiagnosticStates(unittest.TestCase):
    def test_active_complete_dashboard(self):
        text=text_of(diagnostic_deck(run()))
        for label in ['ACTIVE','DATA OK','COMPLETE','Severity','Impact','Priority','Main Business Action KPI','Confidence / evidence','Recommended Action']:
            self.assertIn(label,text)
    def test_material_priority_dashboard(self):
        a=run([fixture(**{N:10000,'% CVR':0})]);self.assertIn('MATERIAL',text_of(diagnostic_deck(a)))
    def test_incomplete_uses_known_issue(self):
        a=run([fixture(**{N:10000,'% CVR':0,P:np.nan})]);text=text_of(diagnostic_deck(a))
        self.assertIn('Leading known issue',text);self.assertNotIn('Main Business Action KPI',text)
        self.assertIn('Not assessed',text);self.assertIn('Required data review',text);self.assertIn('Unavailable metrics: PNA',text)
    def test_unknown_dashboard(self):
        a=run([fixture('A','July'),fixture('B')]);deck=diagnostic_deck(a);text=text_of(PresentationSlice(deck.slides[0]))
        self.assertIn('UNKNOWN',text);self.assertIn('MISSING DATA',text);self.assertIn('INCOMPLETE',text);self.assertIn('Missing current evidence',text)
    def test_flowers_hygiene_dashboard(self):
        settings=Settings(partner_sbu='FLOWERS',as_of_date=date(2026,9,7))
        a=analyze(load_frames(pd.DataFrame([fixture(**{U:.4,N:10000,'% CVR':0})]),pd.DataFrame([fixture(**{'Store Sub Business Unit':'FLOWERS'})]),settings))
        self.assertEqual(a.stores.iloc[0]['Main Driver'],'CVR')
        self.assertIn('Hygiene only — excluded from business Priority',text_of(diagnostic_deck(a)))
    def test_closed_dashboard(self):
        a=run([zero(**{'Lifecycle Status':'closed'})]);text=text_of(diagnostic_deck(a))
        self.assertIn('CLOSED',text);self.assertIn('NO ROUTINE RECOVERY',text);self.assertIn('Lifecycle / activity evidence',text)
        self.assertNotIn('MATERIAL',text);self.assertNotIn('Main Business Action KPI',text)
    def test_no_activity_keeps_zero_and_review(self):
        a=run([fixture('A','July'),zero()]);self.assertEqual(a.stores.iloc[0]['Opportunity Score'],0)
        text=text_of(diagnostic_deck(a));self.assertIn('NO ACTIVITY',text);self.assertTrue(a.reviews)
    def test_critical_dashboard_raw_denominator_not_material(self):
        a=run([fixture(**{D:0,G:0,N:1,T:5,P:0,C:1,'% CVR':0})]);text=text_of(diagnostic_deck(a))
        self.assertIn('Critical Review',text);self.assertIn('Cancels 20.0% · n=5',text)
        self.assertIn('BELOW MAIN THRESHOLD',text)
    def test_main_differs_from_worst(self):
        a=run([fixture(**{D:0,G:0,N:100000,T:1,P:0,C:1,'% CVR':.18})]);r=a.stores.iloc[0]
        self.assertEqual(r['Main Driver'],'CVR');self.assertEqual(r['Severity Driver'],'Cancels')
        text=text_of(diagnostic_deck(a));self.assertIn('Main Business Action KPI\nCVR',text);self.assertIn('Worst KPI / Severity Driver\nCancels',text)
    def test_main_differs_from_largest_exposure(self):
        a=run([fixture(**{D:0,G:0,N:100000,T:500,P:0,C:250,'% CVR':.18})]);r=a.stores.iloc[0]
        self.assertEqual(r['Main Driver'],'Cancels');self.assertEqual(r['Impact Source'],'CVR')
        text=text_of(diagnostic_deck(a));self.assertIn('Main Business Action KPI\nCancels',text);self.assertIn('Largest Exposure KPI\nCVR',text)
    def test_trend_available_and_no_fabricated_missing_points(self):
        a=run([fixture('A','March',**{'% CVR':.15}),fixture('A','May',**{'% CVR':.1}),fixture('A',**{'% CVR':.05})])
        trend=trend_data(a,a.stores.iloc[0]);self.assertIsNotNone(trend)
        self.assertEqual(trend[1].notna().sum(),3);self.assertTrue(pd.isna(trend[1].loc['2026-04-01']))
        self.assertNotIn('Trend unavailable',text_of(diagnostic_deck(a)))
    def test_trend_unavailable(self):
        a=run([fixture(**{'% CVR':.05})]);self.assertIsNone(trend_data(a,a.stores.iloc[0]))
        self.assertIn('Trend unavailable',text_of(diagnostic_deck(a)))


class PresentationQualifications(unittest.TestCase):
    def test_partial_period_and_challenger_every_slide(self):
        a=run([fixture(month='September')],[fixture(month='September')],allow_partial_period=True,priority_model='simple_product')
        with tempfile.TemporaryDirectory() as tmp:
            deck=Presentation(render(a,Path(tmp)/'experimental.pptx'))
        for slide in deck.slides:
            text='\n'.join(s.text for s in slide.shapes if s.has_text_frame)
            self.assertIn('PARTIAL PERIOD — PROVISIONAL ANALYSIS',text)
            self.assertIn('EXPERIMENTAL / CHALLENGER MODEL',text)
    def test_standard_default_paired_unmarked_and_challenger_exported(self):
        a=run();self.assertEqual(a.settings.priority_model,'paired')
        self.assertIn('Simple Product Priority',a.stores)
        with tempfile.TemporaryDirectory() as tmp:
            deck=Presentation(render(a,Path(tmp)/'standard.pptx'))
        text=text_of(deck);self.assertNotIn('EXPERIMENTAL / CHALLENGER MODEL',text);self.assertNotIn('PARTIAL PERIOD',text)
        self.assertIn('The 5.0 threshold is provisional and will be calibrated during Beta validation.',text)


class PresentationSlice:
    def __init__(self,slide): self.slides=[slide]


class Dev6GoldenCompatibility(unittest.TestCase):
    def test_untouched_scoring_and_lifecycle_source(self):
        import hashlib
        expected=json.loads((Path(__file__).parent/'dev6_golden.json').read_text())
        for name,digest in expected['source_hashes'].items():
            self.assertEqual(hashlib.sha256((Path(__file__).parents[1]/name).read_bytes()).hexdigest(),digest)
    def test_80_store_scores_flags_actions_rank_against_dev6(self):
        expected=json.loads((Path(__file__).parent/'dev6_golden.json').read_text())
        result=run(expected['rows'])
        actual=result.stores.set_index('Store ID')
        for record in expected['expected']:
            for key,value in record.items():
                if key=='Store ID':continue
                got=actual.loc[record['Store ID'],key]
                if isinstance(value,float):self.assertAlmostEqual(got,value,places=10,msg=key)
                else:self.assertEqual(got,value,key)
        self.assertEqual(result.opportunities['Store ID'].tolist(),expected['ranked_ids'])
        for actual,old in zip(result.actions,expected['expected_actions']):
            self.assertEqual(actual['Store ID'],old['Store ID']);self.assertEqual(actual['Action'],old['Action']);self.assertEqual(actual['Metric'],old['Metric'])
