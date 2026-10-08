"""Generate synthetic V2.1 examples through the actual, frozen engine."""
from pathlib import Path
from dataclasses import replace
import json
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from analysis import run_analysis
from config import Settings
from business_classification import classify
from triage_exports import export_triage, write_json
from triage_presentation import render
from triage_fixtures import new_visual_fixtures
from beta_fixtures import large_partner
from test_regression import run
from presentation_hardening_fixtures import current_partial
from ppt_core import add_text
from pptx import Presentation


def build(destination):
    destination=Path(destination);destination.mkdir(parents=True,exist_ok=True)
    cases=[('Uploaded_Beauty_Partner',run_analysis(Settings()))] + new_visual_fixtures() + [
        ('Partial_Current_Coverage',current_partial()),('Thirty_Stores',large_partner(30)),('Healthy_Empty',run())]
    inventory=[]
    for name,a in cases:
        t=classify(a)
        filename='Partner_Business_Review_V9_Beta1_1.pptx' if name=='Uploaded_Beauty_Partner' else 'V9_Beta1_1_'+name+'.pptx'
        path=destination/filename
        export_triage(t,destination/(path.stem+'_audit'))
        render(t,path)
        prs=Presentation(path)
        if name!='Uploaded_Beauty_Partner':
            add_text(prs.slides[0], 'SYNTHETIC UX FIXTURE · Actual frozen-engine scores · Illustrative data', .8, 6.73, 11.5, .2, size=10)
            prs.save(path)
        record={'fixture':name,'file':filename,'slides':len(prs.slides),**t.summary}
        inventory.append(record)
        print(name,record['slides'],record['business_classification_counts'],flush=True)
    write_json(destination/'fixture_inventory.json',inventory)
    return inventory


if __name__=='__main__':
    build(sys.argv[1] if len(sys.argv)>1 else Path(__file__).parents[1]/'validation/v2_1_examples')
