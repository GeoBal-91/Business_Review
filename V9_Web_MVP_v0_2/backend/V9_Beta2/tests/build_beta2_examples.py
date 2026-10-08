from pathlib import Path
import sys,json,hashlib,copy
root=Path(__file__).resolve().parents[1];sys.path[:0]=[str(root),str(root/'tests')]
import pandas as pd
from pptx import Presentation
from analysis import run_analysis
from config import Settings
from business_classification import classify
from triage_presentation import render
from triage_exports import export_triage
from test_beta2 import analyze_sbu
from test_regression import fixture,run,zero
from data_pipeline import G,U
from presentation_hardening import top_gmv_rows
from presentation_polish import build_top_gmv
out=root/'validation';out.mkdir(exist_ok=True)
records=[]
for name,a in [('JEWELRY',run_analysis(Settings())),('SUPERMARKET_synthetic',analyze_sbu('SUPERMARKET')),('FLOWERS_synthetic',analyze_sbu('FLOWERS',[fixture(**{U:.1})]))]:
 t=classify(a);path=out/f'V9_Beta2_{name}.pptx';render(t,path);export_triage(t,out/f'V9_Beta2_{name}_audit')
 p=Presentation(path);records.append({'file':path.name,'slides':len(p.slides),'source':'provided partner workbook' if name=='JEWELRY' else 'synthetic partner; NEW production benchmark','generated':True})
# Real source deck: replace only Top GMV slide using real displayed GMV/status evidence.
e=json.loads((root/'tests/abc_evidence.json').read_text());p=Presentation(root/'tests/fixtures/abc_reference.pptx')
r=e['display_rows'];a=run([fixture(r[0]['store'],**{G:r[0]['gmv']}),zero(r[1]['store'])]);shown=top_gmv_rows(a)
shown['Commercial Status']=r[0]['commercial'];shown['Operational Status']=r[0]['operational']
old_id=p.slides._sldIdLst[12];build_top_gmv(p,a,shown)
new_id=p.slides._sldIdLst[-1];p.slides._sldIdLst.remove(new_id);p.slides._sldIdLst.insert(12,new_id);p.slides._sldIdLst.remove(old_id);p.part.drop_rel(old_id.rId)
for slide in p.slides:
 for s in slide.shapes:
  if s.has_text_frame:
   for para in s.text_frame.paragraphs:
    for run_ in para.runs:
     run_.text=run_.text.replace('MAIN ACTION KPI','PRIMARY ACTION KPI')
     for old in ['V9 Beta 1.0','V9 Beta 1.1','V9_Beta1_1']:
      run_.text=run_.text.replace(old,'V9_Beta2')
p.core_properties.version='V9_Beta2';path=out/'V9_Beta2_PHARMACY_abc_presentation_replay.pptx';p.save(path)
records.append({'file':path.name,'slides':len(p.slides),'source':'real abc deck; Top GMV slide replay only. Raw input workbook unavailable. Other analytics retained from source deck.','generated':True})
# Empty-state standalone presentation smoke artifact.
p=Presentation();p.slide_width=12191695;p.slide_height=6858000;a=run([zero()]);build_top_gmv(p,a,top_gmv_rows(a));p.core_properties.version='V9_Beta2';p.save(out/'V9_Beta2_Top_GMV_empty_state.pptx')
(root/'reports/deck_generation.json').write_text(json.dumps(records,indent=2))
print(json.dumps(records,indent=2))
