"""Presentation patch acceptance; the Dev6.1 regression suite stays untouched."""
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.colors import to_hex
from pptx import Presentation
from pptx.util import Inches
from test_regression import run, fixture
from test_hardening import disappearance, text_of
from data_pipeline import G, N
from presentation import render, context
from presentation_polish import (
    store_color_mapping, matrix_figure, build_matrix, build_top_gmv,
    top_gmv_concentration, build_guide, current_coverage_incomplete,
    COVERAGE_CVR_NOTE, MATRIX_FOOTNOTE,
)


def blank():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    return prs


def full_deck(result):
    with tempfile.TemporaryDirectory() as tmp:
        return Presentation(render(result, Path(tmp) / 'review.pptx'))


def slide_text(prs, title):
    for slide in prs.slides:
        if any(s.has_text_frame and s.text == title for s in slide.shapes):
            return '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
    raise AssertionError('Missing slide: ' + title)


class StoreMatrix(unittest.TestCase):
    def test_deterministic_color_mapping_independent_of_input_order(self):
        ids = ['North', 'Central', 'East', 'West', 'South']
        self.assertEqual(store_color_mapping(ids), store_color_mapping(ids[::-1]))
        self.assertEqual(store_color_mapping(ids), store_color_mapping(ids + ids))
        a = run([fixture(sid, **{G: 100 + i, '% CVR': .1 + i * .01}) for i, sid in enumerate(ids)])
        b = run([fixture(sid, **{G: 9999 - i, '% CVR': .4 - i * .01}) for i, sid in enumerate(ids)])
        self.assertEqual(store_color_mapping(a.stores['Store ID']), store_color_mapping(b.stores['Store ID']))

    def test_each_plotted_store_gets_distinct_color(self):
        for n in [1, 12, 30, 80]:
            colors = store_color_mapping([str(i) for i in range(n)])
            self.assertEqual(len({x.lower() for x in colors.values()}), n)

    def test_xy_reference_lines_and_no_point_labels(self):
        a = run([fixture('B', **{G: 8000, '% CVR': .09}), fixture('A', **{G: 2000, '% CVR': .31})])
        fig = matrix_figure(a.stores, a.data.benchmark_row)
        try:
            ax = fig.axes[0]
            np.testing.assert_allclose(ax.collections[0].get_offsets(), a.stores[[G, '% CVR']].to_numpy() * [1, 100])
            self.assertEqual(ax.get_xlabel(), 'GMV (€)'); self.assertEqual(ax.get_ylabel(), 'CVR (%)')
            np.testing.assert_allclose(ax.lines[0].get_xdata(), [5000, 5000])
            np.testing.assert_allclose(ax.lines[1].get_ydata(), [20, 20])
            self.assertEqual(len(ax.texts), 0)
            colors = store_color_mapping(a.stores['Store ID'])
            self.assertEqual([to_hex(x).lower() for x in ax.collections[0].get_facecolors()], [colors[sid].lower() for sid in a.stores['Store ID']])
        finally:
            plt.close(fig)

    def test_full_store_legend_readable_across_continuations(self):
        rows = [fixture(f'S{i:02}', **{'Store Name': f'Store {i:02}', 'Store Address Address': f'{100+i} Chavchavadze Avenue, Tbilisi, Georgia'}) for i in range(28)]
        a = run(rows); prs = blank()
        with tempfile.TemporaryDirectory() as tmp:
            build_matrix(prs, a, a.stores, lambda n: str(Path(tmp) / n))
        self.assertGreater(len(prs.slides), 1)
        labels, colors = {}, {}
        for slide in prs.slides:
            self.assertTrue(any(s.has_text_frame and s.text == MATRIX_FOOTNOTE for s in slide.shapes))
            for shape in slide.shapes:
                if shape.name.startswith('Store legend: '):
                    sid = shape.name.split(': ', 1)[1]
                    self.assertGreaterEqual(shape.left / 914400, 8.8)
                    labels[sid] = shape.text
                    self.assertGreaterEqual(shape.text_frame.paragraphs[0].font.size.pt, 10.5)
                elif shape.name.startswith('Store color: '):
                    colors[shape.name.split(': ', 1)[1]] = '#' + str(shape.fill.fore_color.rgb).lower()
        self.assertEqual(set(labels), set(a.stores['Store ID']))
        for row in rows:
            self.assertIn(row['Store Name'], labels[row['Store ID']])
            self.assertIn(row['Store Address Address'], labels[row['Store ID']])
        self.assertEqual(colors, {k: v.lower() for k, v in store_color_mapping(a.stores['Store ID']).items()})


class TopGMV(unittest.TestCase):
    def test_concentration_uses_displayed_gmv_and_total_partner_gmv(self):
        a = run([fixture(f'S{i}', **{G: 100 * i}) for i in range(1, 9)])
        displayed = a.stores.sort_values([G, 'Store ID'], ascending=[False, True], na_position='last').head(6)
        n, share, reason = top_gmv_concentration(a, displayed)
        self.assertEqual(n, 6); self.assertAlmostEqual(share, 3300 / 3600); self.assertFalse(reason)
        prs = blank(); build_top_gmv(prs, a, displayed)
        self.assertIn('Top 6 stores generated 91.7% of total partner GMV in August 2026.', text_of(prs))
        self.assertTrue(any(s.shape_type == 13 for s in prs.slides[0].shapes))

    def test_existing_top_gmv_ranking_and_rows_retained(self):
        a = run([fixture(f'S{i}', **{G: 100 * i}) for i in range(1, 9)])
        text = slide_text(full_deck(a), 'Top GMV Stores')
        names = [line for line in text.splitlines() if line in {f'S{i}' for i in range(1, 9)}]
        self.assertEqual(names, ['S8', 'S7', 'S6', 'S5', 'S4', 'S3'])
        for field in ['Commercial status', 'Operational status', 'GMV']:
            self.assertIn(field, text)

    def test_fewer_than_six_and_single_store(self):
        for n in [1, 2, 5]:
            a = run([fixture(str(i)) for i in range(n)])
            count, share, _ = top_gmv_concentration(a, a.stores)
            self.assertEqual(count, n); self.assertEqual(share, 1)
            prs = blank(); build_top_gmv(prs, a, a.stores)
            noun = 'store' if n == 1 else 'stores'
            self.assertIn(f'Top {n} {noun} generated 100.0%', text_of(prs))

    def test_zero_missing_or_empty_gmv_not_fabricated(self):
        for value in [0, np.nan]:
            a = run([fixture(**{G: value})])
            self.assertIsNone(top_gmv_concentration(a, a.stores)[1])
            prs = blank(); build_top_gmv(prs, a, a.stores)
            self.assertNotIn('generated', text_of(prs))
        a = run(); self.assertEqual(top_gmv_concentration(a, a.stores.iloc[:0])[0], 0)

    def test_partial_coverage_not_claimed_as_total_gmv_share(self):
        a = disappearance()
        count, share, reason = top_gmv_concentration(a, a.stores)
        self.assertEqual(count, 2); self.assertIsNone(share)
        self.assertIn('incomplete store coverage', reason)


class MethodologyGuide(unittest.TestCase):
    def test_dynamic_actual_benchmark_month(self):
        a = run(benchmarks=[fixture(month='July')], max_benchmark_lag_months=1)
        prs = blank(); build_guide(prs, a); text = text_of(prs)
        self.assertIn('2026-07. SBU peer composition is not available.', text)
        for title in ['Severity', 'Impact', 'Priority', 'Exposure components', 'Selection threshold', 'Missing evidence', 'Benchmark month', 'Uptime', 'Key takeaway', 'Remember']:
            self.assertIn(title, text)

    def test_missing_benchmark_is_explicit(self):
        a = run(benchmarks=[fixture(month='July')])
        prs = blank(); build_guide(prs, a)
        self.assertIn('Unavailable. SBU peer composition is not available.', text_of(prs))

    def test_guide_retains_flowers_uptime_caveat_and_challenger_meaning(self):
        from dataclasses import replace
        from analysis import analyze
        from data_pipeline import load_frames
        from config import Settings
        from datetime import date
        settings = Settings(partner_sbu='FLOWERS', as_of_date=date(2026, 9, 7))
        a = analyze(load_frames(pd.DataFrame([fixture()]), pd.DataFrame([fixture(**{'Store Sub Business Unit': 'FLOWERS'})]), settings))
        prs = blank(); build_guide(prs, a); text = text_of(prs)
        for phrase in ['95% target', 'healthy history', 'FLOWERS uptime: hygiene only, excluded from business scoring', '5-point selection threshold is provisional', 'not a guaranteed recovery forecast', 'separate review list']:
            self.assertIn(phrase, text)
        b = run(priority_model='simple_product'); prs = blank(); build_guide(prs, b)
        self.assertIn('Challenger: Priority = Severity × Impact / 100', text_of(prs))


class PartnerCVRCoverage(unittest.TestCase):
    def test_incomplete_current_coverage_suppresses_direction_keeps_raw_rates(self):
        a = disappearance(); text = slide_text(full_deck(a), 'Executive Summary')
        self.assertIn(COVERAGE_CVR_NOTE, text)
        self.assertIn('Expected 2 · Present 1 · Missing 1', text)
        self.assertIn('Partner CVR\n20.0%', text); self.assertIn('SBU CVR: 20.0%', text)
        self.assertNotIn('Partner CVR is', text)
        self.assertTrue(context(a)['current_partner_coverage_incomplete'])

    def test_complete_current_coverage_preserves_direction(self):
        a = run([fixture(**{'% CVR': .1})]); text = slide_text(full_deck(a), 'Executive Summary')
        self.assertIn('Partner CVR is 10.0 pp worse', text)
        self.assertNotIn(COVERAGE_CVR_NOTE, text)

    def test_historical_gap_does_not_suppress_complete_current_comparison(self):
        a = run([fixture('A', 'June'), fixture('B', 'June'), fixture('A', 'July'), fixture('A'), fixture('B')])
        self.assertFalse(current_coverage_incomplete(a))
        text = slide_text(full_deck(a), 'Executive Summary')
        self.assertIn('Partner CVR is', text)
        self.assertNotIn(COVERAGE_CVR_NOTE, text)
        self.assertIn('MoM: Not comparable', text)


class FrozenDev61(unittest.TestCase):
    def test_frozen_source_files_byte_identical(self):
        root = Path(__file__).parents[1]
        baseline = json.loads((root / 'tests/dev6_1_presentation_baseline.json').read_text())
        for name, digest in baseline['frozen_files'].items():
            self.assertEqual(hashlib.sha256((root / name).read_bytes()).hexdigest(), digest, name)

    def test_rendering_does_not_modify_numerical_records_or_actions(self):
        a = disappearance()
        frames = {name: getattr(a, name).copy(deep=True) for name in ['stores', 'metrics', 'opportunities']}
        monthly = a.data.monthly.copy(deep=True)
        actions = json.dumps(a.actions, sort_keys=True)
        full_deck(a)
        for name, before in frames.items():
            pd.testing.assert_frame_equal(before, getattr(a, name), check_exact=True)
        pd.testing.assert_frame_equal(monthly, a.data.monthly, check_exact=True)
        self.assertEqual(actions, json.dumps(a.actions, sort_keys=True))

    def test_recommendations_content_and_layout_xml_unchanged(self):
        baseline = json.loads((Path(__file__).parent / 'dev6_1_presentation_baseline.json').read_text())
        prs = full_deck(run(baseline['recommendation_rows']))
        actual = [hashlib.sha256(s._element.xml.encode()).hexdigest() for s in prs.slides
                  if any(x.has_text_frame and x.text.startswith('Priority Opportunities & Actions') for x in s.shapes)]
        self.assertEqual(actual, baseline['recommendation_slide_xml_sha256'])
