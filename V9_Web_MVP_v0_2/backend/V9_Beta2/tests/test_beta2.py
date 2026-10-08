"""Beta2 compatibility and presentation regressions. No score recalibration."""
import unittest
from pathlib import Path
from dataclasses import replace
from datetime import date
import json
import numpy as np
import pandas as pd
from pandas.testing import assert_frame_equal
from pptx import Presentation
from pptx.util import Inches
from config import Settings
from data_pipeline import load_frames, DataValidationError, G, U, D, N, T, P, C
from analysis import analyze
from business_classification import classify
from sbu_catalog import supported_sbus
from presentation_hardening import top_gmv_rows
from presentation_polish import build_top_gmv, top_gmv_concentration
from test_regression import run, fixture, zero

ROOT = Path(__file__).resolve().parents[1]
BENCHMARK = pd.read_excel(ROOT/'subbusiness_unit_data.xlsx')


def analyze_sbu(sbu, rows=None):
    settings = Settings(partner_sbu=sbu, as_of_date=date(2026,9,14))
    return analyze(load_frames(pd.DataFrame(rows or [fixture()]), BENCHMARK, settings))


class Catalog(unittest.TestCase):
    def test_schema_and_count(self):
        self.assertEqual(BENCHMARK.shape, (516,29))
        self.assertEqual(len(supported_sbus()),43)
    def test_unknown_fails_without_fallback(self):
        with self.assertRaises(DataValidationError) as cm:
            analyze_sbu('NOT_IN_BENCHMARK')
        self.assertIn('No rows for the requested partner/SBU',str(cm.exception))
    def test_discovery_normalization_and_no_static_enum(self):
        frame=pd.DataFrame({'Store Sub Business Unit':[' experimental ','FLOWERS',None,'', 'EXPERIMENTAL']})
        self.assertEqual(supported_sbus(frame),('EXPERIMENTAL','FLOWERS'))
    def test_bad_catalog_schema(self):
        for frame in [pd.DataFrame(),pd.DataFrame({'Store Sub Business Unit':[None,' ']})]:
            with self.assertRaises(DataValidationError):supported_sbus(frame)


def sbu_case(sbu):
    def test(self):
        a=analyze_sbu(sbu)
        selected=a.data.sbu_raw if hasattr(a.data,'sbu_raw') else None
        expected=BENCHMARK[BENCHMARK['Store Sub Business Unit']==sbu]
        self.assertEqual(a.settings.partner_sbu,sbu)
        self.assertEqual(a.data.benchmark_period,pd.Timestamp('2026-08-01'))
        self.assertEqual(len(a.data.selected_sbu),12)
        # Exact same-month source values, never a different SBU or a fallback month.
        row=expected[expected.Month=='August'].iloc[0]
        self.assertAlmostEqual(a.data.benchmark_row[G],row[G])
        if pd.isna(row['% CVR']):
            self.assertTrue(pd.isna(a.data.benchmark_row['% CVR']))
        else:
            self.assertAlmostEqual(a.data.benchmark_row['% CVR'],row['% CVR'])
        self.assertEqual(len(classify(a).stores),1)
    return test
for sbu in supported_sbus(BENCHMARK):
    setattr(Catalog,'test_resolves_'+sbu,sbu_case(sbu))


class TopGMV(unittest.TestCase):
    def check(self,a,expected):
        before=a.stores.copy(deep=True)
        shown=top_gmv_rows(a)
        self.assertEqual(shown['Store ID'].tolist(),expected)
        assert_frame_equal(before,a.stores,check_exact=True)
        prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
        build_top_gmv(prs,a,shown)
        return '\n'.join(s.text for s in prs.slides[0].shapes if s.has_text_frame)
    def test_one_positive_one_zero_no_activity(self):
        a=run([fixture('Positive'),zero('Zero')]);text=self.check(a,['Positive'])
        self.assertIn('Top 1 store generated 100.0%',text)
        self.assertNotIn('Zero',text)
    def test_multiple_positives_zero_statuses(self):
        a=run([fixture('B',**{G:3000.}),fixture('A'),zero('No activity'),zero('Closed',**{'Lifecycle Status':'closed'}),fixture('Data review',**{G:0.,U:np.nan})])
        self.check(a,['B','A'])
    def test_all_zero_empty_state(self):
        text=self.check(run([zero(),zero('B')]),[])
        self.assertIn('No positive-GMV stores available for the current period.',text)
    def test_one_positive_only(self):self.check(run(),['A'])
    def test_more_than_six_positive(self):
        a=run([fixture(str(i),**{G:float(i)}) for i in range(1,9)])
        self.check(a,['8','7','6','5','4','3'])
        self.assertAlmostEqual(top_gmv_concentration(a,top_gmv_rows(a))[1],33/36)
    def test_unknown_missing_current_positive_excluded(self):
        a=run([fixture('A'),fixture('B'),fixture('C',month='July')])
        a.stores[G]=a.stores[G].astype(float)
        a.stores.loc[a.stores['Store ID']=='B','Store Status']='UNKNOWN'
        a.stores.loc[a.stores['Store ID']=='C',G]=9999.
        self.check(a,['A'])
        a.stores.loc[a.stores['Store ID']=='B','Store Status']='MISSING_CURRENT'
        self.check(a,['A'])
    def test_nonfinite_missing_negative_and_outside_population(self):
        a=run([fixture(str(i),**{G:100.}) for i in range(6)])
        a.stores.loc[0,G]=np.inf;a.stores.loc[1,G]=-np.inf;a.stores.loc[2,G]=np.nan;a.stores.loc[3,G]=-1.
        a.stores.loc[4,'Store ID']='not-in-current-population'
        self.check(a,['5'])
    def test_reported_denominator_not_eligible_sum(self):
        a=run([fixture('A'),fixture('B')])
        a.stores.loc[a.stores['Store ID']=='B','Store Status']='UNKNOWN'
        shown=top_gmv_rows(a)
        self.assertEqual(top_gmv_concentration(a,shown)[1],.5)
    def test_abc_real_presentation_values(self):
        evidence=json.loads((ROOT/'tests/abc_evidence.json').read_text())
        rows=evidence['display_rows']
        a=run([fixture(rows[0]['store'],**{G:rows[0]['gmv']}),zero(rows[1]['store'])])
        text=self.check(a,[rows[0]['store']])
        self.assertIn('€11,211',text)
        self.assertIn('Top 1 store generated 100.0%',text)
        self.assertNotIn(rows[1]['store'],text)


class UptimeAndFlowers(unittest.TestCase):
    def test_zero_and_blank_partial_coverage(self):
        a=run([fixture('Zero',**{U:0.}),fixture('Missing',**{U:np.nan})])
        row=a.data.monthly.iloc[0]
        self.assertEqual(row[U],0.);self.assertEqual(row['Uptime Known Store Count'],1)
        self.assertEqual(row['Operating Store Count'],2);self.assertEqual(row['Uptime Value Coverage'],.5)
        self.assertTrue(pd.isna(a.stores.set_index('Store ID').loc['Missing',U]))
    def test_all_blank_not_assessed(self):
        a=run([fixture(**{U:np.nan})]);self.assertTrue(pd.isna(a.data.monthly.iloc[0][U]))
        self.assertEqual(a.data.monthly.iloc[0]['Uptime Known Store Count'],0)
    def test_contradictory_activity_not_closed(self):
        a=run([fixture(**{'Lifecycle Status':'closed'})])
        self.assertEqual(a.stores.iloc[0]['Store Status'],'Active')
        self.assertEqual(a.stores.iloc[0]['Data Quality'],'Conflict')
    def test_flowers_uptime_never_commercial(self):
        results=[analyze_sbu('FLOWERS',[fixture(**{U:u})]) for u in [0.,.1,.95,np.nan]]
        base=results[0]
        for a in results:
            uptime=a.metrics[a.metrics.metric=='Uptime'].iloc[0]
            self.assertFalse(uptime['eligible'])
            for key in ['Opportunity Score','Business Impact Score','Main Driver']:
                self.assertEqual(a.stores.iloc[0][key],base.stores.iloc[0][key])
            ca,cb=classify(a),classify(base)
            for key in ['Business Classification','Classification Reason']:
                if key in ca.stores:self.assertEqual(ca.stores.iloc[0][key],cb.stores.iloc[0][key])
            self.assertFalse(any(x['Metric']=='Uptime' for x in a.actions))


class FrozenBaseline(unittest.TestCase):
    def test_byte_identical_analytical_modules(self):
        import hashlib
        expected=json.loads((ROOT/'tests/beta2_frozen_modules.json').read_text())
        for name,digest in expected.items():
            self.assertEqual(hashlib.sha256((ROOT/name).read_bytes()).hexdigest(),digest,name)
    def test_frozen_output_fixtures_exact(self):
        from beta_fixtures import flowers
        cases={
            'default':run(),
            'partial_uptime':run([fixture(),fixture('B',**{U:float('nan')})]),
            'zero_uptime':run([fixture(**{U:0.}),fixture('B',**{U:float('nan')})]),
            'missing_current':run([fixture(),fixture('B',month='July')]),
            'zero_activity':run([zero()]),'positive_zero':run([fixture(),zero('B')]),
            'closed_conflict':run([fixture(**{'Lifecycle Status':'closed'})]),
            'golden_80':run(json.loads((ROOT/'tests/dev6_golden.json').read_text())['rows']),
            'flowers':flowers()}
        expected=json.loads((ROOT/'tests/beta2_frozen_outputs.json').read_text())
        for name,a in cases.items():
            with self.subTest(case=name):
                payload={k:getattr(a,k).to_json(orient='split',double_precision=15,date_format='iso') for k in ['stores','metrics','opportunities']}
                payload.update(actions=a.actions,reviews=a.reviews,classification=classify(a).stores.to_json(orient='split',double_precision=15,date_format='iso'))
                self.assertEqual(json.loads(json.dumps(payload,default=str)),expected[name])
