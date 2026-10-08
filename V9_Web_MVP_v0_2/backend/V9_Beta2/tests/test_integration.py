import unittest
import tempfile
import subprocess
import sys
import json
from pathlib import Path
from unittest.mock import patch
import pandas as pd
from pptx import Presentation
from test_regression import run,fixture,zero
from main import export_analysis
from presentation import render,context
from data_pipeline import U,N,G


class Integration(unittest.TestCase):
    def test_critical_flags_do_not_add_score(self):
        rows=[fixture(**{U:.4,'% CVR':.01})]
        expected=run(rows).stores.iloc[0]['Opportunity Score']
        with patch('scoring._raw_flag',return_value=''):
            self.assertEqual(run(rows).stores.iloc[0]['Opportunity Score'],expected)
    def test_import_does_not_execute_pipeline(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=str(Path(__file__).resolve().parents[1])
            code=f"import sys; sys.path.insert(0,{root!r}); import config,data_pipeline,store_status,scoring,analysis,main,presentation"
            completed=subprocess.run([sys.executable,'-c',code],cwd=tmp,capture_output=True,text=True)
            self.assertEqual(completed.returncode,0,completed.stderr)
            self.assertEqual(list(Path(tmp).iterdir()),[])
    def test_full_precision_exports_and_strict_json(self):
        a=run([fixture(**{'% CVR':.01}),fixture('Missing',**{U:float('nan')})])
        with tempfile.TemporaryDirectory() as tmp:
            export_analysis(a,tmp)
            report=json.loads((Path(tmp)/'run_audit.json').read_text(),parse_constant=lambda x:self.fail(x))
            self.assertEqual(report['incomplete_assessments'],1)
            records=json.loads((Path(tmp)/'metric_records.json').read_text(),parse_constant=lambda x:self.fail(x))
            self.assertTrue(any(r['severity'] is None for r in records))
            exported=pd.read_csv(Path(tmp)/'stores.csv')
            self.assertAlmostEqual(exported.iloc[0]['Opportunity Score'],a.stores.iloc[0]['Opportunity Score'],places=12)
    def test_stress_deck_contains_every_action_and_unknown_store(self):
        rows=[fixture('Opportunity '+str(i),**{N:10000,'% CVR':0}) for i in range(7)]
        rows += [fixture('Missing',month='July'),zero('Closed',**{'Lifecycle Status':'closed'})]
        a=run(rows)
        self.assertTrue(any(i.code == "incomparable_share" for i in a.data.issues))
        self.assertTrue(pd.isna(context(a)["latest_share"]))
        with tempfile.TemporaryDirectory() as tmp:
            path=render(a,Path(tmp)/'stress.pptx')
            deck=Presentation(path)
            text='\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
            self.assertIn('Top GMV Stores',text)
            self.assertIn('UNKNOWN',text.upper()); self.assertIn('Not assessed',text)
            self.assertNotRegex(text.lower(),r'\bnan\b')
            for action in a.actions: self.assertIn(action['Action'],text)
            self.assertEqual(sum(any(s.has_text_frame and s.text.startswith('Priority Opportunities & Actions') for s in slide.shapes) for slide in deck.slides),2)
    def test_stale_no_common_period_deck(self):
        a=run(benchmarks=[fixture(month='July')])
        with tempfile.TemporaryDirectory() as tmp:
            path=render(a,Path(tmp)/'stale.pptx')
            text='\n'.join(s.text for slide in Presentation(path).slides for s in slide.shapes if s.has_text_frame)
            self.assertIn('benchmark unavailable',text)
            self.assertIn('Not assessed',text)
    def test_empty_action_exports_have_schema(self):
        a=run()
        with tempfile.TemporaryDirectory() as tmp:
            export_analysis(a,tmp)
            self.assertEqual(len(pd.read_csv(Path(tmp)/'actions.csv')),0)
            self.assertIn('Action',pd.read_csv(Path(tmp)/'reviews.csv').columns)
            self.assertEqual(len(pd.read_csv(Path(tmp)/'store_history.csv')),1)

    def test_cohort_metadata_matches_rendering(self):
        a=run([fixture(),fixture('Missing',month='July')])
        text=context(a)['uptime_coverage_note']
        self.assertIn('1/1 operating store',text)
        self.assertIn('Missing current store records: 1',text)

if __name__=='__main__': unittest.main()
