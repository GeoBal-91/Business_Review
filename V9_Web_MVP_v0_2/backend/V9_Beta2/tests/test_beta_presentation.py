"""Beta partner-facing flow and exact AnalysisResult lineage acceptance."""
import json
import tempfile
import unittest
from pathlib import Path
import pandas as pd
import numpy as np
from pptx import Presentation
from test_regression import run, fixture, zero
from test_presentation_patch import blank, full_deck, slide_text
from test_hardening import text_of, diagnostic_deck
from test_presentation_navigation import titles
from beta_fixtures import large_partner, flowers, one_opportunity
from presentation_hardening_fixtures import recommendation_rich, current_partial, historical_gap
from beta_presentation import (build_opportunities, build_methodology, overall_action_focus,
                               overview_pages, overview_rows, overview_cells, overview_status,
                               build_overview, plural, ACTION_TITLE, METHOD_TITLE, VERSION)
from presentation_hardening import top_gmv_rows, COVERAGE_BENCHMARK_NOTE
from formatting import format_score, format_currency, format_percent
from data_pipeline import G, U


def action_deck(a):
    prs = blank(); build_opportunities(prs, a); return prs


def slide_body(slide):
    return '\n'.join(s.text for s in slide.shapes if s.has_text_frame)


class BetaFlow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.a = recommendation_rich(inactive_count=2)
        cls.prs = full_deck(cls.a)

    def test_context_before_methodology_and_methodology_before_actions(self):
        names = titles(self.prs)
        expected = ['Partner Business Review', 'PHARMACY_OTC Business Development',
                    'PHARMACY_OTC Performance Trend', 'Partner vs PHARMACY_OTC',
                    'Growth vs PHARMACY_OTC', 'Partner Share of PHARMACY_OTC',
                    'Executive Summary', 'Partner Performance Evolution', 'Growth Drivers',
                    'Conversion Performance', 'Operational Health']
        self.assertEqual(names[:11], expected)
        index = names.index(METHOD_TITLE)
        self.assertEqual(names[index-1], 'Top GMV Stores')
        self.assertTrue(names[index+1].startswith(ACTION_TITLE))

    def test_removed_standalone_sections_absent(self):
        for name in ['Biggest Opportunities', 'Recommendations', 'How to Read This Review']:
            self.assertNotIn(name, titles(self.prs))

    def test_action_plan_and_review_remain_distinct(self):
        names = titles(self.prs)
        plan = names.index('Action Plan')
        review = next(i for i, x in enumerate(names) if x.startswith('Operational & Data Review'))
        self.assertLess(plan, review); self.assertLess(review, names.index('Discussion'))
        body = slide_text(self.prs, 'Action Plan')
        for action in self.a.actions:
            self.assertIn(action['Store']+' / '+action['Metric'], body)
            self.assertIn(action['Action'], body)
        self.assertIn('To agree', body)

    def test_thank_you_then_appendix_then_overview_then_all_diagnostics(self):
        names = titles(self.prs); end = names.index('Thank You')
        self.assertEqual(names[end+1], 'APPENDIX — Store Diagnostics')
        self.assertEqual(names[end+2], 'Store Overview')
        self.assertEqual(names[end+3:], ['Store Diagnostics']*len(self.a.stores))
        original = diagnostic_deck(self.a)
        self.assertEqual([s._element.xml for s in list(self.prs.slides)[end+3:]],
                         [s._element.xml for s in original.slides])

    def test_beta_visible_version_no_stale_release_labels(self):
        body = text_of(self.prs)
        self.assertIn(VERSION, body)
        self.assertNotRegex(body, r'(?i)dev6[._]1[._]2')

    def test_scores_remain_in_diagnostics_not_merged_main(self):
        main = text_of(action_deck(self.a))
        self.assertNotIn('Severity', main); self.assertNotIn('Impact', main)
        diagnostics = diagnostic_deck(self.a)
        for slide, (_, row) in zip(diagnostics.slides, self.a.stores.sort_values('Store ID').iterrows()):
            body = slide_body(slide)
            for label, col in [('Severity','Performance Severity Score'), ('Impact','Business Impact Score')]:
                self.assertIn(label+'\n'+format_score(row[col]), body)

    def test_render_does_not_mutate_any_analysis_payload(self):
        a = current_partial()
        before = {k:getattr(a,k).copy(deep=True) for k in ['stores','metrics','opportunities']}
        records = json.dumps([a.actions,a.reviews,a.summary], default=str, sort_keys=True)
        full_deck(a)
        for k, frame in before.items(): pd.testing.assert_frame_equal(frame,getattr(a,k),check_exact=True)
        self.assertEqual(records,json.dumps([a.actions,a.reviews,a.summary],default=str,sort_keys=True))


class BetaOpportunities(unittest.TestCase):
    def assert_cards(self, a, groups):
        prs = action_deck(a)
        shown = [[s.name.split(': ',1)[1] for s in slide.shapes if s.name.startswith('Opportunity store: ')] for slide in prs.slides]
        self.assertEqual([len(x) for x in shown], groups)
        self.assertEqual([sid for group in shown for sid in group], a.opportunities['Store ID'].tolist())
        body = text_of(prs)
        for rank, action in zip(a.opportunities['Rank'], a.actions):
            self.assertIn(f"#{rank} · {action['Store']}", body)
            for field in ['Issue','Action','Intended Effect']:
                self.assertEqual(body.count(action[field]), sum(other[field] == action[field] for other in a.actions))
        return prs

    def test_zero_honest_empty_state_no_focus_metrics(self):
        a=run();prs=self.assert_cards(a,[0]);body=text_of(prs)
        self.assertIn('No material business opportunities',body)
        self.assertIn('Review operational/data observations separately.',body)
        for name in ['Primary Action KPI','Stores Requiring Action','Priority Range','Critical / Review Items']:
            self.assertNotIn(name,body)

    def test_one_opportunity(self):
        prs=self.assert_cards(one_opportunity(),[1])
        self.assertIn('1 store',text_of(prs));self.assertNotIn('1 stores',text_of(prs))

    def test_four_opportunities_one_page(self):
        self.assert_cards(recommendation_rich(store_count=4),[4])

    def test_five_opportunities_three_plus_two(self):
        self.assert_cards(recommendation_rich(store_count=5),[3,2])

    def test_six_opportunities_three_plus_three(self):
        self.assert_cards(recommendation_rich(),[3,3])

    def test_page_titles_and_overall_focus_repeat_consistently(self):
        a=recommendation_rich();prs=action_deck(a)
        self.assertEqual(titles(prs),[ACTION_TITLE+' · 1/2',ACTION_TITLE+' · 2/2'])
        focus=overall_action_focus(a)
        for slide in prs.slides:
            body=slide_body(slide)
            for expected in ['OVERALL ACTION FOCUS','Full selected opportunity set',
                             'Primary Action KPI\n'+a.actions[0]['Metric'], 'Stores Requiring Action\n6 stores',
                             'Highest '+format_score(a.actions[0]['Priority'],5),
                             'Lowest '+format_score(a.actions[-1]['Priority'],5),
                             plural(focus['review_count'],'item')+' for these stores', 'Separate from business Priority']:
                self.assertIn(expected,body)
            self.assertNotRegex(body,r'\b(high|low|Low|Medium|High)\b')

    def test_focus_uses_distinct_selected_stores_and_excludes_unrelated_review_items(self):
        a=recommendation_rich(inactive_count=2)
        focus=overall_action_focus(a);ids={r['Store ID'] for r in a.actions}
        self.assertEqual(focus['store_count'],len(ids))
        self.assertEqual(focus['review_count'],sum(r['Store ID'] in ids for r in a.reviews))
        self.assertLess(focus['review_count'],len(a.reviews))


class BetaMethodology(unittest.TestCase):
    def test_partner_facing_paired_model_threshold_and_qualifications(self):
        a=run();prs=blank();build_methodology(prs,a);body=text_of(prs)
        for phrase in ['SEVERITY — 0–100','IMPACT — 0–100','PRIORITY — 0–100',
                       'KPI Severity × KPI Impact / 100','dominant-KPI protection',
                       'strongest paired business case','Priority ≥ 5.0',
                       'The 5.0 threshold is provisional and will be calibrated during Beta validation.',
                       'Higher Priority = stronger case for action','incomplete assessment, not zero',
                       'Critical observations are reviewed separately','95% target','healthy history',
                       'not a recovery promise','Not a financial forecast']:
            self.assertIn(phrase,body)
        self.assertNotRegex(body,r'(Severity|Impact|Priority)\s+\d+(\.\d+)?%')

    def test_dynamic_benchmark_period_and_peer_note(self):
        a=run(benchmarks=[fixture(month='July')],max_benchmark_lag_months=1)
        prs=blank();build_methodology(prs,a)
        self.assertIn('Benchmark: 2026-07. SBU peer composition is not available.',text_of(prs))

    def test_missing_benchmark_and_challenger_qualified(self):
        a=run(benchmarks=[fixture(month='July')],priority_model='simple_product')
        prs=blank();build_methodology(prs,a);body=text_of(prs)
        self.assertIn('Benchmark: Unavailable',body)
        self.assertIn('EXPERIMENTAL / CHALLENGER MODEL',body)
        self.assertIn('Store Severity × Store Impact / 100',body)

    def test_flowers_hygiene_only_preserved(self):
        a=flowers();prs=blank();build_methodology(prs,a);body=text_of(prs)
        self.assertIn('FLOWERS: hygiene only; excluded from business Priority.',body)
        self.assertTrue(all(action['Metric']!='Uptime' for action in a.actions))
        self.assertTrue(any(r['Type']=='Opening-hours hygiene' for r in a.reviews))
        self.assertTrue((a.metrics.loc[a.metrics.metric=='Uptime','paired_priority']==0).all())


class BetaOverview(unittest.TestCase):
    def test_fewer_than_five_omitted(self):
        for n in [1,4]:
            a=run([fixture(str(i)) for i in range(n)]);prs=blank();build_overview(prs,a)
            self.assertEqual(len(prs.slides),0)

    def test_five_present(self):
        a=run([fixture(str(i)) for i in range(5)]);prs=blank();build_overview(prs,a)
        self.assertEqual(titles(prs),['Store Overview'])
        self.assertEqual(len(overview_pages(a)[0]),5)

    def test_twelve_rows_one_page_at_readable_font(self):
        a=run([fixture(str(i)) for i in range(12)]);prs=blank();build_overview(prs,a)
        self.assertEqual(len(prs.slides),1)
        names=[s for s in prs.slides[0].shapes if s.name.startswith('Overview store: ')]
        self.assertEqual(len(names),12)
        self.assertTrue(all(s.text_frame.paragraphs[0].font.size.pt>=10.5 for s in names))

    def test_thirteen_paginated(self):
        a=run([fixture(str(i)) for i in range(13)]);prs=blank();build_overview(prs,a)
        self.assertEqual(titles(prs),['Store Overview · 1/2','Store Overview · 2/2'])
        self.assertEqual([len(p) for p in overview_pages(a)],[12,1])

    def test_twentyfive_through_thirty_no_stores_lost(self):
        for n in [25,26,27,28,29,30]:
            a=run([fixture(str(i)) for i in range(n)])
            pages=overview_pages(a);ids=[row['Store ID'] for group in pages for row,_,_ in group]
            self.assertEqual(set(ids),set(a.stores['Store ID']));self.assertEqual(len(ids),n)
            self.assertEqual(len(pages),3);self.assertTrue(all(len(p)<=12 for p in pages))

    def test_cells_match_current_analytical_rows_and_existing_states(self):
        a=large_partner(30);prs=blank();build_overview(prs,a)
        shapes={s.name:s for slide in prs.slides for s in slide.shapes}
        for _,r in a.stores.iterrows():
            names=['Overview store: ']+[f'Overview column {j}: ' for j in range(1,8)]
            actual=[shapes[name+str(r['Store ID'])].text for name in names]
            expected=[r['Store Address Address'],r['Operational Status'] if r['Store Status']=='Active' else r['Store Status'],
                      format_currency(r[G]),format_percent(r['% CVR']),format_percent(r['% PNA Orders']),
                      format_percent(r['% Cancels DTP']),format_percent(r[U]),format_score(r['Opportunity Score'],5)]
            self.assertEqual(actual,expected)
        body=text_of(prs)
        for term in ['Unknown','Closed','No Activity','Data review','Not assessed']:self.assertIn(term,body)

    def test_display_order_no_analysis_mutation(self):
        a=large_partner(18);before=a.stores.copy(deep=True);opps=a.opportunities.copy(deep=True)
        shown=overview_rows(a)
        self.assertEqual(shown.head(len(opps))['Store ID'].tolist(),opps['Store ID'].tolist())
        pd.testing.assert_frame_equal(a.stores,before,check_exact=True)
        pd.testing.assert_frame_equal(a.opportunities,opps,check_exact=True)

    def test_missing_conversion_keeps_operational_status_separate(self):
        a=run([fixture(**{'% CVR':np.nan})]);row=a.stores.iloc[0]
        self.assertEqual(overview_status(row),row['Operational Status'])
        self.assertEqual(overview_cells(row,5)[3],'Not assessed')
        self.assertEqual(overview_cells(row,5)[7],'Not assessed')

    def test_long_names_get_more_room_without_truncation(self):
        label='Long store address '+('Additional location detail '*8)
        a=run([fixture(str(i),**{'Store Address Address':label+str(i)}) for i in range(13)])
        pages=overview_pages(a);self.assertGreater(len(pages),2)
        self.assertEqual(sum(len(p) for p in pages),13)
        self.assertTrue(all(cells[0]==row['Store Address Address'] for group in pages for row,cells,_ in group))

    def test_overview_store_links_target_unchanged_diagnostics(self):
        a=recommendation_rich();prs=full_deck(a)
        diagnostic_slides=[s for s in prs.slides if slide_body(s).startswith('Store Diagnostics\n')]
        targets=dict(zip(a.stores.sort_values('Store ID')['Store ID'],diagnostic_slides))
        links=[s for slide in prs.slides for s in slide.shapes if s.name.startswith('Overview store: ')]
        self.assertEqual(len(links),len(a.stores))
        for shape in links:self.assertEqual(shape.click_action.target_slide,targets[shape.name.split(': ',1)[1]])


class PreservedQualifications(unittest.TestCase):
    def test_current_partial_suppresses_all_directional_comparisons(self):
        prs=full_deck(current_partial())
        for title in ['Executive Summary','Partner vs PHARMACY_OTC','Operational Health']:
            body=slide_text(prs,title)
            self.assertIn(COVERAGE_BENCHMARK_NOTE,body)
            self.assertNotRegex(body.lower(),r'\b(better|worse|above|below)\b|in line')

    def test_complete_current_retains_direction_despite_historical_gap(self):
        for a in [run(),historical_gap()]:
            prs=full_deck(a)
            for title in ['Executive Summary','Partner vs PHARMACY_OTC','Operational Health']:
                body=slide_text(prs,title);self.assertIn('In line with',body);self.assertNotIn(COVERAGE_BENCHMARK_NOTE,body)

    def test_top_gmv_missing_current_and_observed_zero_excluded(self):
        self.assertEqual(top_gmv_rows(current_partial())['Store ID'].tolist(),['Small'])
        a=run([fixture('Observed zero',**{G:0}),fixture('Observed positive')])
        self.assertEqual(top_gmv_rows(a)['Store ID'].tolist(),['Observed positive'])

    def test_singular_plural_main_deck_no_priority_bands(self):
        a=one_opportunity();body=text_of(full_deck(a))
        self.assertIn('1 store · Reference section',body)
        self.assertNotRegex(body,r'\b1 stores\b|\b1 items\b|\b1 opportunities\b')
        self.assertNotRegex(body,r'\b(Low|Medium|High) Priority\b')
