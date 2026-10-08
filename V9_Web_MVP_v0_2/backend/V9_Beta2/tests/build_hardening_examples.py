"""Generate Dev6.1.2 synthetic UX examples; no business input files changed."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parents[1]))
from pptx import Presentation
from ppt_core import add_text
from presentation import render
from main import export_analysis
from presentation_hardening import build_recommendations
from pptx.util import Inches
from presentation_hardening_fixtures import recommendation_rich, current_partial, no_recommendations


def examples(output):
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    fixtures=[('Complete_Coverage_Recommendations',recommendation_rich(inactive_count=2)),
              ('Partial_Current_Coverage',current_partial()),
              ('Zero_Recommendations',no_recommendations()),
              ('Four_Recommendations',recommendation_rich(store_count=4))]
    outputs=[]
    for name,result in fixtures:
        ppt=output/(name+'.pptx')
        render(result,ppt)
        prs=Presentation(ppt)
        add_text(prs.slides[0],'ILLUSTRATIVE UX FIXTURE — SYNTHETIC DATA',.8,6.8,11.5,.3,size=12,bold=True)
        prs.save(ppt)
        export_analysis(result,output/(name+'_audit'))
        outputs.append(ppt)
        print(f'{ppt.name}: {len(prs.slides)} slides; {len(result.actions)} business actions; {len(result.reviews)} separate review items',flush=True)
    result = recommendation_rich(store_count=5)
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    build_recommendations(prs, result)
    for slide in prs.slides:
        add_text(slide, 'ILLUSTRATIVE UX FIXTURE — SYNTHETIC DATA', .67, 1.22, 12, .23, size=10, bold=True)
    ppt = output / 'Five_Recommendations_Pagination.pptx'
    prs.save(ppt)
    export_analysis(result, output / 'Five_Recommendations_Pagination_audit')
    outputs.append(ppt)
    print(f'{ppt.name}: {len(prs.slides)} slides; 5 business actions (3+2)', flush=True)
    return outputs


if __name__=='__main__':examples(sys.argv[1])
