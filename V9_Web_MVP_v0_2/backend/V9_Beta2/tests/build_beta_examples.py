"""Generate all requested Beta visual fixtures and normal/additive audit exports."""
import sys
import json
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parents[1]))
from pptx import Presentation
from presentation import render
from ppt_core import add_text
from main import export_analysis
from beta_fixtures import visual_fixtures


def examples(folder):
    folder = Path(folder); folder.mkdir(parents=True, exist_ok=True)
    records = []
    for name, result in visual_fixtures():
        path = folder / ('V9_Beta1_0_' + name + '.pptx')
        render(result, path)
        prs = Presentation(path)
        add_text(prs.slides[0], 'ILLUSTRATIVE UX FIXTURE — SYNTHETIC DATA', .8, 6.70, 11.5, .3, size=12, bold=True)
        prs.save(path)
        export_analysis(result, folder / (path.stem + '_audit'))
        entry = {'fixture': name, 'file': path.name, 'slides': len(prs.slides), 'stores': len(result.stores),
                 'opportunities': len(result.actions), 'reviews': len(result.reviews)}
        records.append(entry); print(entry, flush=True)
    (folder / 'fixture_inventory.json').write_text(json.dumps(records, indent=2) + '\n')
    return records


if __name__ == '__main__':
    examples(sys.argv[1])
