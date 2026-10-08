from pathlib import Path
import subprocess,os
from pptx import Presentation
root=Path(__file__).resolve().parent
bin=Path('/Users/giorgibalarjishvili/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override');out=root/'previews';out.mkdir(exist_ok=True)
for pptx in sorted((root/'decks').glob('*.pptx')):
 subprocess.run([str(bin/'soffice'),'--headless','--convert-to','pdf','--outdir',str(out),str(pptx)],env=dict(os.environ,FONTCONFIG_FILE=str(root/'fonts.conf')),check=True)
 for n,s in enumerate(Presentation(pptx).slides,1):
  title=next((x.text for x in s.shapes if x.has_text_frame and x.text.strip()),'')
  if title=='Thank You' or ' Share of ' in title:
   subprocess.run([str(bin/'pdftoppm'),'-f',str(n),'-l',str(n),'-scale-to','1600','-png','-singlefile',str(out/(pptx.stem+'.pdf')),str(out/f'{pptx.stem}-{n}')],check=True)
 print('Rendered',pptx.stem,flush=True)
