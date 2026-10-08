from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Cm
from zipfile import ZipFile
from PIL import Image
import json,hashlib,copy
root=Path(__file__).resolve().parent;project=root.parents[1];oldroot=project/'validation/summary_reference/decks'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def title(slide):return next((s.text for s in slide.shapes if s.has_text_frame and s.text.strip()),'')
def clean(x):
 if isinstance(x,dict):return {k:('PATH' if k in ('input_file','output_file','subbusiness_file') else clean(v)) for k,v in x.items()}
 if isinstance(x,list):return [clean(v) for v in x]
 return x
reports=[]
for path in sorted((root/'decks').glob('*.pptx')):
 a=Presentation(oldroot/path.name);b=Presentation(path);assert len(a.slides)==len(b.slides)
 assert [title(s) for s in a.slides]==[title(s) for s in b.slides]
 affected=[]
 for n,(before,after) in enumerate(zip(a.slides,b.slides),1):
  heading=title(after)
  if ' Share of ' in heading:
   assert len(before.shapes)==len(after.shapes)
   for old,new in zip(before.shapes,after.shapes):
    if old.shape_type==13 and old.left==Inches(4.8) and old.top==Inches(1.45):
     assert new.width==Cm(20.0);assert (old.left,old.top,old.height)==(new.left,new.top,new.height)
     assert old.image.blob==new.image.blob
     chart={'old_width_cm':old.width/360000,'new_width_cm':new.width/360000,'x_cm':new.left/360000,'y_cm':new.top/360000,'height_cm':new.height/360000}
     new.width=old.width
    assert old._element.xml==new._element.xml,(path.name,n,new.name)
   affected.append(n)
  elif heading=='Thank You':
   affected.append(n);texts=[s.text.replace('\x0b','\n') for s in after.shapes if s.has_text_frame and s.text]
   period=next(s.text for s in b.slides[0].shapes if s.has_text_frame and s.top==Inches(3.25))
   assert texts==['Thank You','For Your Partnership','Let’s Turn The Agreed Actions\nInto Measurable Progress',period],texts
   pics=[s for s in after.shapes if s.shape_type==13];assert len(pics)==1
   pic=pics[0];asset=project/'backend/V9_Beta2/assets/Thank_You.png';assert pic.image.blob==asset.read_bytes()
   assert abs(pic.width/pic.height-1.5)<1e-6
   assert not any((pic.crop_left,pic.crop_top,pic.crop_right,pic.crop_bottom))
   assert pic.left+pic.width==b.slide_width and pic.top==0 and pic.height==Inches(7.37)
   oldfooter=next(s for s in before.shapes if s.top==Inches(7.37));newfooter=next(s for s in after.shapes if s.top==Inches(7.37));assert oldfooter._element.xml==newfooter._element.xml
   for s in after.shapes:assert 0<=s.left and 0<=s.top and s.left+s.width<=b.slide_width and s.top+s.height<=b.slide_height
  else:assert before._element.xml==after._element.xml,(path.name,n,heading)
 with ZipFile(oldroot/(path.stem+'_audit.zip')) as a,ZipFile(root/'decks'/(path.stem+'_audit.zip')) as b:
  assert set(a.namelist())==set(b.namelist())
  for name in a.namelist():
   x,y=a.read(name),b.read(name)
   if name.endswith('.json'):x,y=clean(json.loads(x)),clean(json.loads(y))
   assert x==y,(path.name,name)
  count=len(a.namelist())
 reports.append({'case':path.stem,'changed_slides_only':affected,'chart_width_only':chart,'reporting_period':period,'audit_entries_unchanged':count,'all_other_slides_xml_identical':True})
allowed={str((project/'web/presentation_update.py').resolve()),str((project/'backend/V9_Beta2/assets/Thank_You.png').resolve())}
for directory,files in json.loads((root/'before_hashes.json').read_text()).items():
 for name,checksum in files.items():
  path=Path(directory)/name
  if str(path) not in allowed:assert sha(path)==checksum,str(path)
assert Image.open(project/'backend/V9_Beta2/assets/Thank_You.png').getchannel('A').getextrema()==(0,255)
(root/'verification.json').write_text(json.dumps(reports,indent=2));print(json.dumps(reports,indent=2));print('PASS: only authorized source/asset changes; frozen analytical reference unchanged.')
