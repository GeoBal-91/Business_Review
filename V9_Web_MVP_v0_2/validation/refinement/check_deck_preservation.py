from pathlib import Path
from collections import defaultdict
from copy import deepcopy
import json
from pptx import Presentation
root=Path(__file__).resolve().parent
reports=[]
def indexed(prs):
    counts=defaultdict(int);result={}
    for slide in prs.slides:
        title=next((s.text for s in slide.shapes if s.has_text_frame and s.text.strip()),'')
        counts[title]+=1;result[(title,counts[title])]=slide
    return result
for case in json.loads((root/'final/generation.json').read_text()):
    name=case['case'];a,b=[indexed(Presentation(root/p/(name+'.pptx'))) for p in ['before','final']]
    unchanged=charts=0
    for key,old in a.items():
        if key[0].startswith('Action Plan') or (key[0].startswith('Operational & Data Review') and 'Store Coverage' not in key[0]):continue
        new=b[key]
        assert old._element.xml==new._element.xml,(name,key,'slide')
        unchanged+=1
        for x,y in zip([s.chart for s in old.shapes if s.has_chart],[s.chart for s in new.shapes if s.has_chart]):
            ox,ny=deepcopy(x._chartSpace),deepcopy(y._chartSpace)
            if 'Performance & Conversion Evolution' in key[0]:
                for marker in ny.xpath('.//c:lineChart/c:ser/c:marker/c:spPr'):
                    marker.getparent().remove(marker)
            assert ox.xml==ny.xml,(name,key,'chart')
            charts+=1
    reports.append({'case':name,'unchanged_slide_xml_including_combined_slide':unchanged,'unchanged_chart_xml_except_requested_marker_style':charts})
(root/'deck_preservation.json').write_text(json.dumps(reports,indent=2))
print('PASS',json.dumps(reports))
