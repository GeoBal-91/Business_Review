from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
import json,sys
ROOT=Path('/tmp/partner-refinement-sep22')
sys.path.insert(0,str(ROOT/'tests'))
from refinement_fixtures import results
source=Path(sys.argv[1])
sys.path.insert(0,str(source))
from web.services import invoke
from pptx import Presentation
out=ROOT/'validation/refinement'/sys.argv[2];out.mkdir(parents=True,exist_ok=True)
reports=[]
for name,triage,raw in results():
    buffer=BytesIO();raw.to_excel(buffer,index=False)
    logo=(source/'backend/V9_Beta2/Partner Logo.png').read_bytes()
    result=invoke('generate',buffer.getvalue(),logo,triage.analysis.settings.partner_sbu,partner_name='PetFood' if name!='FLOWERS_missing' else 'ყვავილების სახლი')
    assert result['ok'],result
    (out/(name+'.pptx')).write_bytes(result['pptx'])
    (out/(name+'_audit.zip')).write_bytes(result['audit'])
    (out/(name+'_input.xlsx')).write_bytes(buffer.getvalue())
    prs=Presentation(BytesIO(result['pptx']))
    affected=[]
    for i,s in enumerate(prs.slides,1):
        title=next((x.text for x in s.shapes if x.has_text_frame and x.text.strip()),'')
        if title.startswith(('Action Plan','Operational & Data Review')) or 'Performance & Conversion Evolution' in title:
            affected.append({'page':i,'title':title})
    reports.append({'case':name,'source':str(source),'sbu':triage.analysis.settings.partner_sbu,'selected_actions':len(triage.actions),'reviews':len(triage.analysis.reviews),'review_types':sorted({r['Type'] for r in triage.analysis.reviews}),'slides':len(prs.slides),'affected':affected})
    print(name,len(triage.actions),len(triage.analysis.reviews),'worker PASS',flush=True)
(out/'generation.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2))
