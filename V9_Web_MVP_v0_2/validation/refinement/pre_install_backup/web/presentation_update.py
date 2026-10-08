"""Web-only slide update over the frozen V2.1 deck.

The engine, original comparison slides, store selection and navigation remain
owned by Beta2. Only the requested slides are inserted, replaced or removed.
"""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import os

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from data_pipeline import G, D, N, U, present
from formatting import (format_currency_2, format_currency, format_number, format_percent,
                        get_row_value, relative_change, pp_change, delta_info)
from ppt_core import add_title, add_footer, add_text, add_kpi_card, set_background
from presentation import context, add_disclosures
from presentation_hardening import benchmark_display, COVERAGE_BENCHMARK_NOTE
from theme import FONT, GREEN, YELLOW, GREY, DARK, MID_GREY
from web.version import PRESENTATION_VERSION

VERSION = PRESENTATION_VERSION
COMBINED_TITLE = 'Partner Performance & Conversion Evolution'
COMPARISONS = (
    ('AOV', '€ AOV Derived', 'currency', True),
    ('CVR', '% CVR', 'percent', True),
    ('Uptime', U, 'percent', True),
    ('PNA', '% PNA Orders', 'percent', False),
    ('Cancels', '% Cancels DTP', 'percent', False),
)
GROWTH = (('GMV Growth', G), ('Orders Growth', D), ('AOV Growth', '€ AOV Derived'))


def title(slide):
    return next((s.text for s in slide.shapes if s.has_text_frame and s.text.strip()), '')


def page(prs, heading, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, heading, subtitle)
    add_footer(slide)
    return slide


def number(value):
    """Blank chart/workbook cells for unknown values, including NaN."""
    return float(value) if present(value) else None


def element(tag, **attrs):
    node = OxmlElement('c:' + tag)
    for key, value in attrs.items():
        node.set(key, str(value))
    return node


def style_chart(chart):
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(10)
    chart.font.color.rgb = DARK
    chart.chart_style = 10
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.size = Pt(10)
        axis.tick_labels.font.color.rgb = GREY
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.color.rgb = MID_GREY
    chart.value_axis.has_major_gridlines = False
    blanks = chart._chartSpace.xpath('./c:chart/c:dispBlanksAs')
    if blanks:
        blanks[0].set('val', 'gap')
    else:
        chart._chartSpace.chart.append(element('dispBlanksAs', val='gap'))


def pair_chart(slide, values, labels, x, y, w, h, fmt, bounds=None):
    """Two literal values, one shared scale per pair; never normalize metrics."""
    data = CategoryChartData(number_format=fmt)
    data.categories = labels
    # Series-level fills and explicit negative handling render signed growth
    # correctly in both Office and LibreOffice. Point-level fills can mirror
    # negative bars in LibreOffice. Blank cells keep the pair on separate rows.
    data.add_series(labels[0], [number(values[0]), None])
    data.add_series(labels[1], [None, number(values[1])])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y),
                                   Inches(w - 1.05), Inches(h), data).chart
    style_chart(chart)
    chart.has_legend = False
    chart.category_axis.reverse_order = True
    chart.category_axis._element.find('{*}tickLblPos').set('val', 'low')
    chart.value_axis.visible = False
    chart.category_axis.format.line.fill.background()
    chart.value_axis.minimum_scale = bounds[0] if bounds else 0
    if bounds:
        chart.value_axis.maximum_scale = bounds[1]
    elif any(present(v) and v > 0 for v in values):
        chart.value_axis.maximum_scale = max(v for v in values if present(v)) * 1.3
    else:
        chart.value_axis.maximum_scale = 1
    plot = chart.plots[0]
    plot.gap_width = 65
    plot.overlap = 100
    # Fixed value column avoids Office/LibreOffice wrapping auto-sized labels.
    plot.has_data_labels = False
    for i, color in enumerate((GREEN, YELLOW)):
        series = chart.series[i]
        series.invert_if_negative = False
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
        series.format.line.fill.background()
    if bounds:
        chart.category_axis.format.line.color.rgb = MID_GREY
    for i, value in enumerate(values):
        shown = ('Not assessed' if not present(value) else
                 format_currency_2(value) if fmt.startswith('€') else
                 f'{value:+.1%}' if fmt.startswith('+') else format_percent(value))
        add_text(slide, shown, x + w - 1.0, y + h * (.25 + .5 * i) - .11,
                 1.03, .40, size=10 if not present(value) else 11, color=GREY if not present(value) else DARK)
    return chart


def comparison_values(ctx):
    return [(name, key, kind, higher, ctx['partner_comparison'].get(key),
             ctx['sbu_comparison'].get(key)) for name, key, kind, higher in COMPARISONS]


def build_partner_visual(prs, ctx):
    period = ctx['comparison_period'].strftime('%B %Y')
    incomplete = (ctx['current_partner_coverage_incomplete']
                  and ctx['comparison_period'] == ctx['partner_latest_period'])
    slide = page(prs, f"Partner vs {ctx['PARTNER_SBU']} · Visual",
                 f'{period} · Comparable performance metrics · Each metric uses its own scale')
    for i, (name, key, kind, higher, partner, sbu) in enumerate(comparison_values(ctx)):
        y = 1.56 + i * .94
        label = ctx['conversion_label'] if name == 'CVR' else name
        add_text(slide, label + (' (€)' if kind == 'currency' else ' (%)'), .85, y + .22, 1.9, .4,
                 size=13, bold=True)
        pair_chart(slide, [partner, sbu], ['Partner', 'SBU'], 2.65, y, 6.8, .89,
                   '€#,##0.00' if kind == 'currency' else '0.0%')
        gap = benchmark_display(partner, sbu, incomplete, metric_type='relative' if kind == 'currency' else 'pp',
                                higher_is_better=higher, label=ctx['PARTNER_SBU'])
        add_text(slide, gap['text'], 9.75, y + .15, 2.75, .65, size=11, bold=True, color=gap['color'])
    note = (COVERAGE_BENCHMARK_NOTE + ' Raw values shown.' if incomplete else
            f'Benchmark comparison uses the latest common period available in both datasets: {period}.')
    if not ctx['include_uptime']:
        note += ' FLOWERS uptime is opening-hours hygiene, excluded from business Priority.'
    add_text(slide, note, .85, 6.55, 11.65, .60, size=10, color=GREY)
    return slide


def growth_values(ctx):
    """Exactly the original Growth vs SBU helper and guarded previous rows."""
    return [(label,
             relative_change(ctx['partner_comparison'][key], get_row_value(ctx['partner_comp_previous'], key)),
             relative_change(ctx['sbu_comparison'][key], get_row_value(ctx['sbu_comp_previous'], key)))
            for label, key in GROWTH]


def build_growth_visual(prs, ctx):
    slide = page(prs, f"Growth vs {ctx['PARTNER_SBU']} · Visual",
                 f"MoM growth comparison — {ctx['comparison_period'].strftime('%B %Y')} · Shared percentage scale")
    rows = growth_values(ctx)
    available = [v for _, p, s in rows for v in (p, s) if present(v)]
    low, high = min([0] + available), max([0] + available)
    span = high - low or .1
    bounds = (low - .18 * span if low < 0 else 0, high + .3 * span)
    for i, (label, partner, sbu) in enumerate(rows):
        y = 1.7 + i * 1.5
        add_text(slide, label, .85, y + .43, 2.0, .45, size=13, bold=True)
        pair_chart(slide, [partner, sbu], ['Partner', 'SBU'], 2.75, y, 6.65, 1.23,
                   '+0.0%;-0.0%;0.0%', bounds)
        # This is the same displayed gap as the existing table, not a growth model.
        gap = f'{(partner - sbu) * 100:+.1f} pp' if present(partner) and present(sbu) else 'Comparison unavailable'
        add_text(slide, 'Growth Gap\n' + gap, 9.8, y + .29, 2.7, .65, size=12, bold=True)
    return slide


def trend_frame(ctx):
    """Retain the original last-N observed rows, inserting blank calendar slots.

    Reindexing is display alignment only: no aggregation, filling or interpolation.
    Both charts use the partner timeline, independent of SBU coverage.
    """
    frame = ctx['monthly'].tail(ctx['TREND_MONTHS']).copy().set_index('Period')
    if frame.empty:
        return frame
    return frame.reindex(pd.date_range(frame.index.min(), frame.index.max(), freq='MS'))


def combo_chart(slide, frame, bar_key, line_key, bar_label, line_label, y, percent=False,
                *, x=.8, width=11.8, height=2.28, compact=False):
    data = CategoryChartData()
    data.categories = frame.index.strftime('%b %y').tolist()
    data.add_series(bar_label, [number(v) for v in frame[bar_key]], '#,##0')
    data.add_series(line_label, [number(v) for v in frame[line_key]], '0.0%' if percent else '#,##0')
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                   Inches(width), Inches(height), data).chart
    style_chart(chart)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.name = FONT
    chart.legend.font.size = Pt(11)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.tick_labels.number_format = '#,##0'
    bar_known = [number(v) for v in frame[bar_key] if present(v)]
    if not bar_known or max(bar_known) == 0:
        chart.value_axis.maximum_scale = 1
        chart.value_axis.major_unit = 1
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = bar_label
    chart.plots[0].gap_width = 85
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = YELLOW
    chart.series[0].format.line.fill.background()
    # python-pptx has no public combo constructor. Keep the native workbook and
    # move the second series into an OOXML line plot with independent right axis.
    area = chart._chartSpace.chart.plotArea
    bar = area.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}barChart')
    line_series = bar.findall('{http://schemas.openxmlformats.org/drawingml/2006/chart}ser')[1]
    bar.remove(line_series)
    sp = line_series.get_or_add_spPr()
    from pptx.dml.line import LineFormat
    line = LineFormat(sp)
    line.color.rgb = GREEN
    line.width = Pt(2.5)
    marker = element('marker')
    marker.append(element('symbol', val='circle'))
    marker.append(element('size', val=5))
    line_series.insert(4, marker)  # after spPr, per the line-series OOXML order
    line_series.append(element('smooth', val=0))
    line_plot = element('lineChart')
    line_plot.append(element('grouping', val='standard'))
    line_plot.append(line_series)
    line_plot.append(element('marker', val=1))
    line_plot.append(element('smooth', val=0))
    line_plot.append(element('axId', val=901001))
    line_plot.append(element('axId', val=901002))
    area.insert(1, line_plot)
    cat, val = deepcopy(chart.category_axis._element), deepcopy(chart.value_axis._element)
    for axis, axis_id, cross_id, position in ((cat, 901001, 901002, 'b'), (val, 901002, 901001, 'r')):
        axis.find('{*}axId').set('val', str(axis_id))
        axis.find('{*}crossAx').set('val', str(cross_id))
        axis.find('{*}axPos').set('val', position)
    cat.find('{*}delete').set('val', '1')
    cat.find('{*}crosses').set('val', 'max')
    val.find('{*}tickLblPos').set('val', 'high')
    val.find('{*}numFmt').set('formatCode', '0%' if percent else '#,##0')
    val.find('{*}numFmt').set('sourceLinked', '0')
    # The right axis must not inherit a zero-count scale from the left axis.
    for child in list(val.find('{*}scaling')):
        if child.tag.rsplit('}', 1)[-1] == 'max':
            val.find('{*}scaling').remove(child)
    for child in list(val):
        if child.tag.rsplit('}', 1)[-1] == 'majorUnit':
            val.remove(child)
    line_known = [number(v) for v in frame[line_key] if present(v)]
    scaling = val.find('{*}scaling')
    if not line_known:
        val.find('{*}tickLblPos').set('val', 'none')
        scaling.insert(scaling.index(scaling.find('{*}min')), element('max', val=1))
    elif max(line_known) == 0:
        scaling.insert(scaling.index(scaling.find('{*}min')), element('max', val=1))
        val.append(element('majorUnit', val=.25 if percent else 1))
    for text in val.xpath('.//a:t'):
        text.text = line_label
    area.append(cat)
    area.append(val)
    # Axis title sizes are explicit, so Office themes cannot enlarge them.
    for tx in chart._chartSpace.xpath('.//c:valAx/c:title/c:tx/c:rich/a:p/a:pPr/a:defRPr'):
        tx.set('sz', '1000')
    if compact:
        chart.legend.font.size = Pt(9)
        chart.category_axis.tick_labels.font.size = Pt(8)
        chart.value_axis.tick_labels.font.size = Pt(8)
        # The section heading and legend supply metric names in the narrow charts.
        for axis in chart._chartSpace.xpath('.//c:valAx'):
            for axis_title in list(axis.findall('{*}title')):
                axis.remove(axis_title)
            for run in axis.xpath('.//a:defRPr'):
                run.set('sz', '800')
    return chart


def meeting_values(ctx):
    """Display existing aggregated outputs with the frozen guarded previous row."""
    latest, previous = ctx['latest_partner'], ctx['partner_previous']
    keys = [G, D, N, '% CVR', '€ AOV Derived']
    changes = {key: (pp_change if key == '% CVR' else relative_change)(
        latest.get(key), get_row_value(previous, key)) for key in keys}
    # Existing helper owns missing/coverage qualification and colors. Only the
    # directional presentation wording changes on this one slide.
    gap = benchmark_display(latest.get('% CVR'), ctx['SBU_CVR'],
        ctx['current_partner_coverage_incomplete'], label=ctx['PARTNER_SBU'])
    gap = dict(gap)
    gap['text'] = (gap['text'].replace('better than ', 'above ')
                  .replace('worse than ', 'below '))
    if gap['text'] not in {'Benchmark unavailable', COVERAGE_BENCHMARK_NOTE}:
        gap['text'] += ' benchmark'
    insights = []
    if present(changes[G]):
        insights.append(f"GMV changed {changes[G]:+.1%} vs previous month.")
    else:
        insights.append('GMV MoM comparison unavailable.')
    if present(changes[D]) and present(changes['€ AOV Derived']):
        if abs(changes[D]) > abs(changes['€ AOV Derived']):
            insights.append('Order volume moved more sharply than basket size.')
        else:
            insights.append('Basket-size movement was larger than the change in order volume.')
    bottom = ('Partner CVR: ' if gap['text'] in {'Benchmark unavailable', COVERAGE_BENCHMARK_NOTE}
              else 'Partner CVR is ') + gap['text'] + ('' if gap['text'].endswith('.') else '.')
    return latest, changes, gap, insights, bottom


def meeting_card(slide, x, y, label, value, delta, gap=None):
    """Reuse existing rounded KPI cards, with room for wrapped benchmark labels."""
    start = len(slide.shapes)
    add_kpi_card(slide, x, y, 2.55, 1.85, label, value, delta['text'], delta['color'],
                 gap['text'] if gap else None, gap['color'] if gap else GREY)
    shapes = list(slide.shapes)[start:]
    for i, shape in enumerate(shapes):
        shape.name = f'Meeting KPI {label} {i}'
    # Preserve typography while giving unavailable/long-SBU text wrapping space.
    shapes[3].height = Inches(.34)
    if gap:
        shapes[4].top = Inches(y + 1.32)
        shapes[4].height = Inches(.48)


def build_combined(prs, ctx):
    slide = page(prs, COMBINED_TITLE, 'Monthly GMV and delivered orders · Partner sessions and conversion')
    frame = trend_frame(ctx)
    if frame.empty:
        add_text(slide, 'Not assessed — monthly data unavailable', .85, 2, 11.6, .5, size=14)
        return slide
    add_text(slide, 'GMV & Orders', .85, 1.5, 6.1, .3, size=13, bold=True)
    combo_chart(slide, frame, G, D, 'GMV (€)', 'Orders', 1.84,
                width=6.4, height=2.13, compact=True)
    add_text(slide, f"Partner Sessions × Partner {ctx['conversion_label']}", .85, 4.24, 6.1, .3, size=13, bold=True)
    combo_chart(slide, frame, N, '% CVR', 'Partner Sessions', 'Partner ' + ctx['conversion_label'],
                4.58, percent=True, width=6.4, height=2.13, compact=True)
    latest, changes, gap, insights, bottom = meeting_values(ctx)
    for x, y, label, key, formatter in [
        (7.45, 1.65, 'GMV', G, format_currency),
        (10.35, 1.65, 'Orders', D, format_number),
        (7.45, 4.4, 'Partner Sessions', N, format_number),
        (10.35, 4.4, 'Partner CVR', '% CVR', format_percent),
    ]:
        delta = delta_info(changes[key], metric_type='pp' if key == '% CVR' else 'relative')
        meeting_card(slide, x, y, label, formatter(latest.get(key)), delta,
                     gap if key == '% CVR' else None)
    for i, text in enumerate(insights):
        box = add_text(slide, '• ' + text, 7.45, 3.58 + i * .30, 5.45, .26 if i == 0 else .42,
                       size=10.5, color=DARK)
        box.name = f'Meeting interpretation top {i}'
    box = add_text(slide, '• ' + bottom, 7.45, 6.39, 5.45, .53, size=10.5, color=DARK)
    box.name = 'Meeting interpretation conversion'
    missing = any(frame[key].isna().any() for key in (G, D, N, '% CVR'))
    if missing:
        unavailable = [label for key, label in ((G, 'GMV'), (D, 'Orders'), (N, 'Partner Sessions'),
                       ('% CVR', 'Partner ' + ctx['conversion_label'])) if frame[key].isna().all()]
        note = (', '.join(unavailable) + ': Not assessed. ' if unavailable else '')
        note += 'Missing months / measurements are blank. Lines break at missing values.'
        add_text(slide, note, .85, 7.02, 11.65, .23, size=9, color=GREY)
    return slide


def enhance_operational(slide, ctx):
    # Keep every existing string, card, status color and disclosure. Move only the
    # Interpretation heading/body to make room for the three operational charts.
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.text == 'Interpretation':
            shape.top = Inches(5.93)
        elif shape.text.startswith(('Raw rates reflect stores', 'Uptime is evaluated')):
            shape.top, shape.height = Inches(6.34), Inches(.75)
    for i, (label, key, benchmark) in enumerate((
            ('Operating Store Uptime', U, ctx['UPTIME_TARGET']),
            ('Partner PNA', '% PNA Orders', ctx['SBU_PNA']),
            ('Partner Cancels', '% Cancels DTP', ctx['SBU_CANCELS']))):
        if key not in ctx['monthly'].columns:
            continue
        pair_chart(slide, [ctx['latest_partner'].get(key), benchmark],
                   ['Partner', '95% target' if key == U else 'SBU'],
                   .75 + 4.05 * i, 3.92, 3.7, 1.73, '0.0%')


def apply_update(prs, result):
    ctx = context(result)
    old_slides, old_ids = list(prs.slides), list(prs.slides._sldIdLst)
    additions, replacements, removed, new_slides = {}, {}, set(), []
    for i, slide in enumerate(old_slides):
        heading = title(slide)
        builder = None
        if heading == f"Partner vs {ctx['PARTNER_SBU']}":
            builder = build_partner_visual
        elif heading == f"Growth vs {ctx['PARTNER_SBU']}":
            builder = build_growth_visual
        elif heading == 'Partner Performance Evolution':
            builder = build_combined
            removed.add(i)
        elif heading in {'Executive Summary', 'Growth Drivers', 'Conversion Performance'} or heading.startswith('Store Performance Matrix'):
            removed.add(i)
        elif heading == 'Operational Health':
            enhance_operational(slide, ctx)
        if builder:
            new_slide = builder(prs, ctx)
            new_slides.append(new_slide)
            (replacements if i in removed else additions)[i] = prs.slides._sldIdLst[-1]
    add_disclosures(SimpleNamespace(slides=new_slides), result, ctx)
    for slide in new_slides:
        slide.notes_slide.notes_text_frame.text = (
            '[Sources]\nExisting Beta2 analysis output and presentation.context; '
            'input hashes in the unchanged audit export.\n[/Sources]\n'
            + result.settings.conversion_description)
    # Append before removing to preserve unique part names and existing hyperlinks.
    final_ids = []
    for i, item in enumerate(old_ids):
        if i in replacements:
            final_ids.append(replacements[i])
        if i not in removed:
            final_ids.append(item)
        else:
            prs.part.drop_rel(item.rId)
        if i in additions:
            final_ids.append(additions[i])
    for item in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(item)
    for item in final_ids:
        prs.slides._sldIdLst.append(item)
    return prs


def identity_title_templates(partner_name, sbu):
    """Explicit identity slots; analytical labels and arbitrary text are not targets."""
    return {
        'Partner Business Review': f'{partner_name} Business Review',
        f'Partner vs {sbu}': f'{partner_name} vs {sbu}',
        f'Partner vs {sbu} · Visual': f'{partner_name} vs {sbu} · Visual',
        f'Growth vs {sbu}': f'{partner_name} Growth vs {sbu}',
        f'Growth vs {sbu} · Visual': f'{partner_name} Growth vs {sbu} · Visual',
        f'Partner Share of {sbu}': f'{partner_name} Share of {sbu}',
        COMBINED_TITLE: f'{partner_name} Performance & Conversion Evolution',
    }


def fit_identity_title(shape, value):
    """Keep the existing title footprint; wrap/shrink long names, never truncate."""
    import textwrap
    from pptx.enum.text import MSO_AUTO_SIZE
    tf = shape.text_frame
    paragraph = tf.paragraphs[0]
    original_size = paragraph.font.size.pt if paragraph.font.size else 20
    paragraph.text = value
    width = (shape.width - tf.margin_left - tf.margin_right) / 12700
    height = (shape.height - tf.margin_top - tf.margin_bottom) / 12700
    size = original_size
    # Conservative glyph-width estimate reserves enough room for Georgian and
    # mixed scripts. Rendered checks cover realistic long-name layouts.
    while size > 8:
        columns = max(1, int(width / (size * .65)))
        lines = sum(max(1, len(textwrap.wrap(line, width=columns))) for line in value.split('\n'))
        if lines * size * 1.18 <= height:
            break
        size -= .5
    paragraph.font.size = Pt(size)
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    shape.name = 'Presentation identity title'


def apply_display_name(prs, partner_name, sbu):
    """Apply display identity after all analytical slide building is complete."""
    from web.services import normalize_partner_name
    name = normalize_partner_name(partner_name)
    templates = identity_title_templates(name, sbu)
    for slide in prs.slides:
        # Remove only the named legacy identity textbox, if processing an older deck.
        for shape in list(slide.shapes):
            if shape.name == 'Presentation Partner Name':
                slide.shapes._spTree.remove(shape._element)
        heading = next((shape for shape in slide.shapes
                        if shape.has_text_frame and shape.text.strip()), None)
        if heading is not None and heading.text in templates:
            fit_identity_title(heading, templates[heading.text])
    prs.core_properties.title = f'{name} Business Review'[:255]
    prs.core_properties.subject = name[:255]
    return prs


def render(triage, output, *, partner_name=None):
    from triage_presentation import render as render_frozen
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix='web-presentation-') as tmp:
        original = render_frozen(triage, Path(tmp) / 'original.pptx')
        prs = apply_update(Presentation(original), triage.analysis)
        if partner_name is not None:
            apply_display_name(prs, partner_name, triage.analysis.settings.partner_sbu)
        staged = Path(tmp) / 'updated.pptx'
        prs.save(staged)
        # Keep the prior complete output intact if construction or saving fails.
        import shutil
        pending = output.with_name(output.stem + '.tmp.pptx')
        try:
            shutil.copyfile(staged, pending)
            os.replace(pending, output)
        finally:
            pending.unlink(missing_ok=True)
    return output
