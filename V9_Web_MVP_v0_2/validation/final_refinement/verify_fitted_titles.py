from pathlib import Path
import json,re,hashlib
import pandas as pd
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
root=Path(__file__).resolve().parent
records=[]
for name in ['Actions_1','Actions_2','Actions_6','Zero_Zero','Long_Mixed_Unicode']:
 path=root/'decks'/(name+'.pptx');prs=Presentation(path)
 raw=pd.read_excel(root/'decks'/(name+'_input.xlsx'))
 stores=raw.drop_duplicates('Store ID').set_index('Store ID')['Store Address Address'].to_dict()
 titles=[]
 for n,slide in enumerate(prs.slides,1):
  for sh in slide.shapes:
   if not sh.name.startswith('Opportunity store: '):continue
   sid=sh.name.removeprefix('Opportunity store: ')
   assert sh.text.split(' · ',1)[1]==stores[sid]
   assert sh.text_frame.auto_size==MSO_AUTO_SIZE.NONE
   para=sh.text_frame.paragraphs[0];size=para.font.size.pt
   assert size>=10.5
   if stores[sid].startswith('Pet supplies branch'):assert size==12.5
   for r in para.runs:
    if re.search('[\u10a0-\u10ff]',r.text) and re.search('[A-Za-z]',sh.text):assert r.font.name=='Noto Sans Georgian'
   titles.append(dict(slide=n,store_id=sid,full_text=sh.text,font_pt=size,automatic_resizing=False))
 records.append(dict(case=name,pptx_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),titles=titles))
(root/'final_title_checks.json').write_text(json.dumps(records,ensure_ascii=False,indent=2))
print('Full title preservation, local font sizes and disabled automatic resizing PASS')
