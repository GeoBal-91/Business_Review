from pathlib import Path
from io import BytesIO
from datetime import date
import sys,json,hashlib
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'tests'))
from refinement_fixtures import results,inputs,BACKEND,Settings,analyze,load_frames,classify,D,U,G
import pandas as pd
from web.services import invoke
from web.meeting_workspace import output_view
from web.presentation_update import title
from pptx import Presentation
out=Path(__file__).resolve().parent/'decks';out.mkdir(exist_ok=True)
cases=list(results())
name,sbu,raw=next(inputs())
for name,orders,uptime in [('Zero_Missing',0,float('nan')),('Missing_Missing',float('nan'),float('nan')),('Zero_Zero',0,0),('Positive_Missing',100,float('nan'))]:
    source=raw.copy();mask=source['Month'].eq('August') if 'Month' in source else None
    # fixture uses calendar month column, inspect the canonical source below.
    month_col=next(c for c in source if 'Month' in c)
    mask=source[month_col].astype(str).str.contains('August')
    assert mask.any(),source.columns
    source.loc[mask,D]=orders;source.loc[mask,U]=uptime
    if orders == 0: source.loc[mask,G]=0
    result=analyze(load_frames(source,pd.read_excel(BACKEND/'subbusiness_unit_data.xlsx'),Settings(partner_sbu=sbu,as_of_date=date(2026,9,23))))
    cases.append((name,classify(result),source))
records=[]
for name,triage,raw in cases:
    buf=BytesIO();raw.to_excel(buf,index=False)
    display='პეტფუდი — საერთაშორისო პარტნიორი და შინაური ცხოველების მომსახურების ცენტრი' if name=='Actions_6' else 'PetFood'
    result=invoke('generate',buf.getvalue(),(BACKEND/'Partner Logo.png').read_bytes(),triage.analysis.settings.partner_sbu,partner_name=display)
    assert result['ok'],result
    (out/(name+'.pptx')).write_bytes(result['pptx']);(out/(name+'_audit.zip')).write_bytes(result['audit']);(out/(name+'_input.xlsx')).write_bytes(buf.getvalue())
    prs=Presentation(BytesIO(result['pptx']));headings=[title(s) for s in prs.slides]
    view=output_view(triage)
    texts=lambda s:[x.text for x in s.shapes if x.has_text_frame]
    assert 'August 2026' in texts(prs.slides[0]);assert 'V9_Beta2' not in texts(prs.slides[0])
    assert sum('Additional agreed actions' in texts(s) for s in prs.slides)==1
    d=headings.index('Discussion');t=headings.index('Thank You');assert all(h.startswith('Action Plan') for h in headings[d+1:t])
    assert len([x for s in prs.slides for x in s.shapes if x.name.endswith(' recommendation') and x.name.startswith('Plan action')])==len(view.actions)
    excluded={a['Store ID'] for a in triage.actions}-{a['Store ID'] for a in view.actions}
    for slide in prs.slides:
        if title(slide).startswith('Priority Opportunities & Actions'):
            assert not any(x.name=='Opportunity store: '+sid for sid in excluded for x in slide.shapes)
    affected=[dict(page=i,title=title(s)) for i,s in enumerate(prs.slides,1) if i==1 or title(s).startswith(('Priority Opportunities','Operational & Data Review','Action Plan','Discussion','Thank You')) or 'Performance & Conversion Evolution' in title(s)]
    records.append(dict(case=name,actions=len(view.actions),reviews=len(view.analysis.reviews),affected=affected))
    print(name,len(view.actions),len(view.analysis.reviews),'PASS',flush=True)
(out/'generation.json').write_text(json.dumps(records,ensure_ascii=False,indent=2))
