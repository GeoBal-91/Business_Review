from pathlib import Path
import sys,json
from zipfile import ZipFile
root=Path(__file__).resolve().parent;project=root.parents[1]
sys.path[:0]=[str(project),str(project/'backend/V9_Beta2')]
from config import Settings
from analysis import run_analysis
from business_classification import classify
from web.meeting_workspace import output_view,review_summary
from pptx import Presentation
report=[]
for case in json.loads((root/'decks/generation.json').read_text()) + [json.loads((root/'decks/long_names_generation.json').read_text())]:
 name=case['case'];p=root/'decks'/(name+'.pptx');prs=Presentation(p)
 with ZipFile(p) as z:assert len(z.namelist())==len(set(z.namelist()))
 triage=classify(run_analysis(Settings(input_file=root/'decks'/(name+'_input.xlsx'),partner_sbu='FLOWERS' if name=='FLOWERS_missing' else 'PETS')))
 view=output_view(triage);shown_actions=[];shown_reviews=[];shown_opportunities=[]
 for slide in prs.slides:
  vals={x.name:x.text for x in slide.shapes if x.has_text_frame}
  for a in view.actions:
   prefix=f"Plan action {a['Triage Rank']}"
   if prefix+' store' in vals:
    shown_actions.append(a['Store ID'])
    for suffix,key in [('store','Store'),('recommendation','Action'),('primary','Metric'),('Follow-up KPI','Metric')]:assert vals[prefix+' '+suffix]==a[key]
   if 'Opportunity store: '+str(a['Store ID']) in vals:shown_opportunities.append(a['Store ID'])
  if 'Review summary items to review' in vals:
   for label,n in zip(['items to review','stores affected','critical observations'],review_summary(view.analysis.reviews)):assert vals['Review summary '+label]==str(n)
   for i,r in enumerate(view.analysis.reviews):
    prefix=f'Review {i}'
    if prefix+' store' in vals:
     shown_reviews.append(i)
     for suffix,key in [('evidence','Issue'),('action','Action'),('type','Type')]:assert vals[prefix+' '+suffix]==r[key]
     assert str(r['Store ID']) not in '\n'.join(vals.values())
   assert not any(n.endswith(' location') for n in vals)
  for shape in slide.shapes:
   if shape.has_chart:
    for ser in shape.chart._chartSpace.xpath('.//c:lineChart/c:ser'):
     if not ser.xpath('./c:marker/c:spPr'):continue
     line=ser.xpath('./c:spPr/a:ln/a:solidFill/a:srgbClr')[0].get('val')
     assert ser.xpath('./c:marker/c:spPr/a:solidFill/a:srgbClr')[0].get('val')==line
     assert ser.xpath('./c:marker/c:spPr/a:ln/a:solidFill/a:srgbClr')[0].get('val')==line
  if any(n.startswith(('Plan action','Additional action','Review summary')) for n in vals):
   for shape in slide.shapes:
    assert 0<=shape.left and 0<=shape.top and shape.left+shape.width<=prs.slide_width and shape.top+shape.height<=prs.slide_height
 assert shown_actions==[a['Store ID'] for a in view.actions]
 assert shown_opportunities==shown_actions
 assert shown_reviews==list(range(len(view.analysis.reviews)))
 report.append(dict(case=name,actions=len(shown_actions),review_items=len(shown_reviews),passed=True))
(root/'deck_checks.json').write_text(json.dumps(report,indent=2));print('All ten actual worker deck content and pagination checks PASS')
