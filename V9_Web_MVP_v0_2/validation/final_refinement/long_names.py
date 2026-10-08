from pathlib import Path
from io import BytesIO
import sys,json
root=Path(__file__).resolve().parent;project=root.parents[1]
sys.path[:0]=[str(project),str(project/'backend/V9_Beta2')]
import pandas as pd
from web.services import invoke
from web.presentation_update import title
from pptx import Presentation
name='Long_Mixed_Unicode';raw=pd.read_excel(root/'decks/Actions_2_input.xlsx')
names=['PetFood International — Veterinary Nutrition & Animal Care · São Paulo / München / Αθήνα — Central Flagship Store, Rustaveli Avenue 123', 'ცხოველების კვებისა და მოვლის ცენტრი · München / São Paulo / Αθήνα — Rustaveli Avenue 123']
for sid,label in zip(raw['Store ID'].unique(),names):raw.loc[raw['Store ID'].eq(sid),'Store Address Address']=label
buf=BytesIO();raw.to_excel(buf,index=False)
r=invoke('generate',buf.getvalue(),(project/'backend/V9_Beta2/Partner Logo.png').read_bytes(),'PETS',partner_name='PetFood International · თბილისი / München')
assert r['ok'],r
for suffix,data in [('.pptx',r['pptx']),('_audit.zip',r['audit']),('_input.xlsx',buf.getvalue())]:(root/'decks'/(name+suffix)).write_bytes(data)
prs=Presentation(BytesIO(r['pptx']))
for label in names:
 assert any(label in sh.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
record=dict(case=name,actions=2,reviews=6,affected=[dict(page=i,title=title(s)) for i,s in enumerate(prs.slides,1) if i==1 or title(s).startswith(('Priority Opportunities','Action Plan','Operational & Data Review','Discussion','Thank You')) or 'Performance & Conversion Evolution' in title(s)])
(root/'decks/long_names_generation.json').write_text(json.dumps(record,ensure_ascii=False,indent=2))
print('Additional mixed-script long-name worker deck PASS')
