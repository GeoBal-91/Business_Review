"""Presentation-only identity, worker parity and UI state contracts."""
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
import json
import sys
import unittest
import pandas as pd
from pptx import Presentation
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from web.services import invoke, normalize_partner_name, input_key, WebError

LONG_NAME = "PetFood & Friends — Premium Pet Nutrition, Accessories and Veterinary Care Georgia"
LONG_UNICODE_NAME = 'პეტფუდი — შინაური ცხოველების კვებისა და ვეტერინარული მომსახურების ცენტრი'
NAMES = ['PetFood', '  PetFood  ', 'პეტფუდი', LONG_NAME, LONG_UNICODE_NAME]

class PartnerName(unittest.TestCase):
    def test_normalization_and_invalid_values(self):
        for value in NAMES + ["L'été — O'Brien-ის მაღაზია"]:
            self.assertEqual(normalize_partner_name(value), value.strip())
        for value in ['', '  \t\n', None]:
            with self.assertRaises(WebError):
                normalize_partner_name(value)
            with self.assertRaises(WebError):
                invoke('generate', (ROOT/'backend/V9_Beta2/partner_data.xlsx').read_bytes(),
                       (ROOT/'backend/V9_Beta2/Partner Logo.png').read_bytes(), 'JEWELRY', partner_name=value)
        self.assertEqual(input_key(b'x', b'y', 'A', 'B', ' PetFood '), input_key(b'x', b'y', 'A', 'B', 'PetFood'))
        self.assertNotEqual(input_key(b'x', b'y', 'A', 'B', 'PetFood'), input_key(b'x', b'y', 'A', 'B', 'პეტფუდი'))

    def test_direct_worker_rejects_missing_name(self):
        from web.worker import execute
        from tempfile import TemporaryDirectory
        with TemporaryDirectory() as tmp:
            for value in ('', '  ', None):
                response = execute(dict(mode='generate', sbu='JEWELRY', partner_name=value), Path(tmp))
                self.assertFalse(response['ok'])
                self.assertEqual(response['issues'][0]['code'], 'web_partner_name_required')

    def test_worker_names_precedence_and_complete_audit_parity(self):
        frame = pd.read_excel(ROOT/'backend/V9_Beta2/partner_data.xlsx')
        frame['Partner Name'] = 'Different Export Partner'
        payload = BytesIO(); frame.to_excel(payload, index=False)
        logo = (ROOT/'backend/V9_Beta2/Partner Logo.png').read_bytes()
        outputs = ROOT/'validation/identity_titles'; outputs.mkdir(parents=True, exist_ok=True)
        audit_before = None
        report = []
        for index, name in enumerate(NAMES):
            result = invoke('generate', payload.getvalue(), logo, 'JEWELRY', partner_name=name)
            self.assertTrue(result['ok'], result)
            self.assertEqual(result['metadata']['partner'], 'Different Export Partner')
            self.assertEqual(result['presentation_display_name'], name.strip())
            deck = Presentation(BytesIO(result['pptx']))
            identity = next(s for s in deck.slides[0].shapes if s.name == 'Presentation identity title')
            self.assertEqual(identity.text, name.strip() + ' Business Review')
            self.assertFalse(any(s.name == 'Presentation Partner Name' for slide in deck.slides for s in slide.shapes))
            self.assertLessEqual(identity.left + identity.width, deck.slide_width)
            self.assertLessEqual(identity.top + identity.height, deck.slide_height)
            texts = [s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame]
            self.assertNotIn('Different Export Partner', '\n'.join(texts))
            headings = [next((s.text for s in slide.shapes if s.has_text_frame and s.text.strip()), '') for slide in deck.slides]
            self.assertEqual(headings.count(name.strip() + ' Performance & Conversion Evolution'), 1)
            for absent in ['Executive Summary', 'Partner Performance Evolution', 'Growth Drivers', 'Conversion Performance']:
                self.assertNotIn(absent, headings)
            self.assertFalse(any(h.startswith('Store Performance Matrix') for h in headings))
            for present in [name.strip() + suffix for suffix in [' vs JEWELRY', ' vs JEWELRY · Visual', ' Growth vs JEWELRY', ' Growth vs JEWELRY · Visual', ' Share of JEWELRY']] + ['Operational Health']:
                self.assertIn(present, headings)
            with ZipFile(BytesIO(result['audit'])) as z:
                audit = {n:z.read(n) for n in z.namelist()}
            # Isolated workers necessarily use different temporary request/output paths.
            provenance = json.loads(audit['run_audit.json'])
            for key in ('input_file', 'output_file'):
                provenance['settings'][key] = Path(provenance['settings'][key]).name
            audit['run_audit.json'] = json.dumps(provenance, sort_keys=True).encode()
            if audit_before is None:
                audit_before = audit
            self.assertEqual(audit, audit_before, 'Display identity changed audit outputs')
            stem = ['PetFood', 'PetFood_trimmed', 'Georgian', 'Long_Name', 'Long_Unicode'][index]
            (outputs/(stem+'.pptx')).write_bytes(result['pptx'])
            (outputs/(stem+'_audit.zip')).write_bytes(result['audit'])
            report.append(dict(input=name, display=result['presentation_display_name'], headings=headings, audit_files=len(audit)))
        (outputs/'worker_checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))

    def test_exact_title_templates_and_non_title_preservation(self):
        sys.path.insert(0, str(ROOT/'backend/V9_Beta2'))
        from web.presentation_update import apply_display_name
        from ppt_core import add_title, add_text
        from pptx.util import Inches
        from copy import deepcopy
        original_titles = ['Partner Business Review', 'Partner vs BEAUTY_PERSONAL_CARE',
            'Partner vs BEAUTY_PERSONAL_CARE · Visual', 'Partner Share of BEAUTY_PERSONAL_CARE',
            'Partner Performance & Conversion Evolution', 'Growth vs BEAUTY_PERSONAL_CARE',
            'Growth vs BEAUTY_PERSONAL_CARE · Visual', 'Operational Health']
        for name in ['PetFood', LONG_NAME, LONG_UNICODE_NAME, 'Partner & Friends']:
            deck = Presentation(); deck.slide_width=Inches(13.333); deck.slide_height=Inches(7.5)
            bodies=[]
            for heading in original_titles:
                slide=deck.slides.add_slide(deck.slide_layouts[6]);add_title(slide,heading,'Partner benchmark subtitle')
                body=add_text(slide,'Partner CVR · Partner Sessions · Partner Uptime',.8,2,10,.5)
                bodies.append(body._element.xml)
            legacy=add_text(deck.slides[0],name,.82,4.5,11.5,1.65);legacy.name='Presentation Partner Name'
            apply_display_name(deck,name,'BEAUTY_PERSONAL_CARE')
            expected=[f'{name} Business Review', f'{name} vs BEAUTY_PERSONAL_CARE',
                f'{name} vs BEAUTY_PERSONAL_CARE · Visual', f'{name} Share of BEAUTY_PERSONAL_CARE',
                f'{name} Performance & Conversion Evolution', f'{name} Growth vs BEAUTY_PERSONAL_CARE',
                f'{name} Growth vs BEAUTY_PERSONAL_CARE · Visual', 'Operational Health']
            self.assertEqual([slide.shapes[0].text for slide in deck.slides],expected)
            for slide,xml in zip(deck.slides,bodies):
                self.assertEqual(slide.shapes[2]._element.xml,xml)
                self.assertEqual(slide.shapes[1].text,'Partner benchmark subtitle')
                self.assertFalse(any(s.name=='Presentation Partner Name' for s in slide.shapes))
                self.assertLessEqual(slide.shapes[0].top+slide.shapes[0].height,slide.shapes[1].top)

    def test_ui_name_required_and_changes_reuse_validation(self):
        with patch('web.services.invoke', wraps=invoke) as bridge:
            app = AppTest.from_file(str(ROOT/'app.py'), default_timeout=180).run()
            app.file_uploader[0].set_value(('partner.xlsx', (ROOT/'backend/V9_Beta2/partner_data.xlsx').read_bytes(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))
            app.file_uploader[1].set_value(('logo.png', (ROOT/'backend/V9_Beta2/Partner Logo.png').read_bytes(), 'image/png'))
            app.selectbox[0].select('JEWELRY').run()
            generate = lambda: next(b for b in app.button if b.label == 'Generate Review')
            self.assertTrue(generate().disabled)
            app.text_input[0].set_value('  ').run()
            self.assertTrue(generate().disabled)
            app.text_input[0].set_value('  PetFood  ').run()
            self.assertFalse(generate().disabled)
            generate().click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.session_state['result']['presentation_display_name'], 'PetFood')
            # Filename is carried by Streamlit's media URL, so check the app helper separately.
            from ast import parse, FunctionDef, Module, fix_missing_locations
            tree = parse((ROOT/'app.py').read_text())
            fn = next(n for n in tree.body if isinstance(n, FunctionDef) and n.name == 'download_name')
            import re
            ns = {'re': re}; exec(compile(fix_missing_locations(Module(body=[fn], type_ignores=[])), '<download_name>', 'exec'), ns)
            self.assertEqual(ns['download_name']({'partner':'Other', 'period':'August 2026'}, 'Review.pptx', 'PetFood'), 'PetFood_August_2026_Review.pptx')
            self.assertTrue(ns['download_name']({}, 'Review.pptx', 'პეტფუდი').startswith('პეტფუდი'))
            app.text_input[0].set_value('პეტფუდი').run()
            self.assertEqual(len(app.download_button), 0)
            self.assertEqual(sum(c.args[0] == 'validate' for c in bridge.call_args_list), 1)
            next(b for b in app.button if b.label=='Clear uploads and results').click().run()
            self.assertEqual(app.text_input[0].value, '')
            self.assertTrue(generate().disabled)
            self.assertFalse(app.exception)
