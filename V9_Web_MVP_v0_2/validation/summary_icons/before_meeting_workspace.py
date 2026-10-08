"""Editable meeting slides over frozen triage outputs. No analytical decisions."""
import textwrap
import re
from copy import copy
from data_pipeline import D, U
import pandas as pd
from types import SimpleNamespace

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from formatting import format_score
from ppt_core import add_text, add_picture_contain
from presentation import add_disclosures, context
from theme import LIGHT_GREY, MID_GREY, GREY, DARK, GREEN, YELLOW, RED, ASSETS_DIR

PALE_YELLOW = RGBColor(255, 249, 224)
PALE_RED = RGBColor(252, 237, 238)


REVIEW_ICONS = {
    'Critical observation': 'icon_critical_observation.png',
    'Availability review': 'icon_availability_review.png',
    'Data quality': 'icon_data_quality.png',
    'Opening-hours hygiene': 'icon_opening_hours.png',
    'Activity review': 'icon_store_review.png',
    'Closure verification': 'icon_store_review.png',
}


def approved_icon(slide, filename, x, y, size, *, name):
    """Embed the approved PNG unchanged, preserving its aspect ratio."""
    path = ASSETS_DIR / filename
    if not path.is_file():
        raise FileNotFoundError(f'Approved presentation icon is missing: {path}')
    picture = add_picture_contain(slide, path, x, y, size, size)
    picture.name = name
    return picture


def panel(slide, x, y, w, h, fill=LIGHT_GREY, border=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = .08
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border; shape.line.width = Pt(.5)
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    for ref in shape._element.xpath('./p:style/a:effectRef'):
        ref.set('idx', '0')
    return shape


def text(slide, value, x, y, w, h, *, name, size=11, bold=False, color=DARK):
    shape = add_text(slide, value, x, y, w, h, size=size, bold=bold, color=color)
    if shape is not None:
        shape.name = name
        shape.text_frame.paragraphs[0].space_after = Pt(0)
        shape.text_frame.paragraphs[0].line_spacing = 1.05
    return shape


def lines(value, width, size):
    # Conservative pagination estimate, independent of analytical values.
    columns = max(1, int(width * 72 / (size * .60)))
    return sum(max(1, len(textwrap.wrap(part, width=columns))) for part in str(value).split('\n'))


def source_metadata(result, sid):
    """Use the latest observed source period for this exact normalized store ID."""
    raw = result.data.partner_raw
    rows = raw.loc[raw['Store ID'] == sid]
    if not rows.empty:
        rows = rows.loc[rows['Period'] == rows['Period'].max()]
    return rows


def unambiguous(rows, columns):
    values = {str(v).strip() for c in columns if c in rows for v in rows[c]
              if pd.notna(v) and str(v).strip()}
    return next(iter(values)) if len(values) == 1 else ''


def review_identity(result, item):
    rows = source_metadata(result, item['Store ID'])
    store = str(item['Store']).strip() if pd.notna(item.get('Store')) else ''
    if not store or store == str(item['Store ID']):
        store = unambiguous(rows, ('Store Name',)) or 'Store name unavailable'
    return store, ''


def review_summary(items):
    return len(items), len({item['Store ID'] for item in items}), sum(item['Type'] == 'Critical observation' for item in items)


def review_pages(result):
    pages, current, used = [], [], 0
    for index, item in enumerate(result.reviews):
        store, location = review_identity(result, item)
        height = max(.98, .27 + max(
            lines(store, 3.30, 11.5) * .18 + (.26 if location else 0),
            lines(item['Type'], 1.27, 10.5) * .17,
            lines(item['Issue'], 2.58, 11) * .18,
            lines(item['Action'], 3.62, 11) * .18))
        if current and (len(current) == 4 or used + height + .09 > 4.55):
            pages.append(current); current, used = [], 0
        current.append((index, item, store, location, height)); used += height + .09
    if current:
        pages.append(current)
    return pages or [[]]


def build_reviews(prs, result):
    from web.presentation_update import page
    pages = review_pages(result)
    count, stores, critical = review_summary(result.reviews)
    for p, entries in enumerate(pages, 1):
        heading = 'Operational & Data Review' + (f' · {p}/{len(pages)}' if len(pages) > 1 else '')
        slide = page(prs, heading, 'Critical observations and hygiene remain visible independently of business Priority')
        panel(slide, .65, 1.55, 12.02, .54)
        for i, (value, label) in enumerate(((count, 'items to review'), (stores, 'stores affected'), (critical, 'critical observations'))):
            x = .87 + i * 4.01
            text(slide, str(value), x, 1.66, .55, .30, name='Review summary ' + label, size=17, bold=True)
            text(slide, label, x + .66, 1.73, 3.0, .23, name='Review summary label', size=10, color=GREY)
        if not entries:
            panel(slide, .65, 2.40, 12.02, 1.65)
            text(slide, 'No operational or data review items', .95, 2.8, 11.4, .4, name='Review empty state', size=18, bold=True)
            text(slide, 'No items in this list. See the separate review and diagnostic sections.', .95, 3.4, 11.4, .35, name='Review empty note', color=GREY)
            continue
        for label, x, w in [('STORE / LOCATION', .89, 3.3), ('TYPE', 4.44, 1.45), ('EVIDENCE', 6.1, 2.62), ('REVIEW ACTION', 8.91, 3.5)]:
            text(slide, label, x, 2.23, w, .23, name='Review column ' + label, size=9.5, bold=True, color=GREY)
        y = 2.59
        for index, item, store, location, height in entries:
            prefix = f'Review {index}'
            panel(slide, .65, y, 12.02, height, border=MID_GREY)
            color = RED if item['Type'] == 'Critical observation' else YELLOW
            panel(slide, .65, y + .04, .045, height - .08, color)
            name_h = lines(store, 3.30, 11.5) * .18 + .05
            text(slide, store, .89, y + .16, 3.30, name_h, name=prefix + ' store', size=11.5, bold=True)
            if location:
                text(slide, location, .89, y + .19 + name_h, 3.30, .25, name=prefix + ' location', size=9.5, color=GREY)
            badge_h = max(.44, lines(item['Type'], 1.27, 10.5) * .17 + .14)
            panel(slide, 4.28, y + .17, 1.72, badge_h, PALE_RED if color == RED else PALE_YELLOW)
            approved_icon(slide, REVIEW_ICONS[item['Type']], 4.31, y + .21, .28, name=prefix + ' icon')
            text(slide, item['Type'], 4.65, y + .24, 1.27, badge_h - .11, name=prefix + ' type', size=10.5, bold=True, color=RED if color == RED else DARK)
            text(slide, item['Issue'], 6.10, y + .16, 2.58, height - .24, name=prefix + ' evidence', size=11, bold=True)
            text(slide, item['Action'], 8.91, y + .16, 3.52, height - .24, name=prefix + ' action', size=11)
            y += height + .09


def plan_pages(actions):
    # Retain the engine-selected ordered list in its entirety; never rerank/head().
    pages, current = [], []
    for action in actions:
        oversized = lines(action['Store'], 5.3, 13) > 3 or lines(action['Action'], 5.3, 11.5) > 5
        if oversized:
            if current:
                pages.append(current); current = []
            pages.append([action])
        else:
            current.append(action)
            if len(current) == 2:
                pages.append(current); current = []
    if current:
        pages.append(current)
    return pages


def build_plan(prs, triage):
    from web.presentation_update import page
    pages = plan_pages(triage.actions)
    for p, actions in enumerate(pages, 1):
        heading = 'Action Plan' + (f' · {p}/{len(pages)}' if len(pages) > 1 else '')
        slide = page(prs, heading, 'Agree actions, owners and dates during the business review')
        for i, action in enumerate(actions):
            w = 5.87 if len(actions) == 2 else 12.02
            x, y = .65 + i * 6.15, 1.60
            prefix = f"Plan action {action['Triage Rank']}"
            panel(slide, x, y, w, 4.86, border=MID_GREY)
            approved_icon(slide, 'icon_priority_action.png', x + .18, y + .14, .28, name=prefix + ' priority icon')
            text(slide, f"PRIORITY ACTION {action['Triage Rank']}", x + .55, y + .17, w - 2.85, .25, name=prefix + ' rank', size=10, bold=True, color=GREEN)
            text(slide, 'Priority ' + format_score(action['Priority'], 5), x + w - 1.82, y + .14, 1.6, .30, name=prefix + ' priority', size=12, bold=True)
            text(slide, action['Store'], x + .20, y + .53, w - .40, .72, name=prefix + ' store', size=13, bold=True)
            text(slide, 'Primary Action KPI', x + .20, y + 1.30, 1.7, .25, name=prefix + ' primary label', size=10, color=GREY)
            text(slide, action['Metric'], x + 2.0, y + 1.27, w - 2.2, .3, name=prefix + ' primary', size=12, bold=True)
            text(slide, 'Recommended Action (System)', x + .20, y + 1.67, w - .4, .24, name=prefix + ' recommendation label', size=10, bold=True)
            text(slide, action['Action'], x + .20, y + 1.99, w - .4, 1.04, name=prefix + ' recommendation', size=11.5)
            panel(slide, x + .15, y + 3.05, w - .30, .66, PALE_YELLOW)
            approved_icon(slide, 'icon_agreed_action.png', x + .23, y + 3.10, .26, name=prefix + ' agreed icon')
            text(slide, 'Agreed Action (Meeting)', x + .57, y + 3.13, w - .85, .23, name=prefix + ' agreed label', size=10, bold=True)
            text(slide, '[To agree]', x + .28, y + 3.42, w - .56, .23, name=prefix + ' agreed', size=11, color=GREY)
            for j, (label, value) in enumerate((('Owner', '[To agree]'), ('Contact', '[To agree]'), ('ETA', '[To agree]'), ('Follow-up KPI', action['Metric']))):
                col, row = j % 2, j // 2
                xx = x + .20 + col * (w / 2)
                yy = y + 3.94 + row * .44
                label_w = 1.25 if label == 'Follow-up KPI' else .68
                text(slide, label, xx, yy, label_w, .25, name=prefix + ' ' + label + ' label', size=9, color=GREY)
                text(slide, value, xx + label_w + .08, yy, w / 2 - label_w - .4, .3, name=prefix + ' ' + label, size=10.5, bold=label == 'Follow-up KPI')
        panel(slide, .65, 6.72, 12.02, .43, PALE_YELLOW)
        approved_icon(slide, 'icon_next_checkin.png', .83, 6.80, .28, name='Next check-in icon')
        text(slide, 'Next check-in:', 1.20, 6.83, 1.23, .22, name='Next check-in label', size=10, bold=True)
        text(slide, '[To agree]', 2.48, 6.82, 9.8, .25, name='Next check-in', size=11, color=GREY)


def output_view(triage):
    """Filter only action/review outputs using the application's reporting month.

    Shallow copies preserve every analytical object and original ordered selection.
    Absent current rows are missing current evidence; never fall back to history.
    """
    raw = triage.analysis.data.partner_raw
    current = raw.loc[raw['Period'] == triage.analysis.data.reporting_period]
    eligible = set(current.loc[current[U].notna() | (current[D].notna() & current[D].ne(0)), 'Store ID'])
    view = copy(triage)
    view.analysis = copy(triage.analysis)
    view.actions = [a for a in triage.actions if a['Store ID'] in eligible]
    view.analysis.reviews = [r for r in triage.analysis.reviews if r['Store ID'] in eligible]
    return view


def build_additional_plan(prs):
    from web.presentation_update import page
    slide = page(prs, 'Action Plan', 'Additional agreed actions')
    for i in range(2):
        x, y, w = .65 + i * 6.15, 1.60, 5.87
        prefix = f'Additional action {i + 1}'
        panel(slide, x, y, w, 4.86, border=MID_GREY)
        text(slide, f'ACTION {i + 1}', x + .2, y + .18, w - .4, .3, name=prefix, size=10, bold=True, color=GREEN)
        for label, offset, height in [('Action / Topic', .70, .7), ('Agreed Action', 1.85, 1.0)]:
            inset = .54 if label == 'Agreed Action' else .2
            if label == 'Agreed Action':
                approved_icon(slide, 'icon_agreed_action.png', x + .20, y + offset - .03, .27, name=prefix + ' agreed icon')
            text(slide, label, x + inset, y + offset, w - inset - .2, .25, name=prefix + label + ' label', size=11, bold=True)
            panel(slide, x + .15, y + offset + .35, w - .3, height, PALE_YELLOW)
            text(slide, '[To agree]', x + .28, y + offset + .48, w - .56, .35, name=prefix + label, size=12, color=GREY)
        for j, label in enumerate(['Owner', 'Contact', 'ETA', 'Follow-up KPI']):
            xx = x + .2 + (j % 2) * (w / 2)
            yy = y + 3.63 + (j // 2) * .57
            text(slide, label, xx, yy, 2.4, .22, name=prefix + label + ' label', size=9, color=GREY)
            text(slide, '[To agree]', xx, yy + .24, 2.4, .26, name=prefix + label, size=11)
    panel(slide, .65, 6.72, 12.02, .43, PALE_YELLOW)
    approved_icon(slide, 'icon_next_checkin.png', .83, 6.80, .28, name='Next check-in icon')
    text(slide, 'Next check-in:', 1.20, 6.83, 1.23, .22, name='Next check-in label', size=10, bold=True)
    text(slide, '[To agree]', 2.48, 6.82, 9.8, .25, name='Next check-in', size=11, color=GREY)


def apply_meeting_workspace(prs, triage):
    from web.presentation_update import title
    # Earlier adapters removed/reordered slides. Compact part names before
    # appending so python-pptx's count-based allocator cannot reuse a live name.
    # Relationships reference parts, so navigation targets remain intact.
    prs.part.rename_slide_parts([item.rId for item in prs.slides._sldIdLst])
    triage = output_view(triage)
    old_slides, old_ids = list(prs.slides), list(prs.slides._sldIdLst)
    review_indices = [i for i, slide in enumerate(old_slides) if title(slide) == 'Operational & Data Review' or title(slide).startswith('Operational & Data Review · ') and 'Store Coverage' not in title(slide)]
    plan_indices = [i for i, slide in enumerate(old_slides) if title(slide) == 'Action Plan' or title(slide).startswith('Action Plan · ')]
    opportunity_indices = [i for i, slide in enumerate(old_slides) if title(slide).startswith('Priority Opportunities & Actions')]
    discussion = next(i for i, slide in enumerate(old_slides) if title(slide) == 'Discussion')
    start = len(prs.slides)
    from triage_presentation import build_actions
    build_actions(prs, triage)
    from web.presentation_update import fit_identity_title
    for slide in list(prs.slides)[start:]:
        for shape in slide.shapes:
            if shape.name.startswith('Opportunity store: '):
                original_name = shape.name
                original_size = shape.text_frame.paragraphs[0].font.size
                fit_identity_title(shape, shape.text)
                # Font size is already calculated. Avoid renderer-dependent
                # automatic scaling of mixed-script fallback fonts.
                from pptx.enum.text import MSO_AUTO_SIZE
                shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
                paragraph = shape.text_frame.paragraphs[0]
                value = paragraph.text
                # Montserrat has no Georgian glyphs. Pin only mixed-script
                # title runs to a Georgian face instead of renderer-dependent
                # fallback shaping. Other text keeps the deck typography.
                georgian = r'[\u10a0-\u10ff\u1c90-\u1cbf]'
                if paragraph.font.size < original_size and re.search(georgian, value) and re.search(r'[A-Za-z]', value):
                    paragraph.clear()
                    for part in re.split('(' + georgian + '+)', value):
                        if part:
                            run = paragraph.add_run()
                            run.text = part
                            if re.match(georgian, part):
                                run.font.name = 'Noto Sans Georgian'
                shape.name = original_name
    opportunity_ids = list(prs.slides._sldIdLst)[start:]
    review_start = len(prs.slides)
    build_reviews(prs, triage.analysis)
    review_ids = list(prs.slides._sldIdLst)[review_start:]
    plan_start = len(prs.slides)
    build_plan(prs, triage)
    build_additional_plan(prs)
    plan_ids = list(prs.slides._sldIdLst)[plan_start:]
    new_slides = list(prs.slides)[start:]
    add_disclosures(SimpleNamespace(slides=new_slides), triage.analysis, context(triage.analysis))
    for slide in new_slides:
        slide.notes_slide.notes_text_frame.text = ('Presentation-only layout. Recommended actions, priority and review evidence come unchanged from V9_Beta2. Meeting fields are editable text, initially unagreed.\n' + triage.analysis.settings.conversion_description)
    removed = set(review_indices + plan_indices + opportunity_indices)
    final = []
    for i, item in enumerate(old_ids):
        if opportunity_indices and i == opportunity_indices[0]:
            final.extend(opportunity_ids)
        if review_indices and i == review_indices[0]:
            final.extend(review_ids)
        if i in removed:
            prs.part.drop_rel(item.rId)
        else:
            final.append(item)
        if i == discussion:
            final.extend(plan_ids)
    for item in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(item)
    for item in final:
        prs.slides._sldIdLst.append(item)
    return prs
