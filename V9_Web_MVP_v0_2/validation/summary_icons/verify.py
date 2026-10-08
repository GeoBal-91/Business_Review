from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
from PIL import Image
from io import BytesIO
from lxml import etree
import json,hashlib,copy,re
root=Path(__file__).resolve().parent;project=root.parents[1];assets=project/'backend/V9_Beta2/assets';oldroot=project/'validation/icon_integration/decks'
expected={'items to review':'icon_data_quality.png','stores affected':'icon_store_review.png','critical observations':'icon_critical_observation.png'}
def canonical(s):
 e=copy.deepcopy(s._element)
 for n in e.iter():
  if n.tag.endswith('}cNvPr'):
   n.set('id','0')
   if re.fullmatch(r'(Rounded Rectangle|Rectangle|Picture|TextBox) \d+',n.get('name','')):n.set('name',re.sub(r' \d+$','',n.get('name')))
  for k in list(n.attrib):
   if k.endswith('}embed'):n.set(k,'IMAGE')
 return etree.tostring(e)
def clean(x):
 if isinstance(x,dict):return {k:('PATH' if k in ('input_file','output_file','subbusiness_file') else 'VERSION' if k=='presentation_version' else clean(v)) for k,v in x.items()}
 if isinstance(x,list):return [clean(v) for v in x]
 return x
reports=[]
for path in sorted((root/'decks').glob('*.pptx')):
 old=Presentation(oldroot/path.name);new=Presentation(path);assert len(old.slides)==len(new.slides)
 pages=[]
 for page,(a,b) in enumerate(zip(old.slides,new.slides),1):
  added=[s for s in b.shapes if s.name.startswith('Review summary ') and s.name.endswith(' icon')]
  remaining=[s for s in b.shapes if s not in added]
  assert len(a.shapes)==len(remaining)
  for before,after in zip(a.shapes,remaining):
   if after.name=='Review summary label':
    assert before.text==after.text
    copy_after=copy.deepcopy(after._element)
    after.left=before.left;after.width=before.width
    assert canonical(before)==canonical(after),(path.name,page,after.name)
    after._element.getparent().replace(after._element,copy_after)
   else:assert canonical(before)==canonical(after),(path.name,page,after.name)
   if before.shape_type==13:assert before.image.blob==after.image.blob
  if added:
   assert len(added)==3;pages.append(page)
   shapes={s.name:s for s in b.shapes};labels=[s for s in b.shapes if s.name=='Review summary label']
   for label,filename in expected.items():
    icon=shapes['Review summary '+label+' icon'];number=shapes['Review summary '+label];txt=next(s for s in labels if s.text==label)
    assert icon.shape_type==13 and icon.image.blob==(assets/filename).read_bytes()
    assert icon.width==icon.height and not any((icon.crop_left,icon.crop_right,icon.crop_top,icon.crop_bottom))
    assert number.left+number.width<icon.left and icon.left+icon.width<txt.left
    assert Image.open(BytesIO(icon.image.blob)).getchannel('A').getextrema()==(0,255)
    for text in b.shapes:
     if text.has_text_frame and text.text:
      assert min(icon.left+icon.width,text.left+text.width)<=max(icon.left,text.left) or min(icon.top+icon.height,text.top+text.height)<=max(icon.top,text.top),(page,label,text.name)
 with ZipFile(oldroot/(path.stem+'_audit.zip')) as a,ZipFile(root/'decks'/(path.stem+'_audit.zip')) as b:
  assert set(a.namelist())==set(b.namelist())
  for n in a.namelist():
   x,y=a.read(n),b.read(n)
   if n.endswith('.json'):x,y=clean(json.loads(x)),clean(json.loads(y))
   assert x==y,(path.name,n)
  count=len(a.namelist())
 reports.append({'case':path.stem,'slides':len(new.slides),'summary_icon_pages':pages,'summary_icons':len(pages)*3,'unchanged_audit_entries':count,'only_summary_icons_and_label_spacing_changed':True})
for name,entry in json.loads((project/'validation/icon_integration/transparency_verification.json').read_text()).items():assert hashlib.sha256((assets/name).read_bytes()).hexdigest()==entry['corrected_sha256']
(root/'verification.json').write_text(json.dumps(reports,indent=2));print(json.dumps(reports,indent=2));print('PASS: exact mapping, native pictures, alpha, number/icon/label order, no text overlap, row icons and other content unchanged; pagination and audits unchanged; all eight asset hashes unchanged.')
