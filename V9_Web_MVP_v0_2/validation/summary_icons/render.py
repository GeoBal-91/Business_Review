from pathlib import Path
import subprocess,json,os,sys
from pptx import Presentation
root=Path(__file__).resolve().parent
bin=Path('/Users/giorgibalarjishvili/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override')
previews=root/'previews';previews.mkdir(exist_ok=True)
for pptx in sorted((root/'decks').glob('*.pptx')):
    if sys.argv[1:] and pptx.stem not in sys.argv[1:]:continue
    prs=Presentation(pptx)
    case={'case':pptx.stem,'affected':[]}
    for n,slide in enumerate(prs.slides,1):
        heading=next((x.text for x in slide.shapes if x.has_text_frame and x.text.strip()),'')
        if n==1 or heading.startswith(('Action Plan','Operational & Data Review','Priority Opportunities','Discussion','Thank You')) or 'Performance & Conversion Evolution' in heading:case['affected'].append({'page':n})
    name=case['case'];pdf=previews/(name+'.pdf');pptx=root/'decks'/(name+'.pptx')
    if not pdf.exists() or pdf.stat().st_mtime < pptx.stat().st_mtime:
        subprocess.run([str(bin/'soffice'),'--headless','--convert-to','pdf','--outdir',str(previews),str(pptx)],env=dict(os.environ,FONTCONFIG_FILE=str(root/'fonts.conf')),check=True)
    for item in case['affected']:
        n=item['page'];dest=previews/f'{name}-{n}'
        if dest.with_suffix('.png').exists() and dest.with_suffix('.png').stat().st_mtime>pdf.stat().st_mtime:continue
        subprocess.run([str(bin/'pdftoppm'),'-f',str(n),'-l',str(n),'-scale-to','1600','-png','-singlefile',str(pdf),str(dest)],check=True)
    print('Rendered',name,flush=True)
