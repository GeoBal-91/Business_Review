"""Render the existing Glovo narrative from one explicit analysis result."""
from pathlib import Path
from tempfile import TemporaryDirectory
import os
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.util import Inches
from data_pipeline import period_row, get_previous_month_row, get_yoy_row, G, U
from config import UPTIME_TARGET
from coverage_guard import guarded_previous, comparison, coverage_summary, NOT_COMPARABLE


def context(result):
    d, s = result.data, result.settings
    monthly, sbu = d.monthly, d.selected_sbu
    p = d.reporting_period
    latest = period_row(monthly, p)
    common = sorted(set(monthly.Period) & set(sbu.Period))
    comp = common[-1] if common else p
    empty = pd.Series({c: np.nan for c in monthly.columns})
    def row(frame, period):
        found = period_row(frame, period)
        return found if found is not None else empty
    sp = sbu.Period.max() if not sbu.empty else p
    share = monthly[['Period', G]].merge(sbu[['Period', G]], on='Period', suffixes=('_partner', '_sbu'))
    share['Partner Share'] = share[G + '_partner'] / share[G + '_sbu'].replace(0, np.nan)
    share.loc[share['Partner Share'] > 1, 'Partner Share'] = np.nan
    shperiod = share.Period.max() if not share.empty else p
    def share_at(period):
        r = period_row(share, period)
        return r['Partner Share'] if r is not None else np.nan
    from presentation_polish import current_coverage_incomplete
    benchmark = d.benchmark_row
    b = lambda key: benchmark.get(key, np.nan) if benchmark is not None else np.nan
    return dict(current_partner_coverage_incomplete=current_coverage_incomplete(result),
        conversion_label=s.conversion_label, PARTNER_SBU=s.partner_sbu, TREND_MONTHS=s.trend_months,
        SBU_CVR=b('% CVR'), SBU_AOV=b('€ AOV Derived'), SBU_PNA=b('% PNA Orders'), SBU_CANCELS=b('% Cancels DTP'),
        UPTIME_TARGET=UPTIME_TARGET, include_uptime=s.include_uptime,
        monthly=monthly, selected_sbu=sbu, latest_partner=latest, latest_sbu=row(sbu, sp),
        partner_latest_period=p, sbu_latest_period=sp, comparison_period=comp,
        partner_comparison=row(monthly, comp), sbu_comparison=row(sbu, comp),
        partner_previous=guarded_previous(d, p), sbu_previous=get_previous_month_row(sbu, sp),
        partner_comp_previous=guarded_previous(d, comp), sbu_comp_previous=get_previous_month_row(sbu, comp),
        sbu_yoy_row=get_yoy_row(sbu, sp), share_data=share, share_period=shperiod,
        latest_share=share_at(shperiod), previous_share=share_at(shperiod-pd.DateOffset(months=1)) if comparison(d,shperiod,shperiod-pd.DateOffset(months=1)).comparable else np.nan,
        share_yoy=share_at(shperiod-pd.DateOffset(years=1)) if comparison(d,shperiod,shperiod-pd.DateOffset(years=1)).comparable else np.nan,
        benchmark_period_label=d.benchmark_period.strftime('%B %Y') if d.benchmark_period is not None else 'benchmark unavailable',
        uptime_coverage_note=f"Uptime coverage: {int(latest.get('Uptime Known Store Count', 0))}/{int(latest.get('Operating Store Count', 0))} operating {'store' if latest.get('Operating Store Count', 0) == 1 else 'stores'}. Missing current store records: {int((result.stores['Data Quality'] == 'Missing current data').sum())}.")


def render(result, output=None):
    import slides_market, slides_partner, slides_stores
    from ppt_core import add_title, add_text, add_footer
    output = Path(output or result.settings.output_file)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.core_properties.version = 'V9_Beta2'
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    ctx = context(result)
    with TemporaryDirectory(prefix='dev61-charts-') as tmp:
        charts = lambda name: str(Path(tmp) / name)
        if not result.data.selected_sbu.empty and set(result.data.monthly.Period) & set(result.data.selected_sbu.Period):
            slides_market.build(prs, ctx, charts)
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            add_title(slide, 'Partner Business Review', result.summary['reporting_period'])
            add_text(slide, 'Market comparison unavailable: no common partner / SBU period.', .75, 2, 11.8, 1, size=20)
            add_footer(slide)
        from beta_presentation import VERSION
        add_text(prs.slides[0], VERSION, .80, 7.03, 11.5, .20, size=9)
        slides_partner.build(prs, ctx, charts)
        slides_stores.build(prs, result, charts)
        add_disclosures(prs, result, ctx)
        for slide in prs.slides:
            if result.settings.cvr_mode == 'delivered_per_session':
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                run.text = run.text.replace('CVR', 'Delivered / sessions')
            slide.notes_slide.notes_text_frame.text = ('[Sources]\nUser-supplied partner_data.xlsx and subbusiness_unit_data.xlsx; '
                'periods and input hashes in accompanying audit export. Brand assets supplied with Dev5.\n[/Sources]\n'
                + result.settings.conversion_description)
        temp = output.with_name(output.stem + '.tmp.pptx')
        try:
            prs.save(temp)
            os.replace(temp, output)
        finally:
            temp.unlink(missing_ok=True)
    return output


def add_coverage_slides(prs, result):
    from ppt_core import add_title, add_text, add_footer
    affected=result.data.coverage[result.data.coverage['Missing Stores'] > 0]
    if affected.empty:
        return
    for start in range(0,len(affected),4):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        pages=(len(affected)+3)//4
        add_title(slide,'Operational & Data Review · Store Coverage'+(f' · {start//4+1}/{pages}' if pages>1 else ''),
                  'Raw totals are retained. Incomplete-period growth comparisons are suppressed.')
        add_footer(slide)
        for i,(_,row) in enumerate(affected.iloc[start:start+4].iterrows()):
            y=1.55+i*1.25
            add_text(slide,row['Period'].strftime('%B %Y')+' · '+coverage_summary(result.data,row['Period']),.7,y,11.9,.35,size=13,bold=True)
            names=row['Missing Store Names']
            if len(names)>210: names=names[:207]+'… (full list in coverage.csv)'
            add_text(slide,'Missing: '+names,.7,y+.45,11.9,.6,size=12)


def add_disclosures(prs, result, ctx):
    from ppt_core import add_text
    from theme import ORANGE, GREY, DARK
    warnings=[]
    if any(i.code=='partial_period' for i in result.data.issues):
        warnings.append('PARTIAL PERIOD — PROVISIONAL ANALYSIS')
    if result.settings.priority_model=='simple_product':
        warnings.append('EXPERIMENTAL / CHALLENGER MODEL')
    period=result.data.reporting_period
    mom=comparison(result.data,period,period-pd.DateOffset(months=1))
    affected=result.data.coverage[result.data.coverage['Missing Stores']>0]
    coverage=coverage_summary(result.data,period)
    if 'Partial Store Coverage' not in coverage and not affected.empty:
        coverage='Historical periods have Partial Store Coverage · See coverage details'
    for slide in prs.slides:
        texts=[x.text for x in slide.shapes if x.has_text_frame and x.text.strip()]
        title=texts[0] if texts else ''
        # Dataset/model qualification appears on every slide, including the cover.
        if warnings:
            from diagnostics import box, AMBER
            box(slide,.63,.01,12.10,.27,AMBER)
            add_text(slide,'  |  '.join(warnings),.65,.035,12,.23,size=10,bold=True,color=DARK)
        if not affected.empty and title not in {'Store Diagnostics','Thank You','Discussion'} and 'Business Development' not in title and not title.endswith('Performance Trend'):
            add_text(slide,coverage,.67,1.22,12,.23,size=10,bold=True,color=ORANGE)
        if mom.reason==NOT_COMPARABLE and title in {'Executive Summary','Operational Health'}:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text=='Comparison unavailable':
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs: run.text='MoM: Not comparable — incomplete store coverage'
        if title.startswith('Growth vs ') and not comparison(result.data,ctx['comparison_period'],ctx['comparison_period']-pd.DateOffset(months=1)).comparable:
            add_text(slide,'Partner MoM: '+NOT_COMPARABLE,.75,6.65,11.8,.4,size=12,bold=True,color=GREY)
        if title.startswith('Partner Share of ') and not affected.empty:
            add_text(slide,'Share comparisons involving incomplete store coverage are not comparable.',.75,6.35,11.8,.3,size=10,color=GREY)
