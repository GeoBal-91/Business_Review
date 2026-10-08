"""Existing Dev5 slide narrative/layout, now invoked with explicit per-run context."""
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from theme import DARK, GREEN, GREY, LIGHT_GREY, LOGO_FILE, MPL_GREEN, MPL_YELLOW, PARTNER_LOGO_FILE, RED, YELLOW
from formatting import benchmark_gap_info, delta_info, format_currency, format_currency_2, format_number, format_percent, get_row_value, pp_change, relative_change, valid_positive_rate
from presentation_hardening import benchmark_display, COVERAGE_BENCHMARK_NOTE
from ppt_core import add_chart_image, add_footer, add_kpi_card, add_picture_contain, add_sbu_icon, add_table_header, add_takeaway, add_text, add_title, save_chart, set_background


def build(prs, ctx, charts):
    chart_path = charts
    PARTNER_SBU = ctx['PARTNER_SBU']
    SBU_AOV = ctx['SBU_AOV']
    SBU_CVR = ctx['SBU_CVR']
    TREND_MONTHS = ctx['TREND_MONTHS']
    benchmark_period_label = ctx['benchmark_period_label']
    comparison_period = ctx['comparison_period']
    latest_partner = ctx['latest_partner']
    latest_sbu = ctx['latest_sbu']
    latest_share = ctx['latest_share']
    monthly = ctx['monthly']
    partner_comp_previous = ctx['partner_comp_previous']
    partner_comparison = ctx['partner_comparison']
    partner_latest_period = ctx['partner_latest_period']
    partner_previous = ctx['partner_previous']
    previous_share = ctx['previous_share']
    sbu_comp_previous = ctx['sbu_comp_previous']
    sbu_comparison = ctx['sbu_comparison']
    sbu_latest_period = ctx['sbu_latest_period']
    sbu_previous = ctx['sbu_previous']
    sbu_yoy_row = ctx['sbu_yoy_row']
    selected_sbu = ctx['selected_sbu']
    share_data = ctx['share_data']
    share_period = ctx['share_period']
    share_yoy = ctx['share_yoy']
    current_coverage_incomplete = ctx['current_partner_coverage_incomplete']
    comparison_coverage_incomplete = current_coverage_incomplete and comparison_period == partner_latest_period


    # ============================================================
    # 28. COVER
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(
        slide,
        YELLOW
    )


    # ============================================================
    # CO-BRANDED LOGO LOCKUP
    # ============================================================

    # Glovo logo
    if os.path.exists(LOGO_FILE):
        add_picture_contain(
            slide,
            LOGO_FILE,
            6.55,   # x
            0.28,   # y
            2.55,   # box width
            0.95    # box height
        )


    # Separator
    separator = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(9.28),
        Inches(0.28),
        Inches(0.025),
        Inches(0.92)
    )

    separator.fill.solid()
    separator.fill.fore_color.rgb = DARK
    separator.line.fill.background()


    # Partner logo
    if os.path.exists(PARTNER_LOGO_FILE):
        add_picture_contain(
            slide,
            PARTNER_LOGO_FILE,
            9.48,   # x
            0.30,   # y
            3.40,   # box width
            0.90    # box height
        )
    else:
        add_text(
            slide,
            "PARTNER LOGO",
            9.48,
            0.58,
            3.40,
            0.30,
            size=9,
            bold=True,
            color=GREY,
            align=PP_ALIGN.CENTER
        )


    # ============================================================
    # COVER TEXT
    # ============================================================

    add_text(
        slide,
        "Partner Business Review",
        0.8,
        2.35,
        11.5,
        0.8,
        size=30,
        bold=True,
        color=DARK
    )

    partner_period_text = (
        f"{monthly['Period'].min().strftime('%B %Y')}"
        f" – "
        f"{partner_latest_period.strftime('%B %Y')}"
    )

    add_text(
        slide,
        partner_period_text,
        0.82,
        3.25,
        10,
        0.5,
        size=16,
        color=DARK
    )

    add_text(
        slide,
        f"Sub-Business Unit: {PARTNER_SBU}",
        0.82,
        3.75,
        10,
        0.4,
        size=12,
        color=DARK
    )

    # 29. SBU BUSINESS DEVELOPMENT
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        f"{PARTNER_SBU} Business Development",
        (
            f"Latest SBU performance — "
            f"{sbu_latest_period.strftime('%B %Y')}"
        )
    )

    add_footer(slide)


    sbu_gmv = latest_sbu["€ GMV"]

    sbu_orders = latest_sbu[
        "# Delivered Orders"
    ]

    sbu_aov = latest_sbu[
        "€ AOV Derived"
    ]

    sbu_cvr = latest_sbu[
        "% CVR"
    ]


    sbu_gmv_mom_change = relative_change(
        sbu_gmv,
        get_row_value(
            sbu_previous,
            "€ GMV"
        )
    )

    sbu_orders_mom_change = relative_change(
        sbu_orders,
        get_row_value(
            sbu_previous,
            "# Delivered Orders"
        )
    )

    sbu_aov_mom_change = relative_change(
        sbu_aov,
        get_row_value(
            sbu_previous,
            "€ AOV Derived"
        )
    )

    sbu_cvr_mom_change = pp_change(
        sbu_cvr,
        get_row_value(
            sbu_previous,
            "% CVR"
        )
    )


    sbu_gmv_mom = delta_info(
        sbu_gmv_mom_change
    )

    sbu_orders_mom = delta_info(
        sbu_orders_mom_change
    )

    sbu_aov_mom = delta_info(
        sbu_aov_mom_change
    )

    sbu_cvr_mom = delta_info(
        sbu_cvr_mom_change,
        metric_type="pp"
    )


    sbu_gmv_yoy = delta_info(
        relative_change(
            sbu_gmv,
            get_row_value(
                sbu_yoy_row,
                "€ GMV"
            )
        ),
        comparison_label="YoY"
    )

    sbu_orders_yoy = delta_info(
        relative_change(
            sbu_orders,
            get_row_value(
                sbu_yoy_row,
                "# Delivered Orders"
            )
        ),
        comparison_label="YoY"
    )

    sbu_aov_yoy = delta_info(
        relative_change(
            sbu_aov,
            get_row_value(
                sbu_yoy_row,
                "€ AOV Derived"
            )
        ),
        comparison_label="YoY"
    )

    sbu_yoy_cvr_value = get_row_value(
        sbu_yoy_row,
        "% CVR"
    )

    if valid_positive_rate(
        sbu_yoy_cvr_value
    ):
        sbu_cvr_yoy = delta_info(
            pp_change(
                sbu_cvr,
                sbu_yoy_cvr_value
            ),
            metric_type="pp",
            comparison_label="YoY"
        )
    else:
        sbu_cvr_yoy = {
            "text": "",
            "color": GREY
        }


    card_x = [
        0.45,
        3.5,
        6.55,
        9.6
    ]

    add_kpi_card(
        slide,
        card_x[0],
        1.45,
        2.85,
        1.75,
        "SBU GMV",
        format_currency(sbu_gmv),
        sbu_gmv_mom["text"],
        sbu_gmv_mom["color"],
        sbu_gmv_yoy["text"],
        sbu_gmv_yoy["color"]
    )

    add_kpi_card(
        slide,
        card_x[1],
        1.45,
        2.85,
        1.75,
        "SBU Orders",
        format_number(sbu_orders),
        sbu_orders_mom["text"],
        sbu_orders_mom["color"],
        sbu_orders_yoy["text"],
        sbu_orders_yoy["color"]
    )

    add_kpi_card(
        slide,
        card_x[2],
        1.45,
        2.85,
        1.75,
        "SBU AOV",
        format_currency_2(sbu_aov),
        sbu_aov_mom["text"],
        sbu_aov_mom["color"],
        sbu_aov_yoy["text"],
        sbu_aov_yoy["color"]
    )

    add_kpi_card(
        slide,
        card_x[3],
        1.45,
        2.85,
        1.75,
        "SBU CVR",
        format_percent(sbu_cvr),
        sbu_cvr_mom["text"],
        sbu_cvr_mom["color"],
        sbu_cvr_yoy["text"],
        sbu_cvr_yoy["color"]
    )


    # Growth Composition

    add_text(
        slide,
        "Growth Composition",
        0.65,
        3.65,
        5.4,
        0.35,
        size=12,
        bold=True
    )

    growth_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.65),
        Inches(4.08),
        Inches(5.6),
        Inches(2.15)
    )

    growth_box.fill.solid()
    growth_box.fill.fore_color.rgb = LIGHT_GREY
    growth_box.line.color.rgb = LIGHT_GREY


    growth_items = [
        (
            "GMV",
            sbu_gmv_mom_change
        ),
        (
            "Orders",
            sbu_orders_mom_change
        ),
        (
            "AOV",
            sbu_aov_mom_change
        )
    ]

    growth_y = 4.35

    for name, change in growth_items:

        if change is None:
            continue

        value_text = f"{change:+.1%}"

        value_color = (
            GREEN
            if change >= 0
            else RED
        )

        add_text(
            slide,
            name,
            0.95,
            growth_y,
            1.35,
            0.35,
            size=10,
            bold=True
        )

        add_text(
            slide,
            value_text,
            2.25,
            growth_y,
            1.1,
            0.35,
            size=14,
            bold=True,
            color=value_color
        )

        bar_x = 3.35
        bar_y = growth_y + 0.08

        max_width = 2.25

        magnitude = min(
            abs(change) / 0.15,
            1
        )

        bar_width = max(
            0.12,
            magnitude * max_width
        )

        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(bar_x),
            Inches(bar_y),
            Inches(bar_width),
            Inches(0.16)
        )

        bar.fill.solid()
        bar.fill.fore_color.rgb = (
            GREEN if change >= 0 else RED
        )

        bar.line.fill.background()

        growth_y += 0.53


    # Key Takeaways

    add_text(
        slide,
        "Key Takeaways",
        6.75,
        3.65,
        5.6,
        0.35,
        size=12,
        bold=True
    )

    takeaway_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.75),
        Inches(4.08),
        Inches(5.9),
        Inches(2.15)
    )

    takeaway_box.fill.solid()
    takeaway_box.fill.fore_color.rgb = LIGHT_GREY
    takeaway_box.line.color.rgb = LIGHT_GREY


    sbu_takeaways = []

    if sbu_gmv_mom_change is not None:
        if sbu_gmv_mom_change >= 0:
            sbu_takeaways.append(
                f"GMV increased "
                f"{abs(sbu_gmv_mom_change):.1%} "
                f"vs previous month."
            )
        else:
            sbu_takeaways.append(
                f"GMV decreased "
                f"{abs(sbu_gmv_mom_change):.1%} "
                f"vs previous month."
            )


    if (
        sbu_orders_mom_change is not None
        and sbu_aov_mom_change is not None
    ):
        if (
            sbu_orders_mom_change > 0
            and sbu_aov_mom_change < 0
        ):
            sbu_takeaways.append(
                "Growth was volume-led: orders increased "
                "while average basket softened."
            )

        elif (
            sbu_orders_mom_change > 0
            and sbu_aov_mom_change > 0
        ):
            sbu_takeaways.append(
                "Both order volume and average basket "
                "supported growth."
            )

        elif (
            sbu_orders_mom_change < 0
            and sbu_aov_mom_change > 0
        ):
            sbu_takeaways.append(
                "Higher basket size partly offset "
                "lower order volume."
            )

        else:
            sbu_takeaways.append(
                "Both order volume and basket size "
                "declined vs previous month."
            )


    if sbu_cvr_mom_change is not None:
        if sbu_cvr_mom_change > 0:
            sbu_takeaways.append(
                f"CVR improved by "
                f"{abs(sbu_cvr_mom_change) * 100:.1f} pp "
                f"to {format_percent(sbu_cvr)}."
            )

        elif sbu_cvr_mom_change < 0:
            sbu_takeaways.append(
                f"CVR declined by "
                f"{abs(sbu_cvr_mom_change) * 100:.1f} pp "
                f"to {format_percent(sbu_cvr)}."
            )


    takeaway_y = 4.37

    for text in sbu_takeaways[:3]:

        add_text(
            slide,
            "• " + text,
            7.05,
            takeaway_y,
            5.25,
            0.45,
            size=9.5
        )

        takeaway_y += 0.55


    # ============================================================
    # 30. SBU PERFORMANCE TREND
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        f"{PARTNER_SBU} Performance Trend",
        f"Last {TREND_MONTHS} available months"
    )

    add_footer(slide)

    sbu_trend = (
        selected_sbu
        .tail(TREND_MONTHS)
        .copy()
    )

    sbu_labels = (
        sbu_trend[
            "Period"
        ]
        .dt.strftime("%b %y")
    )

    fig, ax1 = plt.subplots(
        figsize=(10.5, 4.4)
    )

    ax1.bar(
        sbu_labels,
        sbu_trend["€ GMV"],
        color=MPL_YELLOW,
        alpha=0.85,
        label="GMV (€)"
    )

    ax1.set_ylabel("GMV (€)")
    ax1.spines["top"].set_visible(False)

    ax2 = ax1.twinx()

    ax2.plot(
        sbu_labels,
        sbu_trend[
            "# Delivered Orders"
        ],
        marker="o",
        linewidth=2.5,
        color=MPL_GREEN,
        label="Orders"
    )

    ax2.set_ylabel("Orders")
    ax2.spines["top"].set_visible(False)

    plt.xticks(rotation=35)

    plt.title(
        f"{PARTNER_SBU} — GMV & Orders",
        loc="left",
        fontsize=14,
        fontweight="bold"
    )

    handles1, labels1 = ax1.get_legend_handles_labels()
    handles2, labels2 = ax2.get_legend_handles_labels()

    ax1.legend(
        handles1 + handles2,
        labels1 + labels2,
        loc="upper center",
        bbox_to_anchor=(0.5, 1.02),
        ncol=2,
        frameon=False,
        fontsize=9
    )

    plt.tight_layout()

    chart_file = chart_path("sbu_performance_trend.png")

    save_chart(
        fig,
        chart_file
    )

    add_chart_image(
        slide,
        chart_file,
        0.8,
        1.35,
        11.8,
        4.75
    )

    if len(sbu_trend) >= 2:

        first_row = sbu_trend.iloc[0]
        last_row = sbu_trend.iloc[-1]

        trend_gmv_change = relative_change(
            last_row["€ GMV"],
            first_row["€ GMV"]
        )

        trend_orders_change = relative_change(
            last_row[
                "# Delivered Orders"
            ],
            first_row[
                "# Delivered Orders"
            ]
        )

        if (
            trend_gmv_change is not None
            and trend_orders_change is not None
        ):
            add_takeaway(
                slide,
                (
                    f"Across the displayed period, "
                    f"{PARTNER_SBU} GMV changed "
                    f"{trend_gmv_change:+.1%}, while "
                    f"orders changed "
                    f"{trend_orders_change:+.1%}."
                )
            
            )

            add_sbu_icon(
        slide,
        PARTNER_SBU,
        x=11.2,
        y=0.6,
        w=1.5,
        h=1.5
    )


    # ============================================================
    # 31. PARTNER VS SBU
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        f"Partner vs {PARTNER_SBU}",
        (
            ("Reporting period: " if comparison_coverage_incomplete else "Comparable period: ") +
            f"{comparison_period.strftime('%B %Y')}"
        )
    )

    add_footer(slide)


    comparison_metrics = [
        (
            "AOV",
            "€ AOV Derived",
            "currency",
            True
        ),
        (
            "CVR",
            "% CVR",
            "percent",
            True
        ),
        (
            "Uptime",
            "% Uptime",
            "percent",
            True
        ),
        (
            "PNA",
            "% PNA Orders",
            "percent",
            False
        ),
        (
            "Cancels",
            "% Cancels DTP",
            "percent",
            False
        )
    ]


    headers = [
        "KPI",
        "Partner",
        PARTNER_SBU,
        "Gap"
    ]

    x_positions = [
        0.9,
        4.0,
        6.4,
        8.9
    ]

    widths = [
        2.5,
        1.8,
        1.8,
        3.2
    ]

    add_table_header(
        slide,
        headers,
        x_positions,
        widths,
        1.42,
        height=0.42,
        font_size=14
    )

    y = 2.18

    for (
        metric_name,
        metric_column,
        metric_format,
        higher_is_better
    ) in comparison_metrics:

        partner_value = (
            partner_comparison.get(
                metric_column,
                np.nan
            )
        )

        sbu_value = (
            sbu_comparison.get(
                metric_column,
                np.nan
            )
        )

        if metric_format == "currency":
            partner_display = (
                format_currency_2(
                    partner_value
                )
            )

            sbu_display = (
                format_currency_2(
                    sbu_value
                )
            )

            gap_info = benchmark_display(
                partner_value,
                sbu_value,
                comparison_coverage_incomplete,
                metric_type="relative",
                higher_is_better=
                    higher_is_better,
                label=PARTNER_SBU
            )

        else:
            partner_display = (
                format_percent(
                    partner_value
                )
            )

            sbu_display = (
                format_percent(
                    sbu_value
                )
            )

            gap_info = benchmark_display(
                partner_value,
                sbu_value,
                comparison_coverage_incomplete,
                metric_type="pp",
                higher_is_better=
                    higher_is_better,
                label=PARTNER_SBU
            )

        add_text(
            slide,
            metric_name,
            x_positions[0],
            y,
            widths[0],
            0.42,
            size=13,
            bold=True,
            valign=MSO_ANCHOR.MIDDLE
        )

        add_text(
            slide,
            partner_display,
            x_positions[1],
            y,
            widths[1],
            0.42,
            size=13,
            valign=MSO_ANCHOR.MIDDLE
        )

        add_text(
            slide,
            sbu_display,
            x_positions[2],
            y,
            widths[2],
            0.42,
            size=13,
            valign=MSO_ANCHOR.MIDDLE
        )

        add_text(
            slide,
            gap_info["text"],
            x_positions[3],
            y,
            widths[3],
            0.60 if comparison_coverage_incomplete else 0.42,
            size=11 if comparison_coverage_incomplete else 12,
            bold=True,
            color=gap_info["color"],
            valign=MSO_ANCHOR.MIDDLE
        )

        y += 0.75


    add_takeaway(
        slide,
        (
            f"{COVERAGE_BENCHMARK_NOTE} Raw values use {comparison_period.strftime('%B %Y')}."
            if comparison_coverage_incomplete else
            f"Benchmark comparison uses the latest "
            f"common period available in both datasets: "
            f"{comparison_period.strftime('%B %Y')}."
        )
    )


    # ============================================================
    # 32. GROWTH VS SBU
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        f"Growth vs {PARTNER_SBU}",
        (
            f"MoM growth comparison — "
            f"{comparison_period.strftime('%B %Y')}"
        )
    )

    add_footer(slide)


    growth_metrics = [
        (
            "GMV Growth",
            "€ GMV"
        ),
        (
            "Orders Growth",
            "# Delivered Orders"
        ),
        (
            "AOV Growth",
            "€ AOV Derived"
        )
    ]

    headers = [
        "KPI",
        "Partner",
        PARTNER_SBU,
        "Growth Gap"
    ]

    x_positions = [
        0.9,
        4.0,
        6.4,
        8.9
    ]

    widths = [
        2.5,
        1.8,
        1.8,
        3.2
    ]

    add_table_header(
        slide,
        headers,
        x_positions,
        widths,
        1.42,
        height=0.42,
        font_size=14
    )

    y = 2.2

    for (
        metric_name,
        column
    ) in growth_metrics:

        partner_growth = None
        sbu_growth = None

        if partner_comp_previous is not None:
            partner_growth = relative_change(
                partner_comparison[column],
                partner_comp_previous[column]
            )

        if sbu_comp_previous is not None:
            sbu_growth = relative_change(
                sbu_comparison[column],
                sbu_comp_previous[column]
            )

        if (
            partner_growth is not None
            and sbu_growth is not None
        ):
            growth_gap = (
                partner_growth
                - sbu_growth
            )

            gap_color = (
                GREEN
                if growth_gap >= 0
                else RED
            )

            gap_text = (
                f"{growth_gap * 100:+.1f} pp"
            )

        else:
            gap_text = ""
            gap_color = GREY

        partner_text = (
            f"{partner_growth:+.1%}"
            if partner_growth is not None
            else ""
        )

        sbu_text = (
            f"{sbu_growth:+.1%}"
            if sbu_growth is not None
            else ""
        )

        add_text(
            slide,
            metric_name,
            x_positions[0],
            y,
            widths[0],
            0.42,
            size=13,
            bold=True
        )

        add_text(
            slide,
            partner_text,
            x_positions[1],
            y,
            widths[1],
            0.42,
            size=13
        )

        add_text(
            slide,
            sbu_text,
            x_positions[2],
            y,
            widths[2],
            0.42,
            size=13
        )

        add_text(
            slide,
            gap_text,
            x_positions[3],
            y,
            widths[3],
            0.42,
            size=12,
            bold=True,
            color=gap_color
        )

        y += 0.85


    # ============================================================
    # 33. PARTNER SHARE OF SBU
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        f"Partner Share of {PARTNER_SBU}",
        (
            f"Partner GMV as a share of "
            f"{PARTNER_SBU} GMV — "
            f"{share_period.strftime('%B %Y')}"
        )
    )

    add_footer(slide)


    share_mom = delta_info(
        pp_change(
            latest_share,
            previous_share
        ),
        metric_type="pp",
        higher_is_better=True
    )

    share_yoy_delta = delta_info(
        pp_change(
            latest_share,
            share_yoy
        ),
        metric_type="pp",
        higher_is_better=True,
        comparison_label="YoY"
    )

    add_kpi_card(
        slide,
        0.8,
        1.6,
        3.6,
        2.0,
        f"Share of {PARTNER_SBU} GMV",
        format_percent(latest_share),
        share_mom["text"],
        share_mom["color"],
        share_yoy_delta["text"],
        share_yoy_delta["color"]
    )


    share_chart_data = (
        share_data
        .tail(TREND_MONTHS)
        .copy()
    )

    share_labels = (
        share_chart_data[
            "Period"
        ]
        .dt.strftime("%b %y")
    )

    fig, ax = plt.subplots(
        figsize=(7.2, 3.7)
    )

    x_values = np.arange(
        len(share_chart_data)
    )

    share_values = (
        share_chart_data[
            "Partner Share"
        ] * 100
    )

    ax.plot(
        x_values,
        share_values,
        marker="o",
        linewidth=2.5,
        color=MPL_GREEN,
        label="Partner Share"
    )

    ax.fill_between(
        x_values,
        share_values,
        alpha=0.08,
        color=MPL_GREEN
    )

    ax.set_xticks(x_values)

    ax.set_xticklabels(
        share_labels,
        rotation=35
    )

    ax.set_ylabel(
        f"Share of {PARTNER_SBU} GMV (%)"
    )

    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

    ax.legend(
        frameon=False,
        loc="upper center"
    )

    plt.title(
        "Partner Share Evolution",
        loc="left",
        fontsize=13,
        fontweight="bold"
    )

    plt.tight_layout()

    chart_file = chart_path("partner_sbu_share.png")

    save_chart(
        fig,
        chart_file
    )

    add_chart_image(
        slide,
        chart_file,
        4.8,
        1.45,
        7.6,
        4.6
    )

    if pd.notna(previous_share) and pd.notna(latest_share):

        share_change_pp = (
            latest_share
            - previous_share
        ) * 100

        if share_change_pp > 0:
            share_takeaway = (
                f"Partner gained "
                f"{abs(share_change_pp):.2f} pp "
                f"of {PARTNER_SBU} GMV share "
                f"vs previous month."
            )

        elif share_change_pp < 0:
            share_takeaway = (
                f"Partner lost "
                f"{abs(share_change_pp):.2f} pp "
                f"of {PARTNER_SBU} GMV share "
                f"vs previous month."
            )

        else:
            share_takeaway = (
                f"Partner share of "
                f"{PARTNER_SBU} GMV "
                f"was unchanged vs previous month."
            )

        add_takeaway(
            slide,
            share_takeaway
        )


    # ============================================================
    # 34. EXECUTIVE SUMMARY
    # ============================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    set_background(slide)

    add_title(
        slide,
        "Executive Summary",
        (
            f"Latest partner performance — "
            f"{partner_latest_period.strftime('%B %Y')}"
        )
    )

    add_footer(slide)


    latest_gmv = latest_partner["€ GMV"]

    latest_orders = latest_partner[
        "# Delivered Orders"
    ]

    latest_aov = latest_partner[
        "€ AOV Derived"
    ]

    latest_cvr = latest_partner.get(
        "% CVR",
        np.nan
    )

    previous_gmv = get_row_value(
        partner_previous,
        "€ GMV"
    )

    previous_orders = get_row_value(
        partner_previous,
        "# Delivered Orders"
    )

    previous_aov = get_row_value(
        partner_previous,
        "€ AOV Derived"
    )

    previous_cvr = get_row_value(
        partner_previous,
        "% CVR"
    )


    gmv_delta = delta_info(
        relative_change(
            latest_gmv,
            previous_gmv
        )
    )

    orders_delta = delta_info(
        relative_change(
            latest_orders,
            previous_orders
        )
    )

    aov_delta = delta_info(
        relative_change(
            latest_aov,
            previous_aov
        )
    )

    cvr_delta = delta_info(
        pp_change(
            latest_cvr,
            previous_cvr
        ),
        metric_type="pp"
    )

    aov_gap = benchmark_display(
        latest_aov,
        SBU_AOV,
        current_coverage_incomplete,
        raw_label="SBU AOV: ",
        metric_type="relative",
        higher_is_better=True,
        label=PARTNER_SBU
    )

    cvr_gap = benchmark_display(
        latest_cvr, SBU_CVR, current_coverage_incomplete, raw_label='SBU CVR: ',
        metric_type='pp', higher_is_better=True, label=PARTNER_SBU)


    add_kpi_card(
        slide,
        0.45,
        1.45,
        2.85,
        1.75,
        "GMV",
        format_currency(latest_gmv),
        gmv_delta["text"],
        gmv_delta["color"]
    )

    add_kpi_card(
        slide,
        3.5,
        1.45,
        2.85,
        1.75,
        "Orders",
        format_number(latest_orders),
        orders_delta["text"],
        orders_delta["color"]
    )

    add_kpi_card(
        slide,
        6.55,
        1.45,
        2.85,
        1.75,
        "AOV",
        format_currency_2(latest_aov),
        aov_delta["text"],
        aov_delta["color"],
        aov_gap["text"],
        aov_gap["color"]
    )

    add_kpi_card(
        slide,
        9.6,
        1.45,
        2.85,
        1.75,
        "Partner CVR",
        format_percent(latest_cvr),
        cvr_delta["text"],
        cvr_delta["color"],
        cvr_gap["text"],
        cvr_gap["color"]
    )


    insights = []

    gmv_change = relative_change(
        latest_gmv,
        previous_gmv
    )

    orders_change = relative_change(
        latest_orders,
        previous_orders
    )

    aov_change = relative_change(
        latest_aov,
        previous_aov
    )

    if gmv_change is not None:
        insights.append(
            f"GMV changed "
            f"{gmv_change:+.1%} "
            f"vs previous month."
        )

    if (
        orders_change is not None
        and aov_change is not None
    ):
        if (
            abs(orders_change)
            > abs(aov_change)
        ):
            insights.append(
                "Order volume moved more sharply "
                "than basket size during the month."
            )

        else:
            insights.append(
                "Basket-size movement was larger "
                "than the change in order volume."
            )

    if current_coverage_incomplete:
        from presentation_polish import COVERAGE_CVR_NOTE
        insights.append("Partner vs SBU AOV: " + COVERAGE_BENCHMARK_NOTE)
        insights.append(COVERAGE_CVR_NOTE)
    elif (
        pd.notna(latest_cvr)
        and pd.notna(SBU_CVR)
    ):
        cvr_gap_pp = (
            latest_cvr
            - SBU_CVR
        ) * 100

        if cvr_gap_pp >= 0:
            insights.append(
                f"Partner CVR is "
                f"{abs(cvr_gap_pp):.1f} pp "
                f"better than the selected "
                f"{PARTNER_SBU} benchmark ({benchmark_period_label})."
            )

        else:
            insights.append(
                f"Partner CVR is "
                f"{abs(cvr_gap_pp):.1f} pp "
                f"worse than the selected "
                f"{PARTNER_SBU} benchmark ({benchmark_period_label})."
            )

    y = 3.75

    for insight in insights[:3]:
        add_text(
            slide,
            "• " + insight,
            0.8,
            y,
            11.7,
            0.45,
            size=11.5
        )

        y += 0.62


    # ============================================================
