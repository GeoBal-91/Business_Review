"""Dev6.1.2 presentation qualification and layout; analysis stays read-only."""
import numpy as np
from pptx.dml.color import RGBColor
from pptx.util import Inches
from diagnostics import box, AMBER, BORDER
from formatting import benchmark_gap_info, format_currency_2, format_percent, format_score
from data_pipeline import G
from ppt_core import add_text, add_title, add_footer
from theme import GREEN, GREY, WHITE

COVERAGE_BENCHMARK_NOTE = 'Not assessed — incomplete store coverage.'


def benchmark_display(partner, benchmark, incomplete, *, raw_label=None,
                      metric_type='pp', higher_is_better=True, label='SBU'):
    """Qualify only the displayed aggregate comparison, retaining raw values."""
    if incomplete:
        formatter = format_currency_2 if metric_type == 'relative' else format_percent
        text = raw_label + formatter(benchmark) if raw_label else COVERAGE_BENCHMARK_NOTE
        return {'text': text, 'color': GREY}
    return benchmark_gap_info(partner, benchmark, metric_type=metric_type,
                              higher_is_better=higher_is_better, label=label)


def top_gmv_rows(result):
    """Select positive finite GMV from the current population for this slide only."""
    stores, period = result.stores, result.data.reporting_period
    current_ids = result.data.store_monthly.loc[
        result.data.store_monthly['Period'] == period, 'Store ID']
    eligible = (stores['Store ID'].isin(current_ids)
                & (stores['Period'] == period)
                & (stores['Last Observed Period'] == period)
                & ~stores['Store Status'].astype(str).str.upper().isin(['UNKNOWN', 'MISSING_CURRENT'])
                & (stores['Data Quality'] != 'Missing current data')
                & np.isfinite(stores[G])
                & (stores[G] > 0))
    return stores.loc[eligible].sort_values(
        [G, 'Store ID'], ascending=[False, True], na_position='last').head(6).copy()


def action_focus(result, displayed):
    """Summarize the same displayed action records, without ranking again."""
    if not displayed:
        return None
    store_ids = {a['Store ID'] for a in displayed}
    priorities = [a['Priority'] for a in displayed]
    return {'primary_kpi': displayed[0]['Metric'], 'store_count': len(store_ids),
            'highest_priority': max(priorities), 'lowest_priority': min(priorities),
            'review_count': sum(r['Store ID'] in store_ids for r in result.reviews)}


def recommendation_pages(actions):
    """Display pagination only: up to four together, then three per page."""
    if len(actions) <= 4:
        return [actions[:]]
    return [actions[start:start + 3] for start in range(0, len(actions), 3)]


def matrix_cohort_note(result, plotted):
    """Describe the unchanged plotted cohort and all other known stores."""
    excluded = result.stores.loc[~result.stores['Store ID'].isin(plotted['Store ID'])]
    prefix = f"{len(plotted)} {'store' if len(plotted) == 1 else 'stores'} plotted"
    if excluded.empty:
        return prefix + ' · 0 stores reviewed separately.'
    counts = {}
    for _, row in excluded.iterrows():
        label = row['Store Status']
        if label == 'Active':
            label = 'Active / incomplete' if row['Scoring Status'] != 'Complete' else 'Active / missing matrix inputs'
        counts[label] = counts.get(label, 0) + 1
    if len(counts) == 1:
        label, count = next(iter(counts.items()))
        noun = 'store' if count == 1 else 'stores'
        return f'{prefix} · {count} {label} {noun} reviewed separately.'
    detail = ', '.join(f'{count} {label}' for label, count in sorted(counts.items()))
    return f'{prefix} · {len(excluded)} reviewed separately: {detail}.'


def add_matrix_cohort_note(slides, result, plotted):
    from presentation_polish import MATRIX_FOOTNOTE
    note = matrix_cohort_note(result, plotted)
    for slide in slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text == MATRIX_FOOTNOTE:
                shape.top, shape.height = Inches(6.99), Inches(.27)
        text = add_text(slide, note, .8, 6.69, 11.7, .25, size=9.5, color=GREY)
        text.name = 'Matrix cohort note'


def move_guide_after_cover(prs):
    """Move the existing slide without rebuilding any of its content."""
    def has_title(slide, title):
        return any(s.has_text_frame and s.text == title for s in slide.shapes)
    guide = next((i for i, slide in enumerate(prs.slides) if has_title(slide, 'How to Read This Review')), None)
    cover = next((i for i, slide in enumerate(prs.slides) if has_title(slide, 'Partner Business Review')), None)
    if guide is None or cover is None:
        return
    ids = prs.slides._sldIdLst
    guide_id, cover_id = ids[guide], ids[cover]
    ids.remove(guide_id)
    ids.insert(list(ids).index(cover_id) + 1, guide_id)


def _next_review(slide, y):
    add_text(slide, 'Next review', 9.74, y, 2.65, .3, size=12, bold=True)
    for i, text in enumerate(['Confirm priority stores', 'Assign owner / deadline', 'Track progress next review']):
        yy = y + .40 + i * .29
        add_text(slide, str(i + 1), 9.74, yy, .2, .25, size=10, bold=True, color=GREEN)
        add_text(slide, text, 10.02, yy, 2.40, .27, size=10)


def build_recommendations(prs, result):
    """Keep action order/content and apply only the requested page grouping."""
    actions, settings = result.actions, result.settings
    for displayed in recommendation_pages(actions):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, 'Recommendations',
                  f"{result.summary['reporting_period']} · Actions follow the business KPI selected for each ranked store")
        add_footer(slide)
        box(slide, 9.48, 1.55, 3.19, 5.36, AMBER)
        if not displayed:
            box(slide, .67, 1.55, 8.55, 5.36, WHITE, BORDER)
            box(slide, .92, 2.04, .09, 3.20, AMBER)
            add_text(slide, 'No material business opportunities\nat or above the current threshold.',
                     1.27, 2.28, 7.43, 1.45, size=23, bold=True)
            add_text(slide, 'Review operational/data observations separately.',
                     1.27, 4.00, 7.43, .8, size=16, color=GREY)
            add_text(slide, f'Priority ≥{settings.min_priority:g} remains provisional and requires business calibration.',
                     1.27, 5.77, 7.43, .65, size=11, color=GREY)
            _next_review(slide, 2.04)
            continue
        for i, action in enumerate(displayed):
            compact = len(displayed) == 4
            y = 1.55 + i * (1.34 if compact else 1.79)
            box(slide, .67, y, 8.55, 1.27 if compact else 1.70,
                RGBColor(247, 249, 249) if i % 2 == 0 else WHITE, BORDER)
            header = f"{action['Store']} · {action['Metric']} · Priority {format_score(action['Priority'], settings.min_priority)}"
            add_text(slide, header, .90, y + (.10 if compact else .14), 8.08, .42 if compact else .45,
                     size=12 if compact else 13, bold=True, color=GREEN)
            add_text(slide, action['Issue'], .90, y + (.54 if compact else .64), 8.08, .23,
                     size=10 if compact else 11, color=GREY)
            add_text(slide, action['Action'], .90, y + (.79 if compact else .94), 8.08, .26 if compact else .33,
                     size=11 if compact else 12, bold=True)
            add_text(slide, 'Intended effect: ' + action['Intended Effect'],
                     .90, y + (1.06 if compact else 1.31), 8.08, .21 if compact else .34,
                     size=9.5 if compact else 10.5)
        focus = action_focus(result, displayed)
        add_text(slide, 'Action Focus', 9.74, 1.77, 2.65, .35, size=16, bold=True)
        add_text(slide, 'Primary KPI', 9.74, 2.24, 2.65, .24, size=10, color=GREY)
        add_text(slide, focus['primary_kpi'], 9.74, 2.51, 2.65, .34, size=18, bold=True, color=GREEN)
        add_text(slide, 'Stores requiring action', 9.74, 3.02, 2.65, .24, size=10, color=GREY)
        add_text(slide, str(focus['store_count']), 9.74, 3.29, 2.65, .33, size=17, bold=True)
        add_text(slide, 'Priority range', 9.74, 3.79, 2.65, .24, size=10, color=GREY)
        high = format_score(focus['highest_priority'], settings.min_priority)
        low = format_score(focus['lowest_priority'], settings.min_priority)
        add_text(slide, f'{high} high · {low} low', 9.74, 4.07, 2.65, .30, size=13, bold=True)
        add_text(slide, 'Critical / review items', 9.74, 4.50, 2.65, .24, size=10, color=GREY)
        add_text(slide, f"{focus['review_count']} for these stores", 9.74, 4.77, 2.65, .28, size=12, bold=True)
        add_text(slide, 'Separate from business Priority', 9.74, 5.10, 2.65, .32, size=9.5, color=GREY)
        _next_review(slide, 5.62)
