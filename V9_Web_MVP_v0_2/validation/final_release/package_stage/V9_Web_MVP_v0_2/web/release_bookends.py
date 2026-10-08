"""Final approved cover and closing layouts; display metadata only."""
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from ppt_core import add_text, add_picture_contain, set_background
from theme import ASSETS_DIR, LOGO_FILE, GREEN, DARK, WHITE

CLOSING_YELLOW = RGBColor(255, 214, 54)


def block(slide, name, x, y, w, h, color, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in shape._element.xpath('./p:style/a:effectRef'):
        effect.set('idx', '0')
    if rounded:
        shape.adjustments[0] = .5
    return shape


def text(slide, value, name, x, y, w, h, size, *, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    shape = add_text(slide, value, x, y, w, h, size=size, bold=bold, color=color, align=align)
    shape.name = name
    shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    return shape


def glovo(slide, x, y, w, *, white_pin=False):
    image = add_picture_contain(slide, LOGO_FILE, x, y, w, w * 984 / 3000)
    image.name = 'Bookend Glovo logo'
    if white_pin:
        effect = OxmlElement('a:clrChange')
        effect.set('useA', '1')
        for tag, color in [('clrFrom', 'F1CC37'), ('clrTo', 'FFFFFF')]:
            node = OxmlElement('a:' + tag); rgb = OxmlElement('a:srgbClr'); rgb.set('val', color)
            node.append(rgb); effect.append(node)
        image._element.xpath('.//a:blip')[0].append(effect)
    return image


def build_cover(slide, partner_name, period, sbu):
    import theme
    from web.presentation_update import fit_identity_title
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)
    set_background(slide, WHITE)
    rail = block(slide, 'Cover teal rail', 0, 0, 1.765, 7.5, GREEN)
    rail.fill.gradient()
    rail.fill.gradient_angle = 270
    rail.fill.gradient_stops[0].color.rgb = RGBColor(0, 165, 135)
    rail.fill.gradient_stops[1].color.rgb = RGBColor(0, 183, 152)
    block(slide, 'Cover yellow accent', 1.65, 2.59, .235, 1.53, CLOSING_YELLOW)
    # Ordinary names retain the reference typography; long names wrap into the
    # reserved space above Business Review without moving the rest of the slide.
    long = len(partner_name) > 28
    name = text(slide, partner_name, 'Presentation identity title', 2.21,
                1.95 if long else 2.65, 8.35, 1.48 if long else .78, 44, bold=True)
    if long:
        fit_identity_title(name, partner_name)
        name.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text(slide, 'Business Review', 'Cover review heading', 2.21, 3.49, 8.35, .65, 34)
    text(slide, period, 'Cover reporting period', 2.21, 4.70, 8.35, .45, 20)
    glovo(slide, 8.78, .74, 2.04)
    block(slide, 'Cover logo divider', 11.31, .96, .008, .39, RGBColor(110,110,110))
    logo = add_picture_contain(slide, theme.PARTNER_LOGO_FILE, 11.70, .90, 1.0, .50)
    logo.name = 'Cover uploaded partner logo'
    text(slide, f'Sub-Business Unit: {sbu}', 'Cover SBU', 7.30, 6.82, 5.05, .35, 14, align=PP_ALIGN.RIGHT)


def build_closing(slide, prs, period):
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)
    set_background(slide, CLOSING_YELLOW)
    # Keep the approved transparent PNG intact; native picture effects supply
    # the subtle pale decorative treatment from the final screenshot.
    height = Inches(7.5); width = round(height * 1.5)
    image = slide.shapes.add_picture(str(ASSETS_DIR / 'Thank_You.png'), prs.slide_width - width,
                                    0, width=width, height=height)
    image.name = 'Thank You decorative visual'
    blip = image._element.xpath('.//a:blip')[0]
    duotone = OxmlElement('a:duotone')
    for _ in range(2):
        color = OxmlElement('a:srgbClr'); color.set('val', 'FFFFFF'); duotone.append(color)
    blip.append(duotone)
    opacity = OxmlElement('a:alphaModFix'); opacity.set('amt', '26000'); blip.append(opacity)
    text(slide, 'Thank You', 'Closing main heading', 1.45, 2.51, 7.5, .85, 50, bold=True, color=GREEN)
    text(slide, 'For Your Partnership', 'Closing second heading', 1.45, 3.39, 7.5, .65, 30)
    text(slide, 'Let’s Turn The Agreed Actions\nInto Measurable Progress',
         'Closing message', 1.45, 4.71, 7.8, .85, 20)
    block(slide, 'Closing period pill', 1.45, 6.26, 2.08, .44, GREEN, rounded=True)
    period_shape = text(slide, period, 'Closing reporting period', 1.50, 6.29, 1.98, .35, 17,
                        color=WHITE, align=PP_ALIGN.CENTER)
    period_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    glovo(slide, 8.96, 1.18, 2.0, white_pin=True)
