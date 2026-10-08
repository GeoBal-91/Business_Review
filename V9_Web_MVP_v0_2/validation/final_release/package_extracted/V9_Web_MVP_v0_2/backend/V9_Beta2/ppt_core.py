"""Presentation helpers retained from Dev5, with no import-triggered deck creation."""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from theme import WHITE, YELLOW, GREEN, RED, ORANGE, DARK, GREY, LIGHT_GREY, HEADER_GREY, MID_GREY, FONT, ASSETS_DIR, LOGO_FILE

def set_background(
    slide,
    color=WHITE
):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=12,
    bold=False,
    color=DARK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP
):
    if text is None:
        return None

    text = str(text)

    if text.strip().upper() == "NAN":
        text = "Not assessed"

    if text.strip() == "":
        return None

    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = valign

    p = tf.paragraphs[0]
    p.text = text

    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align

    return box


def add_title(
    slide,
    title,
    subtitle=None
):
    add_text(
        slide,
        title,
        0.65,
        0.35,
        12,
        0.5,
        size=20,
        bold=True
    )

    if subtitle:
        add_text(
            slide,
            subtitle,
            0.67,
            0.9,
            11.8,
            0.35,
            size=10,
            color=GREY
        )


def add_footer(slide):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(7.37),
        Inches(13.333),
        Inches(0.13)
    )

    line.fill.solid()
    line.fill.fore_color.rgb = YELLOW
    line.line.fill.background()
    # The inherited Office shape effect casts a shadow outside a full-bleed footer.
    from pptx.oxml.xmlchemy import OxmlElement
    line._element.spPr.append(OxmlElement("a:effectLst"))
    for effect in line._element.xpath("./p:style/a:effectRef"):
        effect.set("idx", "0")


def add_kpi_card(
    slide,
    x,
    y,
    w,
    h,
    title,
    value,
    delta_text=None,
    delta_color=GREY,
    secondary_text=None,
    secondary_color=GREY
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GREY
    shape.line.color.rgb = LIGHT_GREY

    add_text(
        slide,
        title,
        x + 0.2,
        y + 0.14,
        w - 0.4,
        0.28,
        size=9,
        color=GREY
    )

    add_text(
        slide,
        value,
        x + 0.2,
        y + 0.46,
        w - 0.4,
        0.43,
        size=19 if len(str(value)) <= 16 else 12,
        bold=True
    )

    if delta_text:
        add_text(
            slide,
            delta_text,
            x + 0.2,
            y + 0.95,
            w - 0.4,
            0.27,
            size=8.5,
            bold=True,
            color=delta_color
        )

    if secondary_text:
        add_text(
            slide,
            secondary_text,
            x + 0.2,
            y + 1.26,
            w - 0.4,
            0.27,
            size=8,
            bold=True,
            color=secondary_color
        )


def add_section_card(
    slide,
    x,
    y,
    w,
    h,
    title,
    body
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GREY
    shape.line.color.rgb = LIGHT_GREY

    add_text(
        slide,
        title,
        x + 0.2,
        y + 0.15,
        w - 0.4,
        0.3,
        size=10,
        bold=True
    )

    add_text(
        slide,
        body,
        x + 0.2,
        y + 0.55,
        w - 0.4,
        h - 0.7,
        size=9,
        color=GREY
    )


def add_takeaway(
    slide,
    text
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.75),
        Inches(6.52),
        Inches(11.85),
        Inches(0.55)
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GREY
    shape.line.color.rgb = LIGHT_GREY

    add_text(
        slide,
        text,
        1.0,
        6.67,
        11.3,
        0.25,
        size=9.5,
        bold=True
    )


def save_chart(
    fig,
    filename
):
    fig.savefig(
        filename,
        bbox_inches="tight",
        dpi=180,
        facecolor="white"
    )

    plt.close(fig)


def add_chart_image(
    slide,
    filename,
    x,
    y,
    w,
    h
):
    if os.path.exists(filename):
        slide.shapes.add_picture(
            filename,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h)
        )


# ============================================================
# V9 VISUAL ASSET HELPERS
# ============================================================

VERTICAL_ASSETS = {
    "FOOD": "vertical_food.png",
    "HEALTH_BEAUTY": "vertical_health_beauty.png",
    "HEALTH & BEAUTY": "vertical_health_beauty.png",
    "BEAUTY_PERSONAL_CARE": "vertical_health_beauty.png",
    "SHOPS": "vertical_shops.png",
    "GROCERIES": "vertical_groceries.png",
    "COURIER": "vertical_courier.png",
}

SBU_ASSETS = {
    "FASHION": "sbu_fashion.png",
    "JEWELRY": "sbu_jewelry.png",
    "SERVICES": "sbu_services.png",
    "FLOWERS": "sbu_flowers.png",
    "PETS": "sbu_pets.png",
    "STATIONERY_DIY": "sbu_stationery_diy.png",
    "HOME_DECORATION": "sbu_home_decoration.png",
    "TOYS_GAMES": "sbu_toys_games.png",
    "ELECTRONICS": "sbu_electronics.png",
    "BEAUTY_PERSONAL_CARE": "sbu_beauty_personal_care.png",
    "PHARMACY_OTC": "sbu_pharmacy_otc.png",
    "SEX_SHOP": "sbu_sex_shop.png",
    "BOOKS": "sbu_books.png",
    "VAPES": "sbu_vapes.png",
    "TRADITIONAL_TOBACCO": "sbu_traditional_tobacco.png",
}


def asset_path(filename):
    """Returns an absolute path inside the optional /assets folder."""
    return os.path.join(ASSETS_DIR, filename)


def _existing_asset(filename):
    path = asset_path(filename)
    return path if os.path.exists(path) else None


def add_picture_contain(slide, image_path, x, y, box_w, box_h):
    """Fits an image inside a box without stretching its aspect ratio."""
    if not os.path.exists(image_path):
        return None

    from PIL import Image

    with Image.open(image_path) as img:
        px_w, px_h = img.size

    if not px_w or not px_h:
        return None

    image_ratio = px_w / px_h
    box_ratio = box_w / box_h

    if image_ratio >= box_ratio:
        draw_w = box_w
        draw_h = box_w / image_ratio
    else:
        draw_h = box_h
        draw_w = box_h * image_ratio

    draw_x = x + (box_w - draw_w) / 2
    draw_y = y + (box_h - draw_h) / 2

    return slide.shapes.add_picture(
        str(image_path),
        Inches(draw_x),
        Inches(draw_y),
        width=Inches(draw_w),
        height=Inches(draw_h)
    )


def add_glovo_logo(slide, x, y, w=2.7):
    """Adds the built-in Glovo logo. Falls back to Logo.png if needed."""
    logo = _existing_asset("glovo_logo.png")
    if logo is None and os.path.exists(LOGO_FILE):
        logo = LOGO_FILE

    if logo is None:
        print("WARNING: Glovo logo asset was not found.")
        return None

    return slide.shapes.add_picture(
        logo,
        Inches(x),
        Inches(y),
        width=Inches(w)
    )


def add_vertical_icon(slide, vertical, x, y, w=0.65):
    """Adds a Glovo vertical icon when the requested asset is available."""
    key = str(vertical).strip().upper()
    filename = VERTICAL_ASSETS.get(key)

    if filename is None:
        return None

    path = _existing_asset(filename)
    if path is None:
        print(f"WARNING: Vertical icon not found: {filename}")
        return None

    return slide.shapes.add_picture(
        path,
        Inches(x),
        Inches(y),
        width=Inches(w)
    )


def add_blob_accent(slide, x, y, w=0.8, variant=1):
    """Adds one of six official blob accents. Returns None if unavailable."""
    variant = max(1, min(6, int(variant)))
    path = _existing_asset(f"blob_{variant}.png")
    if path is None:
        return None

    return slide.shapes.add_picture(
        path,
        Inches(x),
        Inches(y),
        width=Inches(w)
    )


def add_arrow_asset(slide, x, y, w=0.8, variant=1):
    """Adds one of six arrow assets."""
    variant = max(1, min(6, int(variant)))
    path = _existing_asset(f"arrow_{variant}.png")
    if path is None:
        return None

    return slide.shapes.add_picture(
        path,
        Inches(x),
        Inches(y),
        width=Inches(w)
    )


def add_phone_mockup(slide, x, y, w=2.2, variant=2):
    """Adds one of the two phone mockup assets."""
    variant = 1 if int(variant) == 1 else 2
    path = _existing_asset(f"phone_mockup_{variant}.png")
    if path is None:
        return None

    return slide.shapes.add_picture(
        path,
        Inches(x),
        Inches(y),
        width=Inches(w)
    )


# ============================================================
# 27. TABLE HELPERS
# ============================================================

def add_table_header(
    slide,
    headers,
    x_positions,
    widths,
    y,
    height=0.5,
    font_size=14
):
    total_x = min(x_positions)

    total_right = max(
        x + w
        for x, w
        in zip(
            x_positions,
            widths
        )
    )

    bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(total_x - 0.12),
        Inches(y - 0.08),
        Inches(
            total_right
            - total_x
            + 0.24
        ),
        Inches(height + 0.15)
    )

    bg.fill.solid()
    bg.fill.fore_color.rgb = HEADER_GREY
    bg.line.color.rgb = HEADER_GREY

    for i, header in enumerate(headers):
        add_text(
            slide,
            header,
            x_positions[i],
            y,
            widths[i],
            height,
            size=font_size,
            bold=True,
            color=DARK,
            valign=MSO_ANCHOR.MIDDLE
        )

    separator = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(total_x),
        Inches(y + height + 0.04),
        Inches(
            total_right
            - total_x
        ),
        Inches(0.015)
    )

    separator.fill.solid()
    separator.fill.fore_color.rgb = MID_GREY
    separator.line.fill.background()

def add_sbu_icon(slide, sbu, x, y, w=1.5, h=1.5):
    key = str(sbu).strip().upper()
    filename = SBU_ASSETS.get(key)

    if filename is None:
        print(f"WARNING: No SBU icon mapping found for {sbu}")
        return None

    path = _existing_asset(filename)

    if path is None:
        print(f"WARNING: SBU icon not found: {filename}")
        return None

    return add_picture_contain(
        slide,
        path,
        x,
        y,
        w,
        h
    )
# ============================================================
