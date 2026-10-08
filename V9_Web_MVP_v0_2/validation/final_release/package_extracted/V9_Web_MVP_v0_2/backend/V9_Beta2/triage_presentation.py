"""Presentation adapter for V2.1. Original Beta slide builders stay byte-identical."""
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import json
import math
import textwrap
from pptx import Presentation
from pptx.util import Inches
from beta_presentation import page, build_methodology, plural, overview_rows, overview_status, link_overview
from diagnostics import box, AMBER, BORDER
from presentation_hardening import _next_review
from formatting import format_score, format_percent, format_currency
from data_pipeline import G, U
from ppt_core import add_text
from theme import GREEN, GREY, WHITE, DARK

VERSION = 'V9_Beta2'
REASONS = {'BASE_PRIORITY': 'Base Priority', 'SINGLE_STORE_OVERRIDE': 'Single-Store Override',
           'MULTI_KPI_ESCALATION': 'Multi-KPI Escalation',
           'SEVERE_PRIMARY_CORROBORATION': 'Severe CVR + Corroboration',
           'BUSINESS_MATERIALITY_OVERRIDE': 'Business Materiality Override'}


def texts(slide):
    return [s.text for s in slide.shapes if s.has_text_frame]


def replace_text(shape, text):
    """Preserve the original run styling while replacing a single text label."""
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ''
    else:
        shape.text = text


def scale_note(row):
    if row['Ultra Low Scale Flag']:
        return 'Ultra-low current scale; human review required. Validate persistence / operational recovery.'
    if row['Low Scale Flag']:
        return 'Low current scale; validate persistence / operational recovery.'
    return 'Human review required; validate supporting evidence.' if row['Human Review Required'] else ''


def action_pages(actions):
    # V2.1 adds reason and causal-evidence metadata. Two complete cards avoid tiny type.
    return [actions[i:i + 2] for i in range(0, len(actions), 2)] or [[]]


def build_actions(prs, triage):
    actions, result = triage.actions, triage.analysis
    pages = action_pages(actions)
    for n, displayed in enumerate(pages, 1):
        title = 'Priority Opportunities & Actions' + (f' · {n}/{len(pages)}' if len(pages) > 1 else '')
        slide = page(prs, title, f"{result.summary['reporting_period']} · V2.1 business triage · Priority ranking remains unchanged")
        box(slide, 9.48, 1.55, 3.19, 5.36, AMBER)
        if not displayed:
            box(slide, .67, 1.55, 8.55, 5.36, WHITE, BORDER)
            add_text(slide, 'No stores classified ACTION\nunder the current V2.1 rules.', 1.03, 2.35, 7.75, 1.1, size=23, bold=True)
            add_text(slide, 'MONITOR, HEALTHY and UNCLASSIFIED remain separate.\nReview operational/data observations and Appendix context.', 1.03, 4.0, 7.75, 1.0, size=16, color=GREY)
            _next_review(slide, 2.04)
            continue
        for i, a in enumerate(displayed):
            single = len(displayed) == 1
            y = 1.55 + i * 2.75
            box(slide, .67, y, 8.55, 5.36 if single else 2.60, WHITE, BORDER)
            shape = add_text(slide, f"#{a['Triage Rank']} · {a['Store']}", .92, y + .12, 8.0, .55 if single else .36,
                             size=18 if single else 12.5, bold=True, color=GREEN)
            shape.name = 'Opportunity store: ' + str(a['Store ID'])
            positions = ([.91, 1.47, 1.91, 2.48, 3.34, 4.30, 4.82] if single else [.54, .83, 1.08, 1.34, 1.77, 2.24, 2.46])
            lines = [f"Priority {format_score(a['Priority'], 5)} · Primary Action KPI: {a['Metric']}",
                     'ACTION · ' + REASONS[a['Action Reason']], a['Issue'],
                     'Likely operational driver: ' + a['Likely Operational Driver'],
                     a['Action'], 'Intended effect: ' + a['Intended Effect'], scale_note(a)]
            heights = [.42, .33, .40, .60, .75, .40, .31] if single else [.23, .19, .20, .37, .40, .18, .13]
            sizes = [21, 14, 15, 13, 17, 13, 10] if single else [12, 10, 10, 9.8, 10.8, 9, 8]
            for j, (line, offset, height, size) in enumerate(zip(lines, positions, heights, sizes)):
                if line:
                    s = add_text(slide, line, .92, y + offset, 8.0, height, size=size, bold=j in {0, 1, 4}, color=GREY if j in {2, 3, 5, 6} else DARK)
                    s.name = f'V2.1 action field {j}: ' + str(a['Store ID'])
        add_text(slide, 'OVERALL ACTION FOCUS', 9.71, 1.78, 2.77, .32, size=13, bold=True)
        total = triage.summary['business_classification_counts']['ACTION']
        add_text(slide, f'{len(actions)} selected of {plural(total, "ACTION store")}', 9.71, 2.15, 2.77, .27, size=9.5, color=GREY)
        priorities = [a['Priority'] for a in actions]
        ids = {a['Store ID'] for a in actions}
        review_count = sum(r['Store ID'] in ids for r in result.reviews)
        labels = [('Primary Action KPI', actions[0]['Metric']), ('Stores Requiring Action', plural(len(ids))),
                  ('Priority Range', f'Highest {format_score(max(priorities), 5)} · Lowest {format_score(min(priorities), 5)}'),
                  ('Critical / Review Items', plural(review_count, 'item'))]
        for i, (label, value) in enumerate(labels):
            yy = 2.55 + .65 * i
            add_text(slide, label, 9.71, yy, 2.77, .20, size=10, color=GREY)
            add_text(slide, value, 9.71, yy + .26, 2.77, .29, size=15 if i < 2 else 11, bold=True)
        add_text(slide, 'Review items are separate from Priority.', 9.71, 5.19, 2.77, .27, size=9, color=GREY)
        _next_review(slide, 5.65)


def build_method(prs, triage):
    slide = build_methodology(prs, triage.analysis)
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.text.startswith('CURRENT BETA THRESHOLD'):
            replace_text(shape, 'BASE ACTION RULE · Priority ≥ 5.0')
        elif shape.text.startswith('The 5.0 threshold'):
            replace_text(shape, 'The 5.0 threshold remains provisional. V2.1 overrides can classify lower-score stores ACTION.')
        elif shape.text.startswith('Assessed, eligible stores'):
            replace_text(shape, 'Priority ranks stores. Separate V2.1 triage assigns ACTION / MONITOR / HEALTHY / UNCLASSIFIED (six ACTION stores shown by default).')
    return slide


def build_plan(prs, triage):
    from slides_stores import table
    table(prs, 'Action Plan', 'Agree owners and dates for the selected V2.1 ACTION stores',
          ['Store / Primary KPI', 'Recommended action', 'Owner / deadline'],
          [[a['Store'] + ' / ' + a['Metric'], a['Action'], 'To agree'] for a in triage.actions],
          [3.6, 6.6, 1.7], per_page=5)


def build_overview(prs, triage):
    result = triage.analysis
    if len(result.stores) < 5:
        return
    meta = triage.stores.set_index('Store ID')
    widths = [3.18, 1.27, 1.29, 1.11, .95, .95, .95, .95, 1.35]
    entries, used, pages = [], 0, []
    for _, row in overview_rows(result).iterrows():
        h = max(.42, .17 * max(len(textwrap.wrap(row['Store Address Address'], 38)), len(textwrap.wrap(overview_status(row), 16))) + .10)
        if entries and (len(entries) >= 10 or used + h > 4.79):
            pages.append(entries); entries, used = [], 0
        entries.append((row, h)); used += h
    if entries:
        pages.append(entries)
    for n, entries in enumerate(pages, 1):
        slide = page(prs, 'Store Overview' + (f' · {n}/{len(pages)}' if len(pages) > 1 else ''),
                     f"{result.summary['reporting_period']} · Operational state and business classification are separate")
        box(slide, .65, 1.57, sum(widths), .39, AMBER)
        x = .65
        for heading, width in zip(['Store', 'Status', 'Business', 'GMV', 'CVR', 'PNA', 'Cancels', 'Uptime', 'Priority'], widths):
            add_text(slide, heading, x + .07, 1.68, width - .12, .22, size=10, bold=True, color=GREEN)
            x += width
        y = 2.03
        for row, h in entries:
            classification = meta.loc[row['Store ID'], 'Business Classification']
            cells = [row['Store Address Address'], overview_status(row), classification, format_currency(row[G]),
                     format_percent(row['% CVR']), format_percent(row['% PNA Orders']), format_percent(row['% Cancels DTP']),
                     format_percent(row[U]), format_score(row['Opportunity Score'], 5)]
            box(slide, .65, y, sum(widths), h - .025, WHITE, BORDER)
            x = .65
            for j, (value, width) in enumerate(zip(cells, widths)):
                shape = add_text(slide, value, x + .07, y + .06, width - .12, h - .08, size=9 if j in {1, 2} else 10,
                                 bold=j in {2, 8}, color=GREEN if j in {0, 2} else DARK)
                shape.name = ('Overview store: ' if j == 0 else f'V2.1 Overview {j}: ') + str(row['Store ID'])
                x += width
            y += h
        add_text(slide, 'Select a store for Diagnostics. ACTION may be below Priority 5.0 through a documented V2.1 override.', .67, 7.01, 12.05, .23, size=9, color=GREY)


def augment_diagnostics(prs, triage):
    diagnostics = [s for s in prs.slides if 'Store Diagnostics' in texts(s)]
    rows = triage.stores.sort_values('Store ID')
    for slide, (_, row) in zip(diagnostics, rows.iterrows()):
        reason = REASONS.get(row['Action Reason'], '')
        label = f"Business: {row['Business Classification']}" + (' · ' + reason if reason else '')
        if row['Business Classification'] == 'ACTION' and scale_note(row):
            label += ' · ' + scale_note(row)
        add_text(slide, label, .65, 1.595, 12.10, .14, size=8.8, bold=True, color=GREEN)
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if shape.text == 'Main Business Action KPI':
                replace_text(shape, 'Primary Action KPI')
            elif shape.text == 'MATERIAL':
                replace_text(shape, 'PRIORITY ≥ 5.0')
            elif shape.text == 'BELOW MAIN THRESHOLD':
                replace_text(shape, 'PRIORITY < 5.0')
            elif shape.text == 'Recommended Action' and row['Business Classification'] != 'UNCLASSIFIED':
                replace_text(shape, 'Likely operational driver: ' + row['Likely Operational Driver'])
            elif shape.text == row['Recommended Focus'] and row['Business Classification'] != 'UNCLASSIFIED':
                replace_text(shape, row['V2.1 Recommended Action'])
        slide.notes_slide.notes_text_frame.text += '\nV2.1 business triage:\n' + json.dumps(
            {k: row[k] for k in ['Business Classification', 'Action Reason', 'Supporting Rule Matches',
                                'Primary Action KPI', 'Likely Operational Driver', 'Operational/Data State']}, ensure_ascii=False)
    link_overview(prs, triage.analysis, diagnostics)


def render(triage, output):
    from presentation import render as render_beta, add_disclosures, context
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix='beta11-render-') as tmp:
        path = render_beta(triage.analysis, Path(tmp) / 'baseline.pptx')
        prs = Presentation(path)
    new_slides = []
    original_slides, original_ids = list(prs.slides), list(prs.slides._sldIdLst)
    replacements, removed = {}, set()
    # Append all replacements before removing anything: python-pptx allocates part
    # names from slide count, so delete-then-add can collide with existing parts.
    for prefix, builder in [('How We Prioritize Opportunities', build_method),
                            ('Priority Opportunities & Actions', build_actions),
                            ('Action Plan', build_plan), ('Store Overview', build_overview)]:
        found = [i for i, slide in enumerate(original_slides) if texts(slide)[0].startswith(prefix)]
        if not found:
            continue
        before = len(prs.slides)
        builder(prs, triage)
        added = list(prs.slides._sldIdLst)[before:]
        new_slides.extend(list(prs.slides)[before:])
        replacements[found[0]] = added
        removed.update(found)
    final_ids = []
    for i, item in enumerate(original_ids):
        final_ids.extend(replacements.get(i, []))
        if i not in removed:
            final_ids.append(item)
        else:
            prs.part.drop_rel(item.rId)
    for item in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(item)
    for item in final_ids:
        prs.slides._sldIdLst.append(item)
    # Disclosure helper runs only on new slides, preserving all original coverage/model warnings.
    add_disclosures(SimpleNamespace(slides=new_slides), triage.analysis, context(triage.analysis))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        run.text = run.text.replace('V9 Beta 1.0', VERSION)
    augment_diagnostics(prs, triage)
    prs.core_properties.version = 'V9_Beta2'
    prs.save(output)
    return output
