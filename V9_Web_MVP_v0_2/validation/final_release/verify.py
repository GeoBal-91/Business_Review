from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches,Cm
from PIL import Image
from collections import Counter
import json,hashlib,ast
root=Path(__file__).resolve().parent;project=root.parents[1];assets=project/'backend/V9_Beta2/assets'
sha=lambda data:hashlib.sha256(data).hexdigest()
review_icons={'Critical observation':'icon_critical_observation.png','Availability review':'icon_availability_review.png','Data quality':'icon_data_quality.png','Opening-hours hygiene':'icon_opening_hours.png','Activity review':'icon_store_review.png','Closure verification':'icon_store_review.png'}
summary={'items to review':(1.1174,1.5174,'icon_data_quality.png'),'stores affected':(5.1534,5.5534,'icon_store_review.png'),'critical observations':(9.1374,9.5374,'icon_critical_observation.png')}
reports=[];review_types=set()
# The external reference already defaults to PETS, while the web bundle defaults
# to JEWELRY. Both runs explicitly receive the identical selected SBU.
def without_default_sbu(path):
 tree=ast.parse(path.read_text())
 tree.body=[n for n in tree.body if not (isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='PARTNER_SBU' for t in n.targets))]
 return ast.dump(tree)
assert without_default_sbu(project/'backend/V9_Beta2/config.py')==without_default_sbu(Path('/Users/giorgibalarjishvili/Desktop/V9_Beta2/config.py'))
for record in json.loads((root/'generation.json').read_text()):
 name=record['case'];prs=Presentation(root/'decks'/(name+'.pptx'));old=Presentation(root/'before_decks'/(name+'.pptx'));assert len(prs.slides)==len(old.slides)
 icon_count=0;unchanged=[];headings=[]
 for n,(before,slide) in enumerate(zip(old.slides,prs.slides),1):
  heading=next((s.text for s in slide.shapes if s.has_text_frame and s.text.strip()),'');headings.append(heading)
  old_heading=next((s.text for s in before.shapes if s.has_text_frame and s.text.strip()),'')
  if n!=1 and heading!='Thank You':
   assert slide._element.xml==before._element.xml,(name,n,'unapproved slide difference')
   assert [(s.name,s.image.blob) for s in slide.shapes if s.shape_type==13]==[(s.name,s.image.blob) for s in before.shapes if s.shape_type==13]
   unchanged.append(n)
  if n==1:
   values={s.name:s for s in slide.shapes};assert values['Presentation identity title'].text==record['partner_name']
   assert values['Cover review heading'].text=='Business Review'
   assert values['Cover reporting period'].text==record['period']
   assert values['Cover SBU'].text=='Sub-Business Unit: '+record['sbu']
   assert Image.open(BytesIO(values['Cover uploaded partner logo'].image.blob)).convert('RGBA').tobytes()==Image.open(project/'backend/V9_Beta2/Partner Logo.png').convert('RGBA').tobytes()
   assert values['Bookend Glovo logo'].image.blob==(project/'backend/V9_Beta2/Logo.png').read_bytes()
   assert str(slide.background.fill.fore_color.rgb)=='FFFFFF'
   assert values['Presentation identity title'].top+values['Presentation identity title'].height<values['Cover review heading'].top
  elif heading=='Thank You':
   txt=[s.text.replace('\x0b','\n') for s in slide.shapes if s.has_text_frame and s.text]
   assert txt==['Thank You','For Your Partnership','Let’s Turn The Agreed Actions\nInto Measurable Progress',record['period']]
   assert str(slide.background.fill.fore_color.rgb)=='FFD636'
   pics={s.name:s for s in slide.shapes if s.shape_type==13};assert set(pics)=={'Thank You decorative visual','Bookend Glovo logo'}
   assert pics['Thank You decorative visual'].image.blob==(assets/'Thank_You.png').read_bytes()
   assert Image.open(BytesIO(pics['Thank You decorative visual'].image.blob)).getchannel('A').getextrema()==(0,255)
  if n==1 or heading=='Thank You':
   for s in slide.shapes:
    assert 0<=s.left and 0<=s.top and s.left+s.width<=prs.slide_width and s.top+s.height<=prs.slide_height,(name,n,s.name)
    if s.shape_type==13:
     w,h=s.image.size;assert abs(s.width/s.height-w/h)<1e-5
     assert not any((s.crop_left,s.crop_top,s.crop_right,s.crop_bottom))
  if ' Share of ' in heading:
   chart=next(s for s in slide.shapes if s.shape_type==13 and s.left==Inches(4.8))
   assert (chart.left,chart.top,chart.width,chart.height)==(Inches(4.8),Inches(1.45),Cm(20),Inches(4.6))
  vals={s.name:s for s in slide.shapes};expected={}
  for sh in slide.shapes:
   if not sh.has_text_frame:continue
   if sh.name.startswith('Plan action ') and sh.name.endswith(' rank'):expected[sh.name[:-5]+' priority icon']='icon_priority_action.png'
   if sh.name.startswith('Plan action ') and sh.name.endswith(' agreed label'):expected[sh.name[:-6]+' icon']='icon_agreed_action.png'
   if sh.name.startswith('Additional action ') and sh.name.endswith('Agreed Action label'):expected[sh.name.removesuffix('Agreed Action label')+' agreed icon']='icon_agreed_action.png'
   if sh.name=='Next check-in label':expected['Next check-in icon']='icon_next_checkin.png'
   if sh.name.startswith('Review ') and sh.name.endswith(' type'):
    expected[sh.name[:-5]+' icon']=review_icons[sh.text];review_types.add(sh.text)
   if sh.name=='Review summary label':
    ix,lx,filename=summary[sh.text];key='Review summary '+sh.text+' icon';expected[key]=filename
    icon=vals[key];assert (icon.left,icon.top,icon.width,icon.height)==(Inches(ix),Inches(1.637),Inches(.34),Inches(.34))
    assert sh.left==Inches(lx)
  for key,filename in expected.items():
   icon=vals[key];assert icon.shape_type==13 and icon.image.blob==(assets/filename).read_bytes()
   assert icon.width==icon.height and not any((icon.crop_left,icon.crop_right,icon.crop_top,icon.crop_bottom));icon_count+=1
  assert sum(s.shape_type==13 and s.name.endswith(' icon') for s in slide.shapes)==len(expected),(name,n)
 assert len([t for s in prs.slides for t in [x.text for x in s.shapes if x.has_text_frame] if t=='Additional agreed actions'])==1
 assert headings.index('Discussion')<headings.index('Thank You')<headings.index('APPENDIX — Store Diagnostics')
 if name=='Empty':assert headings.count('Action Plan')==1
 audit=ZipFile(root/'decks'/(name+'_audit.zip'));files=list((root/'frozen_controls'/name).iterdir())
 assert len(files)==22
 for f in files:
  a,b=f.read_bytes(),audit.read(f.name)
  if f.name=='run_audit.json':
   a,b=json.loads(a),json.loads(b)
   for obj in (a,b):
    for k in ('input_file','output_file','subbusiness_file'):obj['settings'][k]='PATH'
    obj['source_hashes']['config.py']='VERIFIED_DEFAULT_SBU_ONLY_DIFFERENCE'
  assert a==b,(name,f.name,'frozen audit differs')
 reports.append({'case':name,'slides':len(prs.slides),'unchanged_slides':unchanged,'only_cover_and_closing_changed':True,'approved_icon_placements':icon_count,'frozen_audit_files_equal':len(files),'normalized_provenance':['temporary paths','config.py hash: verified default SBU only; explicit identical SBU used']})
assert review_types==set(review_icons)
changed=[]
for directory,files in json.loads((root/'before_hashes.json').read_text()).items():
 for name,checksum in files.items():
  p=Path(directory)/name
  if sha(p.read_bytes())!=checksum:changed.append(str(p.relative_to(project)) if p.is_relative_to(project) else str(p))
assert set(changed)=={'web/presentation_update.py','web/version.py','tests/test_partner_name.py'},changed
(root/'verification.json').write_text(json.dumps({'cases':reports,'review_types':sorted(review_types),'existing_files_changed':changed,'all_backend_and_assets_unchanged':True},indent=2))
print('PASS: only cover/closing differ; all five worker decks preserve other slides, exact share-chart width, icons, pagination, and 110/110 audit files vs frozen engine after documented path/config-hash normalization. All backend/assets unchanged.')
