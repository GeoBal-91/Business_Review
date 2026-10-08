from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from zipfile import ZipFile
import json,hashlib,copy
from lxml import etree
root=Path(__file__).resolve().parent;project=root.parents[1];oldroot=project/'validation/summary_icons/decks'
expected={'items to review':(.870,1.1174,1.5174,'icon_data_quality.png'),'stores affected':(4.880,5.1534,5.5534,'icon_store_review.png'),'critical observations':(8.890,9.1374,9.5374,'icon_critical_observation.png')}
def rect(s):return (s.left,s.top,s.width,s.height)
def wanted(x,y,w,h):return tuple(int(Inches(v)) for v in (x,y,w,h))
def clean(x):
 if isinstance(x,dict):return {k:('PATH' if k in ('input_file','output_file','subbusiness_file') else 'VERSION' if k=='presentation_version' else clean(v)) for k,v in x.items()}
 if isinstance(x,list):return [clean(v) for v in x]
 return x
results=[]
for file in sorted((root/'decks').glob('*.pptx')):
 prs=Presentation(file);old=Presentation(oldroot/file.name);assert len(prs.slides)==len(old.slides);pages=[]
 for i,(a,b) in enumerate(zip(old.slides,prs.slides),1):
  assert len(a.shapes)==len(b.shapes)
  mapping={s.name:s for s in b.shapes}
  if 'Review summary items to review icon' in mapping:
   pages.append(i)
   for label,(nx,ix,lx,png) in expected.items():
    number=mapping['Review summary '+label];icon=mapping['Review summary '+label+' icon'];text=next(s for s in b.shapes if s.name=='Review summary label' and s.text==label)
    assert rect(number)==wanted(nx,1.660,.550,.300),(label,rect(number))
    assert rect(icon)==wanted(ix,1.6370,.340,.340),(label,rect(icon))
    assert rect(text)==wanted(lx,1.730,2.610,.230),(label,rect(text))
    assert icon.shape_type==13 and icon.image.blob==(project/'backend/V9_Beta2/assets'/png).read_bytes()
  for before,after in zip(a.shapes,b.shapes):
   if after.name.startswith('Review summary '):
    after.left=before.left;after.top=before.top
   assert before._element.xml==after._element.xml,(file.name,i,after.name)
 with ZipFile(oldroot/(file.stem+'_audit.zip')) as a,ZipFile(root/'decks'/(file.stem+'_audit.zip')) as b:
  assert set(a.namelist())==set(b.namelist())
  for n in a.namelist():
   x,y=a.read(n),b.read(n)
   if n.endswith('.json'):x,y=clean(json.loads(x)),clean(json.loads(y))
   assert x==y,(file.name,n)
 results.append({'case':file.stem,'exact_coordinate_pages':pages,'slide_count':len(prs.slides),'all_other_shapes_unchanged':True,'audit_entries_equal':len(a.namelist())})
for name,entry in json.loads((project/'validation/icon_integration/transparency_verification.json').read_text()).items():assert hashlib.sha256((project/'backend/V9_Beta2/assets'/name).read_bytes()).hexdigest()==entry['corrected_sha256']
(root/'verification.json').write_text(json.dumps(results,indent=2));print(json.dumps(results,indent=2))
