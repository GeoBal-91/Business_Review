from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
import json, re
from io import BytesIO
root=Path('/Users/giorgibalarjishvili/Desktop/V9_Web_MVP_v0_2/validation/identity_titles')
a=ZipFile(root/'PetFood_BEAUTY.pptx'); b=ZipFile(root/'Long_Unicode_BEAUTY.pptx')
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
slides=sorted((n for n in a.namelist() if re.fullmatch(r'ppt/slides/slide\d+.xml',n)),key=lambda n:int(re.search(r'(\d+)\.xml',n)[1]))
for n in slides:
    trees=[]
    for z in (a,b):
        tree=etree.fromstring(z.read(n))
        for shape in tree.xpath('.//p:sp[p:nvSpPr/p:cNvPr[@name="Presentation identity title"]]',namespaces=ns):
            shape.getparent().remove(shape)
        trees.append(etree.tostring(tree))
    assert trees[0]==trees[1], n
parts=[n for n in a.namelist() if n.startswith(('ppt/charts/','ppt/embeddings/','ppt/media/'))]
for n in parts:
    if n.endswith('.xlsx'):
        za,zb=ZipFile(BytesIO(a.read(n))),ZipFile(BytesIO(b.read(n)))
        def payload(z,p):
            data=z.read(p)
            if p=='docProps/core.xml':
                data=re.sub(rb'(<dcterms:(?:created|modified)[^>]*>)[^<]+',rb'\1TIMESTAMP',data)
            return data
        assert {p:payload(za,p) for p in za.namelist()}=={p:payload(zb,p) for p in zb.namelist()},n
    else:
        assert a.read(n)==b.read(n),n
report={'slides_with_identical_non_identity_content':len(slides),'identical_chart_workbook_media_parts':len(parts),'decks':{}}
for file in [root/'PetFood_BEAUTY.pptx',root/'Long_Unicode_BEAUTY.pptx',root/'Long_Name.pptx']:
    deck=Presentation(file)
    report['decks'][file.name]=[{'slide':i,'title':next((s.text for s in slide.shapes if s.has_text_frame and s.text.strip()),'')} for i,slide in enumerate(deck.slides,1)]
(root/'content_preservation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(f'PASS: all {len(slides)} slide bodies and {len(parts)} chart/workbook/media parts identical across display names.')
